from pathlib import Path
import re
import json
import textwrap

import fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


ROOT = Path(r"D:\Project\tRNA research\biological functions and structural characteristics of tRNA\downloaded_resources")
PDF_DIR = ROOT / "pdfs"
EXTRACTION_DIR = ROOT / "reading_extraction"
OUT_DIR = ROOT / "paper2ppt_output"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "tRNA_literature_Nature_paper2PPT_black_white.pptx"
QA_PATH = OUT_DIR / "QA_report.md"


SLIDE_W = 13.333
SLIDE_H = 7.5
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
GREY = RGBColor(165, 165, 165)
DARK_GREY = RGBColor(38, 38, 38)


PAPERS = [
    {
        "file": "38918637_Structural_basis_of_tRNA_recognition_by_the_m3C_RNA_methyltransferase_METTL6_in_complex_with_Se.pdf",
        "title": "METTL6 通过 SerRS 精准选择 tRNASer 并写入 m3C32",
        "claim": "SerRS 不只是氨酰化酶，而是 METTL6 的 tRNASer 选择因子；METTL6 通过 m3C-RBD 重塑反密码子臂，使 C32 进入催化位点。",
        "figs": [
            "Fig.1 整体复合物：SerRS-tRNA-METTL6 的空间组织",
            "Fig.2 m3C-RBD：夹持反密码子臂",
            "Fig.3 反密码子臂重塑：让 C32/A37 区域可被催化读取",
            "Fig.4 活性中心：C32 与 SAM/SAH 口袋的催化关系",
            "Fig.5 可变臂界面：SerRS 与 METTL6 协同识别 tRNASer",
        ],
    },
    {
        "file": "39075051_Structural_basis_of_tRNA_recognition_by_the_widespread_OB_fold..pdf",
        "title": "OB-fold 蛋白主要读取成熟 tRNA 的 3' CA 末端",
        "claim": "Trbp111 并非主要识别 elbow 或反密码子，而是抓住 tRNA 3' CA；Arc1p 在此基础上加入额外接触增强亲和力。",
        "figs": [
            "Fig.1 Trbp111/相关 OB-fold 与 tRNA 结合能力",
            "Fig.2 共晶结构：Trbp111 二聚体捕获 tRNA 3' 端",
            "Fig.3 突变验证 3' 端接触界面",
            "Fig.4 Arc1p 结构域和点突变影响结合",
            "Fig.5 改变 tRNA 3' 端会削弱 Arc1p 识别",
            "Fig.6 minihelix/microhelix 说明局部末端结构足以贡献识别",
            "Fig.7 OB-fold 核酸识别模式的结构比较",
        ],
    },
    {
        "file": "39111359_A_virally_encoded_tRNA_neutralizes_the_PARIS_antiviral_defence_system..pdf",
        "title": "病毒 tRNALys 变体可中和 PARIS 抗噬菌体防御",
        "claim": "PARIS 感知病毒 Ocr 后释放 AriB，AriB 切割宿主 tRNALys 抑制翻译；T5 编码不易切割的 tRNALys 变体实现逃逸。",
        "figs": [
            "Fig.1 AriA/AriB 组装成大型螺旋桨状免疫复合物",
            "Fig.2 AriA-AriB 界面是防御功能前提",
            "Fig.3 Ocr 触发 AriB 从 PARIS 释放",
            "Fig.4 PARIS 激活导致细胞死亡和翻译抑制",
            "Fig.5 关键证据：PARIS 切 tRNALys，T5 tRNALys 抗切割",
            "Fig.6 PARIS 类系统广泛分化",
        ],
    },
    {
        "file": "40069351_Mechanistic_basis_for_PYROXD1-mediated_protection_of_the_human_tRNA_ligase_complex_against_oxid.pdf",
        "title": "PYROXD1 以氧化还原依赖方式保护 RTCB tRNA 连接酶",
        "claim": "PYROXD1 的 C 端尾部遮挡 RTCB 催化中心；NADH/FAD 状态调控结合与释放，使保护和 Archease 激活可以错时发生。",
        "figs": [
            "Fig.1 RTCB-PYROXD1 复合物整体结构",
            "Fig.2 PYROXD1 CTD 物理遮挡 RTCB 活性中心",
            "Fig.3 FAD/NADH 状态驱动 PYROXD1 变构和 RTCB 招募",
            "Fig.4 定时解离让 Archease 接管并激活 RTCB",
        ],
    },
    {
        "file": "40447571_In_vivo_structure_profiling_reveals_human_cytosolic_and_mitochondrial_tRNA_structurome_and_inte.pdf",
        "title": "DM-DMS-MaPseq 在活细胞中读取 tRNA 结构组和互作组",
        "claim": "体内 tRNA 折叠稳定，但 DMS 可及性与体外差异大，主要反映蛋白/核糖体互作；砷酸盐应激重塑细胞质和线粒体 tRNA 互作。",
        "figs": [
            "Fig.1 方法验证：DM-DMS-MaPseq 捕捉 RNA 可及性",
            "Fig.2 核编码 tRNA 的体内/体外 DMS 图谱差异",
            "Fig.3 DMS 保护信号对应蛋白和核糖体互作",
            "Fig.4 氧化应激改变细胞质 tRNA 互作",
            "Fig.5 线粒体 tRNA 的体内结构图谱",
            "Fig.6 应激下线粒体 tRNA 互作变化",
        ],
    },
    {
        "file": "40563009_Structure_of_fungal_tRNA_ligase_Trl1_with_RNA_reveals_conserved_substrate-binding_principles..pdf",
        "title": "Trl1-LIG 结构揭示 tRNA 剪接末端的保守连接逻辑",
        "claim": "CtTrl1-LIG 与活化 RNA 中间体结构定义了保守 RNA 结合界面，并解释 NTD/CTD 在底物定位、磷酸转移和 2'-P 识别中的分工。",
        "figs": [
            "Fig.1 Trl1 连接反应和 tRNA exon halves 底物",
            "Fig.2 CtTrl1-LIG 与活化 RNA 的结构快照",
            "Fig.3 RNA 结合界面和突变验证",
            "Fig.4 将短 RNA 结构建模到 tRNA 剪接底物",
            "Fig.5 与 T4 Rnl2 比较出保守末端定位逻辑",
            "Fig.6 CTD 影响底物结合和连接效率",
            "Fig.7 保守 Arg 识别 2'-phosphate",
        ],
    },
    {
        "file": "40908366_tRNA_as_an_assembly_chaperone_for_a_macromolecular_transcription-processing_complex..pdf",
        "title": "tRNA 可作为痘病毒 vRNAP 的装配伴侣",
        "claim": "缺少 mcm5s2U34 的 tRNAGln/Arg 被选择性招募，通过非经典反密码子构象桥接转录和 mRNA 加工因子，推动 vRNAP 装配。",
        "figs": [
            "Fig.1 体外重构 vRNAP 装配中间体",
            "Fig.2 tRNA 参与的早期 PIC 结构中间体",
            "Fig.3 tRNA-chaperoned assembly cycle 模型",
            "Fig.4 完整 vRNAP 高分辨率结构",
            "Fig.5 tRNAGln(UUG) 的非经典构象和多点接触",
            "Fig.6 tRNA 修饰状态影响复合物形成",
            "Fig.7 cryo-EM 密度中识别 tRNA 修饰",
            "Fig.8 完整 vRNAP 被整体包装入病毒颗粒",
        ],
    },
    {
        "file": "41407712_AAV-delivered_engineered_suppressor_tRNA_rescues_visual_function_in_mice_with_an_inherited_reti.pdf",
        "title": "AAV 递送 sup-tRNAArg 可恢复 RPE65-R44X 小鼠视觉功能",
        "claim": "工程化 sup-tRNAArg 识别 UGA PTC；scAAV8 递送恢复 RPE65 蛋白、ERG 和视觉行为，且未见明显全局 readthrough 毒性。",
        "figs": [
            "Fig.1 细胞中筛选更强的 sup-tRNAArg readthrough",
            "Fig.2 scAAV8.A4T1 恢复 RPE65 蛋白表达",
            "Fig.3 ERG 和锥细胞保护证明视网膜功能恢复",
            "Fig.4 视觉 cliff 行为改善",
            "Fig.5 H&E、转录组和 ribosome profiling 支持安全性",
        ],
    },
    {
        "file": "41501459_RNA-triggered_Cas12a3_cleaves_tRNA_tails_to_execute_bacterial_immunity..pdf",
        "title": "Cas12a3 被目标 RNA 激活后优先切割 tRNA 3' CCA tail",
        "claim": "Cas12a3 识别 target RNA 后不主要清除 target，而是广泛切 tRNA 保守 3' CCA tail，造成生长停滞和抗噬菌体防御。",
        "figs": [
            "Fig.1 Cas12a3/Cas12a4 是 RNA 激活的 type V nuclease clade",
            "Fig.2 直接证据：Cas12a3 优先切 tRNA 3' CCA tail",
            "Fig.3 四元结构显示 tRNA 被装载到 Cas12a3",
            "Fig.4 比较不同状态解析 tRNA 结合/切割机制",
            "Fig.5 tRNA-like reporter 扩展多重 RNA 检测",
        ],
    },
    {
        "file": "41555020_An_engineered_UGA_suppressor_tRNA_gene_for_disease-agnostic_AAV_delivery..pdf",
        "title": "AAV-NoSTOP(UGA) 将 UGA sup-tRNA 推向泛疾病递送",
        "claim": "NFS context、CuO-CymR 生产抑制系统和载体优化解决 UGA sup-tRNA 包装难题；单次 AAV 给药在两种 LSD 模型中恢复酶活。",
        "figs": [
            "Fig.1 NFS context 增强 UGA-sup-tRNAArg potency",
            "Fig.2 CuO-CymR 在生产阶段压低 sup-tRNA 功能",
            "Fig.3 优化 AAV 产量、纯度和基因组完整性",
            "Fig.4 IduaKI/KI 小鼠中恢复 IDUA 活性",
            "Fig.5 组织表达与 charging efficiency 解释疗效差异",
        ],
    },
    {
        "file": "41807381_Translational_regulation_by_oxidative_desulfuration_of_tRNA_modifications..pdf",
        "title": "氧化脱硫把 tRNA U34 修饰转化为翻译调控开关",
        "claim": "xm5s2U34 在氧化环境中脱硫为 xm5h2U34，削弱部分 tRNA 的氨酰化和 AAR 解码效率，从而重塑氧化压力下的翻译输出。",
        "figs": [
            "Fig.1 2-thiouridine 衍生物及其脱硫产物",
            "Fig.2 spike-in 实验证明脱硫产物可在细胞环境形成",
            "Fig.3 质谱检测细胞质/线粒体 tRNA 的 h2U 衍生物",
            "Fig.4 脱硫修饰降低特定翻译输出",
            "Fig.5 Lys/Gln/Glu 等 tRNA 氨酰化下降",
            "Fig.6 tRNALys 的 AAA/AAG 识别受影响",
            "Fig.7 cryo-EM 解释 mcm5s2U34 与 mcm5h2U34 的解码差异",
            "Fig.8 氧化压力-脱硫修饰-翻译调控模型",
        ],
    },
]


def get_font(size=24, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def normalize_label(label):
    return label.replace("Fig.", "Fig\\.").replace(" ", r"\s*")


def find_main_figure_pages(pdf_path, n_figs):
    doc = fitz.open(pdf_path)
    label_to_pages = {}
    for fig_no in range(1, n_figs + 1):
        patterns = [
            re.compile(rf"(?m)^Fig\.\s*{fig_no}\s*\|"),
            re.compile(rf"(?m)^Figure\s*{fig_no}\s*\|"),
            re.compile(rf"(?m)^Fig\.\s*{fig_no}\b"),
        ]
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text("text") or ""
            if any(p.search(text) for p in patterns):
                pages.append(i)
        if not pages:
            # Fallback: first occurrence anywhere in page text.
            for i, page in enumerate(doc):
                text = page.get_text("text") or ""
                if re.search(rf"Fig\.\s*{fig_no}\b", text):
                    pages.append(i)
                    break
        label_to_pages[f"Fig.{fig_no}"] = pages[:1]
    return label_to_pages


def render_page_thumb(doc, page_index, width_px=640):
    page = doc[page_index]
    pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    ratio = width_px / img.width
    return img.resize((width_px, int(img.height * ratio)), Image.Resampling.LANCZOS)


def make_contact_sheet(paper, asset_path):
    pdf_path = PDF_DIR / paper["file"]
    doc = fitz.open(pdf_path)
    n_figs = len(paper["figs"])
    mapping = find_main_figure_pages(pdf_path, n_figs)

    items = []
    used = set()
    for fig_no in range(1, n_figs + 1):
        pages = mapping.get(f"Fig.{fig_no}", [])
        page_index = pages[0] if pages else min(fig_no - 1, len(doc) - 1)
        key = (fig_no, page_index)
        if key in used:
            continue
        used.add(key)
        items.append((fig_no, page_index))

    cols = 3 if len(items) >= 6 else 2
    tile_w = 500
    label_h = 42
    gap = 18
    thumbs = []
    for fig_no, page_index in items:
        thumb = render_page_thumb(doc, page_index, width_px=tile_w)
        max_h = 520 if len(items) <= 4 else 420
        if thumb.height > max_h:
            ratio = max_h / thumb.height
            thumb = thumb.resize((int(thumb.width * ratio), max_h), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (tile_w, thumb.height + label_h), "white")
        canvas.paste(thumb, ((tile_w - thumb.width) // 2, label_h))
        d = ImageDraw.Draw(canvas)
        d.rectangle((0, 0, tile_w, label_h), fill=(0, 0, 0))
        label = f"Fig. {fig_no} | PDF p.{page_index + 1}"
        d.text((12, 9), label, font=get_font(24), fill=(255, 255, 255))
        thumbs.append(canvas)

    if not thumbs:
        sheet = Image.new("RGB", (1200, 800), "white")
        ImageDraw.Draw(sheet).text((40, 40), "No figure pages found", font=get_font(34), fill=(0, 0, 0))
    else:
        rows = (len(thumbs) + cols - 1) // cols
        row_h = max(t.height for t in thumbs) + gap
        sheet = Image.new("RGB", (cols * tile_w + (cols - 1) * gap, rows * row_h), (245, 245, 245))
        for idx, thumb in enumerate(thumbs):
            x = (idx % cols) * (tile_w + gap)
            y = (idx // cols) * row_h
            sheet.paste(thumb, (x, y))

    asset_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(asset_path, quality=92)
    return {"figure_pages": {k: [p + 1 for p in v] for k, v in mapping.items()}, "asset": str(asset_path)}


def set_slide_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BLACK


def add_textbox(slide, x, y, w, h, text, size=24, bold=False, color=WHITE, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, bullets, size=17, color=WHITE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
        p.space_after = Pt(4)
    return box


def add_picture_fit(slide, image_path, x, y, w, h):
    img = Image.open(image_path)
    iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        pic_w = w
        pic_h = w / img_ratio
    else:
        pic_h = h
        pic_w = h * img_ratio
    left = x + (w - pic_w) / 2
    top = y + (h - pic_h) / 2
    return slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), Inches(pic_w), Inches(pic_h))


def add_thin_rule(slide, y):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(y), Inches(12.1), Inches(0.012))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(90, 90, 90)
    shape.line.fill.background()


def create_deck(manifest):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # 1 cover
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide)
    add_textbox(slide, 0.75, 1.05, 11.8, 1.15, "tRNA 文献阅读汇报", size=46, bold=True)
    add_textbox(slide, 0.78, 2.25, 11.6, 0.6, "功能、结构识别、免疫调控与治疗工程", size=26, color=GREY)
    add_textbox(slide, 0.78, 5.85, 11.6, 0.7, "11 篇文献 | 图文对应 | 黑底白字简版", size=20, color=GREY)

    # 2 argument map
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide)
    add_textbox(slide, 0.6, 0.45, 12.2, 0.55, "总论点：tRNA 是可编程的结构平台，不只是翻译适配器", size=28, bold=True)
    add_thin_rule(slide, 1.15)
    add_bullets(slide, 0.8, 1.55, 5.6, 4.9, [
        "结构识别：不同蛋白读取 3' CA、CCA tail、反密码子臂、可变臂或剪接末端。",
        "动态修饰：m3C32、U34 硫修饰/脱硫、缺失 mcm5s2U34 都会改变 tRNA 的功能命运。",
        "免疫节点：PARIS 和 Cas12a3 通过切割 tRNA 快速压低翻译。",
        "治疗载体：工程化 suppressor tRNA 已能通过 AAV 在体内恢复蛋白功能。",
    ], size=21)
    add_bullets(slide, 7.0, 1.55, 5.5, 4.9, [
        "结构论文：METTL6、OB-fold、Trl1、PYROXD1、Cas12a3、vRNAP。",
        "组学论文：DM-DMS-MaPseq 给出体内 tRNA 结构/互作图谱。",
        "功能论文：PARIS、Cas12a3、氧化脱硫解释翻译调控和抗病毒。",
        "转化论文：两篇 AAV/sup-tRNA 证明递送、表达、charging 和安全性是关键。",
    ], size=21)

    # paper slides
    for i, paper in enumerate(PAPERS, start=1):
        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide)
        add_textbox(slide, 0.45, 0.25, 12.4, 0.55, f"{i}. {paper['title']}", size=23, bold=True)
        add_thin_rule(slide, 0.92)
        asset = Path(manifest[paper["file"]]["asset"])
        add_picture_fit(slide, asset, 0.45, 1.12, 7.55, 5.8)
        add_textbox(slide, 8.25, 1.12, 4.55, 1.35, paper["claim"], size=18, bold=False)
        add_textbox(slide, 8.25, 2.58, 4.55, 0.34, "图像对应", size=17, bold=True, color=GREY)
        add_bullets(slide, 8.25, 2.95, 4.55, 3.95, paper["figs"], size=12 if len(paper["figs"]) > 6 else 13)
        add_textbox(slide, 0.5, 6.95, 12.0, 0.28, "左侧为该论文主文图所在 PDF 页面缩略图；右侧为每张主文图的阅读结论。", size=10, color=GREY)

    # thematic synthesis slides
    synthesis = [
        (
            "结构识别：tRNA 的不同局部标志被不同机器读取",
            [
                "METTL6：C32/反密码子臂 + SerRS 介导的 tRNASer 选择。",
                "OB-fold：成熟 tRNA 的 3' CA 是通用成熟标志。",
                "Trl1：剪接 exon halves 的末端化学状态和 2'-P 决定连接。",
                "Cas12a3：受体茎、T 臂与 3' CCA tail 被装载到 RuvC 位点。",
            ],
        ),
        (
            "免疫与病毒：切 tRNA 是快速关闭翻译的有效策略",
            [
                "PARIS：AriA 感知 Ocr 后释放 AriB，切 tRNALys。",
                "T5：用不易切割的 tRNALys 变体逃逸 PARIS。",
                "Cas12a3：目标 RNA 激活后切多种 tRNA 的 3' CCA tail。",
                "共同逻辑：不需要降解所有 RNA，只要打断 tRNA 池即可压低翻译。",
            ],
        ),
        (
            "修饰与应激：tRNA 修饰把环境信号转成翻译变化",
            [
                "m3C32：影响反密码子环结构和翻译保真度。",
                "U34 硫修饰：提升 AAR 解码和氨酰化；氧化脱硫后效率下降。",
                "vRNAP：选择缺少 mcm5s2U34 的 tRNAGln/Arg 作装配伴侣。",
                "DM-DMS-MaPseq：可在体内读出应激下 tRNA 互作变化。",
            ],
        ),
        (
            "治疗工程：suppressor tRNA 的瓶颈从 readthrough 转向递送与安全",
            [
                "视网膜模型：scAAV8.sup-tRNAArg 恢复 RPE65 和视觉功能。",
                "泛疾病 UGA 平台：NFS + CuO-CymR 改善 AAV 生产和体内活性。",
                "关键变量：组织表达、charging efficiency、正常终止密码子 readthrough。",
                "下一步问题：长期安全性、组织特异递送、不同 PTC context 的可预测性。",
            ],
        ),
    ]
    for title, bullets in synthesis:
        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide)
        add_textbox(slide, 0.7, 0.65, 11.9, 0.75, title, size=30, bold=True)
        add_thin_rule(slide, 1.55)
        add_bullets(slide, 1.05, 2.0, 11.1, 3.8, bullets, size=24)

    # final
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide)
    add_textbox(slide, 0.72, 0.7, 12.0, 0.75, "结论：tRNA 是结构、调控和工程化治疗的交汇点", size=31, bold=True)
    add_thin_rule(slide, 1.6)
    add_bullets(slide, 1.0, 2.1, 11.3, 3.6, [
        "看结构：tRNA 局部构象、末端和修饰是蛋白识别的主要信息层。",
        "看功能：免疫系统和病毒都把 tRNA 当作翻译控制的高杠杆节点。",
        "看应用：suppressor tRNA 已有体内疗效证据，但递送、charging 和安全性决定转化上限。",
        "最值得深挖的方向：体内 tRNA structurome 如何预测工程 tRNA 的功能和副作用。",
    ], size=24)
    add_textbox(slide, 0.92, 6.25, 11.6, 0.42, "建议精读顺序：METTL6 / Cas12a3 / 氧化脱硫 / AAV suppressor tRNA", size=19, color=GREY)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_PATH)
    return len(prs.slides)


def validate_pptx(expected_slides):
    prs = Presentation(PPTX_PATH)
    image_count = sum(1 for slide in prs.slides for shape in slide.shapes if getattr(shape, "shape_type", None) == 13)
    return {"slides": len(prs.slides), "expected_slides": expected_slides, "images": image_count}


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    manifest = {}
    for idx, paper in enumerate(PAPERS, start=1):
        safe = re.sub(r"[^A-Za-z0-9]+", "_", paper["file"].split("_", 1)[0] + "_" + paper["title"])[:80]
        asset_path = ASSET_DIR / f"{idx:02d}_{safe}_figures.jpg"
        manifest[paper["file"]] = make_contact_sheet(paper, asset_path)

    slide_count = create_deck(manifest)
    qa = validate_pptx(slide_count)

    manifest_path = OUT_DIR / "asset_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    QA_PATH.write_text(
        "\n".join([
            "# QA report",
            "",
            f"- PPTX: `{PPTX_PATH}`",
            f"- Slides expected/generated: {qa['expected_slides']} / {qa['slides']}",
            f"- Embedded image shapes: {qa['images']}",
            f"- Figure contact sheets: {len(manifest)}",
            "- Style: black background, white text, minimal decoration.",
            "- Note: figure assets are rendered from PDF pages containing main figure labels, so some thumbnails include nearby article text/captions.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps({"pptx": str(PPTX_PATH), "qa": qa, "manifest": str(manifest_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
