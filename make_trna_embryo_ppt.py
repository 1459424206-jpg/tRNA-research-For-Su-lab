from __future__ import annotations

import json
import math
import shutil
import textwrap
from dataclasses import dataclass
from pathlib import Path

import fitz
import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(r"D:\Project\tRNA research\tRNA-related wet-lab experiments")
OUT = ROOT / "output"
ASSETS = OUT / "assets" / "figures"
LOCAL_FIGS = ROOT / "downloaded_resources" / "figures"
PDFS = ROOT / "downloaded_resources" / "pdfs"

W, H = Inches(13.333), Inches(7.5)

INK = RGBColor(31, 41, 55)
MUTED = RGBColor(99, 113, 128)
LIGHT = RGBColor(246, 248, 250)
LINE = RGBColor(213, 221, 230)
GREEN = RGBColor(35, 122, 92)
TEAL = RGBColor(34, 124, 157)
BLUE = RGBColor(49, 91, 171)
GOLD = RGBColor(161, 112, 38)
RED = RGBColor(174, 67, 67)


@dataclass
class Paper:
    title: str
    pmid: str
    journal: str
    year: str
    role: str
    hypothesis: str
    methods: str
    results: str
    discussion: str
    embryo_angle: str
    source: str
    image: str | None = None
    schematic: str = "axis"


def ensure_dirs() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)


def clean_name(s: str) -> str:
    keep = "".join(c if c.isalnum() else "_" for c in s)
    return "_".join(keep.split("_"))[:80]


def download(url: str, target: Path) -> Path | None:
    if target.exists() and target.stat().st_size > 1000:
        return target
    try:
        r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=45)
        if r.ok and len(r.content) > 1000:
            target.write_bytes(r.content)
            return target
    except Exception:
        pass
    return None


def render_pdf_page(pdf: Path, page_index: int, target: Path, zoom: float = 2.0) -> Path | None:
    try:
        doc = fitz.open(pdf)
        page_index = max(0, min(page_index, doc.page_count - 1))
        pix = doc[page_index].get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        pix.save(target)
        return target
    except Exception:
        return None


def image_path(name: str) -> str | None:
    candidates = list(LOCAL_FIGS.glob(name))
    if candidates:
        target = ASSETS / candidates[0].name
        if not target.exists():
            shutil.copy2(candidates[0], target)
        return str(target)
    return None


def prepare_assets() -> dict[str, str]:
    assets: dict[str, str] = {}
    mappings = {
        "nano_trnaseq": "PMC10791586_41587_2023_1743_Fig1_HTML.jpg",
        "trmt6_hsc": "PMC11231335_41467_2024_50110_Fig3_HTML.jpg",
        "dm_dms": "PMC12125218_41467_2025_59435_Fig1_HTML.jpg",
        "aa_trnaseq": "PMC12368100_41467_2025_62545_Fig1_HTML.jpg",
        "cas12a3": "PMC12851939_41586_2025_9852_Fig1_HTML.jpg",
        "desulfuration": "PMC12976133_41467_2026_70126_Fig1_HTML.jpg",
    }
    for key, pattern in mappings.items():
        p = image_path(pattern)
        if p:
            assets[key] = p

    mim = download(
        "https://cdn.ncbi.nlm.nih.gov/pmc/blobs/7282/8062790/9413dfccc665/gr1.jpg",
        ASSETS / "mim_tRNAseq_gr1.jpg",
    )
    if mim:
        assets["mim_trnaseq"] = str(mim)

    mettl1 = download(
        "https://cdn.ncbi.nlm.nih.gov/pmc/blobs/4be4/11925124/d7fdd6a58501/nihms-2053450-f0006.jpg",
        ASSETS / "mettl1_model.jpg",
    )
    if mettl1:
        assets["mettl1_model"] = str(mettl1)

    oracle_pdf = download(
        "https://www.nature.com/articles/s41467-026-72603-5_reference.pdf",
        ASSETS / "oracle_mouse_early_embryogenesis.pdf",
    )
    if oracle_pdf:
        rendered = render_pdf_page(oracle_pdf, 1, ASSETS / "oracle_mouse_page2.png", 2.0)
        if rendered:
            assets["oracle_mouse"] = str(rendered)

    uga_pdf = next(PDFS.glob("41555020_*.pdf"), None)
    if uga_pdf:
        rendered = render_pdf_page(uga_pdf, 2, ASSETS / "uga_suppressor_page3.png", 2.0)
        if rendered:
            assets["uga_suppressor"] = str(rendered)
    return assets


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = LINE


def set_textbox(shape, text, size=13, bold=False, color=INK, align=None, font="Microsoft YaHei"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align


def add_text(slide, x, y, w, h, text, size=13, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    set_textbox(box, text, size=size, bold=bold, color=color, align=align)
    box.line.fill.background()
    box.fill.background()
    return box


def add_footer(slide, source: str):
    add_text(slide, Inches(0.45), Inches(7.05), Inches(12.4), Inches(0.22), source, size=7.5, color=MUTED)


def add_title(slide, title: str, kicker: str | None = None):
    add_text(slide, Inches(0.45), Inches(0.22), Inches(11.9), Inches(0.5), title, size=24, bold=True, color=INK)
    if kicker:
        add_text(slide, Inches(0.48), Inches(0.75), Inches(11.6), Inches(0.25), kicker, size=9, color=MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(1.02), Inches(12.45), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def add_bullets(slide, x, y, w, h, bullets, size=12):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.space_after = Pt(5)
    box.fill.background()
    box.line.fill.background()
    return box


def fit_picture(slide, path: str, x, y, w, h):
    try:
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        px = x + (w - pw) / 2
        py = y + (h - ph) / 2
        pic = slide.shapes.add_picture(path, px, py, width=pw, height=ph)
        return pic
    except Exception:
        return None


def add_takeaway(slide, text: str):
    band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.35), Inches(12.0), Inches(0.52))
    set_fill(band, RGBColor(236, 245, 242))
    band.line.color.rgb = RGBColor(185, 212, 201)
    set_textbox(band, "Take-home: " + text, size=11.5, bold=True, color=GREEN)


def add_schematic(slide, paper: Paper, x, y, w, h):
    colors = [RGBColor(229, 241, 247), RGBColor(241, 238, 249), RGBColor(239, 247, 239), RGBColor(250, 242, 229)]
    labels = [
        ("假说", paper.hypothesis),
        ("湿实验路线", paper.methods),
        ("关键结果", paper.results),
        ("胚胎启发", paper.embryo_angle),
    ]
    gap = Inches(0.12)
    bw = (w - gap * 3) / 4
    for i, (head, body) in enumerate(labels):
        bx = x + i * (bw + gap)
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, bw, h)
        set_fill(sh, colors[i])
        sh.line.color.rgb = RGBColor(190, 203, 214)
        set_textbox(sh, f"{head}\n{body}", size=10.5, bold=False, color=INK, align=PP_ALIGN.CENTER)
        if i < 3:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, bx + bw - Inches(0.05), y + h / 2 - Inches(0.12), Inches(0.22), Inches(0.24))
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(122, 136, 150)
            arr.line.fill.background()


def add_axis_map(slide, x, y, w, h):
    axes = [
        ("tRNA pool", "丰度 / isodecoder"),
        ("Modification", "m1A / m7G / Ψ / xm5s2U"),
        ("Charging", "aminoacylation"),
        ("Structure", "in vivo folding / interactome"),
        ("Fragments", "tRF / tsRNA"),
        ("Translation", "codon demand / MZT"),
    ]
    cx, cy = x + w / 2, y + h / 2
    r = min(w, h) * 0.32
    center = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - Inches(0.85), cy - Inches(0.42), Inches(1.7), Inches(0.84))
    set_fill(center, RGBColor(235, 246, 241))
    set_textbox(center, "早期胚胎\n翻译重编程", size=12, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    for idx, (a, b) in enumerate(axes):
        ang = 2 * math.pi * idx / len(axes) - math.pi / 2
        bx = cx + r * math.cos(ang) - Inches(1.05)
        by = cy + r * math.sin(ang) - Inches(0.34)
        node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.1), Inches(0.68))
        set_fill(node, RGBColor(247, 249, 251))
        set_textbox(node, f"{a}\n{b}", size=9.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
        line = slide.shapes.add_connector(1, cx, cy, bx + Inches(1.05), by + Inches(0.34))
        line.line.color.rgb = RGBColor(189, 199, 208)
        line.line.width = Pt(1)


def add_paper_slide(prs: Presentation, paper: Paper):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, paper.role, f"{paper.journal}, {paper.year} | PMID {paper.pmid} | {paper.title}")
    visual_x, visual_y, visual_w, visual_h = Inches(0.55), Inches(1.25), Inches(7.45), Inches(4.85)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, visual_x, visual_y, visual_w, visual_h)
    set_fill(frame, RGBColor(252, 253, 254))
    frame.line.color.rgb = RGBColor(205, 215, 224)
    if paper.image and Path(paper.image).exists():
        fit_picture(slide, paper.image, visual_x + Inches(0.12), visual_y + Inches(0.12), visual_w - Inches(0.24), visual_h - Inches(0.24))
    else:
        add_schematic(slide, paper, visual_x + Inches(0.18), visual_y + Inches(1.55), visual_w - Inches(0.36), Inches(1.45))
        add_text(slide, visual_x + Inches(0.35), visual_y + Inches(0.45), visual_w - Inches(0.7), Inches(0.6),
                 "精读示意图：从科学问题到可转化湿实验读出", size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.72), Inches(6.08), Inches(7.1), Inches(0.22), "图：原文关键图/流程图；无开放图时为按论文逻辑重绘的实验示意", size=8, color=MUTED)

    rail = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(1.25), Inches(4.55), Inches(4.85))
    set_fill(rail, RGBColor(248, 250, 252))
    rail.line.color.rgb = RGBColor(214, 223, 232)
    add_bullets(
        slide,
        Inches(8.45),
        Inches(1.42),
        Inches(4.15),
        Inches(4.45),
        [
            f"逻辑 / 假说：{paper.hypothesis}",
            f"方法步骤：{paper.methods}",
            f"关键结果：{paper.results}",
            f"讨论与局限：{paper.discussion}",
            f"对早期胚胎研究的用途：{paper.embryo_angle}",
        ],
        size=10.8,
    )
    add_takeaway(slide, paper.embryo_angle)
    add_footer(slide, paper.source)


def add_section_slide(prs: Presentation, title: str, subtitle: str, points: list[str], color=GREEN):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(249, 251, 250)
    add_text(slide, Inches(0.75), Inches(1.15), Inches(11.8), Inches(0.75), title, size=30, bold=True, color=color)
    add_text(slide, Inches(0.78), Inches(1.92), Inches(10.8), Inches(0.45), subtitle, size=13, color=MUTED)
    add_axis_map(slide, Inches(0.7), Inches(2.55), Inches(5.1), Inches(3.7))
    add_bullets(slide, Inches(6.25), Inches(2.65), Inches(5.8), Inches(3.35), points, size=15)
    add_footer(slide, "来源：本报告综合整理自入选文献；后续逐篇页面列出 PMID 和标题")


def add_matrix_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "文献重筛选：主线保留“湿实验 + 胚胎可迁移性”", "本页说明哪些文献进入逐篇精读，哪些仅作为背景参考")
    rows = [
        ("主线新增", "42062277, 39402326, 38989621, 34797706, 41042673, 26721685, 26721680, 29628141, 29983320", "直接涉及 oocyte / MZT / zygote / preimplantation / ESC 的 tRNA 或 tRF 湿实验"),
        ("本地保留", "37024678, 33581077, 40835813, 40447571, 39096899", "可作为低输入胚胎 tRNA-seq、修饰、charging、结构和小 RNA 捕获的方法工具箱"),
        ("本地保留", "38977676, 39892392, 41807381, 41501459, 35322228, 41555020, 34174184, 37944512", "机制或工程扰动模型，可转化为胚胎中 tRNA 翻译调控的因果验证思路"),
        ("降为背景", "35032425, 37173525", "综述 / commentary，不作为逐篇精读主线；用于方法背景和概念框架"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.62), Inches(1.35), Inches(12.1), Inches(4.25)).table
    table.columns[0].width = Inches(1.6)
    table.columns[1].width = Inches(4.45)
    table.columns[2].width = Inches(6.05)
    for j, htxt in enumerate(["处理", "PMID", "理由"]):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(229, 239, 235)
        cell.text = htxt
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        cell.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    for i, row in enumerate(rows, 1):
        for j, txt in enumerate(row):
            cell = table.cell(i, j)
            cell.text = txt
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(9.5 if j != 2 else 10)
                p.font.color.rgb = INK
    add_takeaway(slide, "主线不是简单按文件名排序，而是围绕“早期胚胎翻译重编程可如何被 tRNA 湿实验解析”。")
    add_footer(slide, "来源：本地文件夹题录 + PubMed 复核；每篇精读页均列出 PMID 和标题")


def add_cover(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 250)
    add_text(slide, Inches(0.7), Inches(0.7), Inches(11.8), Inches(0.55), "tRNA 相关湿实验与早期胚胎发育研究", size=29, bold=True, color=INK)
    add_text(slide, Inches(0.72), Inches(1.35), Inches(11.3), Inches(0.4), "文献精读、方法整合与可执行实验路线", size=16, color=GREEN)
    add_axis_map(slide, Inches(0.85), Inches(2.1), Inches(6.2), Inches(4.4))
    add_bullets(
        slide,
        Inches(7.45),
        Inches(2.25),
        Inches(4.8),
        Inches(3.55),
        [
            "核心问题：MZT/ZGA 期间，tRNA 丰度、修饰、charging 与 tRF 是否共同塑造翻译选择性？",
            "文献范围：本地 tRNA 湿实验文献 + 新增早期胚胎/生殖发育核心研究。",
            "汇报结构：筛选逻辑 → 逐篇精读 → 方法矩阵 → 胚胎研究方案。",
        ],
        size=14,
    )
    add_footer(slide, "汇报用途：正式组会/课题设计讨论；每个文献页均注明 PMID 与标题")


def add_question_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "研究主线：把 tRNA 从“翻译适配器”推进到胚胎命运调控变量", "围绕 oocyte-to-embryo transition / MZT / ZGA 组织证据")
    add_axis_map(slide, Inches(0.55), Inches(1.45), Inches(5.6), Inches(4.4))
    add_bullets(
        slide,
        Inches(6.55),
        Inches(1.55),
        Inches(5.8),
        Inches(4.15),
        [
            "母源 mRNA 清除、ZGA 与细胞周期加速要求快速重排翻译供给。",
            "tRNA 层面可拆成 5 个湿实验读出：丰度、修饰、charging、结构/互作、tRNA-derived fragments。",
            "早期胚胎样本量极低，核心技术瓶颈是低输入、修饰阻断、异构体分辨率和功能扰动。",
            "汇报重点不是泛讲 tRNA，而是判断哪些技术和机制能直接转化为胚胎实验。",
        ],
        size=14,
    )
    add_takeaway(slide, "一个可检验假说：胚胎并非被动继承母源 tRNA，而是在 ZGA 前后主动建立与翻译需求匹配的 tRNA 层调控。")
    add_footer(slide, "来源：本报告综合整理自所有入选文献")


def build_papers(assets: dict[str, str]) -> list[Paper]:
    return [
        Paper(
            "Deciphering tRNA repertoires and translation coordination during mouse early embryogenesis by ORACLE-tRNAseq.",
            "42062277",
            "Nature Communications",
            "2026",
            "ORACLE-tRNAseq 让 mouse early embryo 的 tRNA landscape 可被低输入解析",
            "MZT 期间是否存在从母源到胚胎型 tRNA repertoires 的主动切换，并与 ZGA 和翻译效率耦合。",
            "建立低输入 ORACLE-tRNAseq，从少量 mouse oocytes/embryos 获取 tRNA profiles；结合 Ribo-seq、多组学和染色质信息。",
            "oocyte-to-blastocyst 期间出现 distinct embryonic tRNA repertoires；4-cell stage 伴随 tRNA pseudogene 上调，tRNA activation 与 ZGA、H3K4me3 和 chromatin remodeling 同步。",
            "局限在于相关性仍多于因果；下一步需要 tRNA gene/tRF/修饰酶的胚胎内扰动验证。",
            "这是本报告最直接的胚胎主线论文，可作为 low-input tRNA-seq 方案和分期采样设计的模板。",
            "Source: PMID 42062277 | Deciphering tRNA repertoires and translation coordination during mouse early embryogenesis by ORACLE-tRNAseq.",
            assets.get("oracle_mouse"),
        ),
        Paper(
            "The dynamics and functional impact of tRNA repertoires during early embryogenesis in zebrafish.",
            "39402326",
            "The EMBO Journal",
            "2024",
            "zebrafish 胚胎显示 tRNA pool 重编程会改变解码速率与 maternal mRNA 稳定性",
            "早期胚胎 tRNA pool 是否与 zygotic transcriptome codon demand 匹配，并通过 decoding rate 影响 mRNA turnover。",
            "定量测定 zebrafish MZT 前后 tRNA repertoires；将 tRNA supply、codon composition、translation output 和 mRNA stability 联合分析。",
            "maternal 与 zygotic tRNA pools 明显不同；gastrulation 伴随全局翻译增强，使低供给 tRNA 对应 codon 的解码变慢，并促进某些 maternal mRNA destabilization。",
            "强调 tRNA supply 对转录后调控的功能影响，但需要更多单个 tRNA 或 tRF 的 rescue/knockdown 因果实验。",
            "适合指导我们把胚胎时期分成 MZT 前、ZGA、gastrulation 后，分别测 tRNA 与翻译读出。",
            "Source: PMID 39402326 | The dynamics and functional impact of tRNA repertoires during early embryogenesis in zebrafish.",
            None,
        ),
        Paper(
            "tRNA expression and modification landscapes, and their dynamics during zebrafish embryo development.",
            "38989621",
            "Nucleic Acids Research",
            "2024",
            "tRAM-seq 同时读出 zebrafish embryo 的 tRNA expression 与 modification landscape",
            "不同 tRNA gene copies/isodecoders 和 post-transcriptional modifications 是否在胚胎发育中动态变化。",
            "开发 tRAM-seq 实验与分析流程，覆盖核编码和线粒体 tRNAs，追踪 zebrafish embryo 发育时间序列。",
            "显示胚胎阶段 tRNA 表达和修饰并非静态背景，而是随发育动态改变；为 isodecoder-level 解释提供数据基础。",
            "NAR 论文偏 landscape，功能扰动较少；需要结合 morpholino/CRISPR 或 microinjection 验证候选 tRNA 的因果作用。",
            "可作为胚胎 tRNA modification 初筛平台：先找动态位点，再进入 LC-MS/MS 或酶扰动验证。",
            "Source: PMID 38989621 | tRNA expression and modification landscapes, and their dynamics during zebrafish embryo development.",
            None,
        ),
        Paper(
            "5' Half of specific tRNAs feeds back to promote corresponding tRNA gene transcription in vertebrate embryos.",
            "34797706",
            "Science Advances",
            "2021",
            "特定 5′ tRNA halves 可反馈促进对应 tRNA gene transcription 并影响 vertebrate embryos",
            "5′tRF halves 不只是降解产物，是否能反向调控自身来源 tRNA gene 的转录。",
            "在 zebrafish embryos 中监测 5′tRFlGly/GCC 和 5′tRFlGlu/CTC；用 morpholino knockdown 与 refolded tRNA rescue 做功能验证。",
            "5′tRF 与对应 mature tRNA 动态相关；敲低 5′tRF 下调对应 tRNA 并导致胚胎致死，补回正确折叠 tRNA 可 rescue。",
            "机制涉及 tRNA:DNA / tRF:DNA hybrid，但不同物种和不同 tRF 是否普遍成立还需扩展。",
            "给早期胚胎研究一个清晰因果入口：microinjection tRF mimic/inhibitor + tRNA rescue。",
            "Source: PMID 34797706 | 5' Half of specific tRNAs feeds back to promote corresponding tRNA gene transcription in vertebrate embryos.",
            None,
        ),
        Paper(
            "Epididymal dynamics and preimplantation roles of a sperm-enriched 5' fragment of tRNA-valine.",
            "41042673",
            "Cell Reports",
            "2025",
            "sperm-enriched tRFValCAC 连接 epididymal delivery 与 preimplantation embryo transcriptome",
            "精子成熟获得的 tRFValCAC 是否进入受精后胚胎并调控植入前发育。",
            "追踪 epididymis EV-to-sperm delivery；解析 hnRNPAB binding；在 preimplantation embryos 中抑制 tRFValCAC 并测转录组/发育表型。",
            "hnRNPAB 参与 tRFValCAC 装载与精子富集；胚胎中抑制 tRFValCAC 改变 RNA processing/translation 相关基因并影响早期发育。",
            "需要区分精子携带效应与胚胎内新生成 tRF；剂量、时间窗和 off-target 是关键控制。",
            "对哺乳动物早期胚胎最可操作：zygote microinjection tRF inhibitor/mimic 后看 2-cell、4-cell、blastocyst。",
            "Source: PMID 41042673 | Epididymal dynamics and preimplantation roles of a sperm-enriched 5' fragment of tRNA-valine.",
            None,
        ),
        Paper(
            "Biogenesis and function of tRNA fragments during sperm maturation and fertilization in mammals.",
            "26721685",
            "Science",
            "2016",
            "epididymosome 递送 tRNA fragments 为父源环境影响早期胚胎提供机制",
            "父代饮食能否改变成熟精子小 RNA，并通过受精后胚胎基因调控影响后代表型。",
            "比较 testicular 与 mature sperm small RNAs；检测 epididymosome RNA cargo；在 ESC/embryos 中测试 Gly-tRF 对 MERVL genes 的抑制。",
            "5′ Gly-tRF 在 epididymal maturation 中增加；epididymosomes 可递送 RNA；Gly-tRF 可抑制 ESC 和 embryos 中 MERVL 相关基因。",
            "经典论文提出机制框架，但精确靶点和胚胎阶段特异性仍需更高分辨率方法重做。",
            "适合纳入父源小 RNA 与 ZGA/2-cell-like program 的交叉研究。",
            "Source: PMID 26721685 | Biogenesis and function of tRNA fragments during sperm maturation and fertilization in mammals.",
            None,
        ),
        Paper(
            "Sperm tsRNAs contribute to intergenerational inheritance of an acquired metabolic disorder.",
            "26721680",
            "Science",
            "2016",
            "精子 tsRNA 注入 zygote 可重塑早期胚胎与子代代谢表型",
            "高脂饮食诱导的父源代谢表型是否由 sperm tsRNA 介导，而非 DNA methylation alone。",
            "建立 paternal HFD mouse model；测 sperm tsRNA profile 和 RNA modifications；将 sperm tsRNA fractions 注入正常 zygotes。",
            "HFD 改变 30-34 nt 5′ tRNA halves；注入 tsRNA fraction 可诱导 F1 代谢异常，并改变 early embryos 与 islets 的基因表达。",
            "fraction 注入不能完全定位单个 tsRNA；后续研究应使用合成修饰 tsRNA 和单分子功能验证。",
            "为胚胎实验提供强因果模式：small RNA fraction/mimic zygote injection + long-term phenotype。",
            "Source: PMID 26721680 | Sperm tsRNAs contribute to intergenerational inheritance of an acquired metabolic disorder.",
            None,
        ),
        Paper(
            "Pseudouridylation of tRNA-Derived Fragments Steers Translational Control in Stem Cells.",
            "29628141",
            "Cell",
            "2018",
            "PUS7 介导 tRF pseudouridylation 控制 stem cell translation 与胚层分化",
            "tRF 的 Ψ 修饰是否赋予其调控 translation initiation 的功能，从而影响 early embryogenesis/stem cell commitment。",
            "鉴定 PUS7-dependent tRF network；测试 tRF 对 translation initiation complex 的作用；在 ESC differentiation 与 HSC commitment 中验证。",
            "PUS7 修改并激活特定 tRF；PUS7 loss 破坏 tRF-mediated translational control，导致 protein biosynthesis 上升和 germ layer specification 缺陷。",
            "多在 stem cell 系统，移植到胚胎需验证时空表达、修饰状态和 microinjection rescue。",
            "提示胚胎 tRF 研究必须同时测 sequence 和 modification，而不仅是 small RNA abundance。",
            "Source: PMID 29628141 | Pseudouridylation of tRNA-Derived Fragments Steers Translational Control in Stem Cells.",
            None,
        ),
        Paper(
            "Mettl1/Wdr4-Mediated m7G tRNA Methylome Is Required for Normal mRNA Translation and Embryonic Stem Cell Self-Renewal and Differentiation.",
            "29983320",
            "Molecular Cell",
            "2018",
            "METTL1/WDR4-m7G tRNA methylome 是 ESC 翻译稳态和分化能力的必要条件",
            "tRNA m7G modification 是否通过 codon-specific ribosome pausing 影响 ESC self-renewal 和 differentiation。",
            "建立 m7G tRNA MeRIP-seq 与 TRAC-seq；构建 Mettl1/Wdr4 knockout mESC；分析 ribosome occupancy 与分化表型。",
            "22 个 tRNAs 在 RAGGU motif 发生 m7G；Mettl1 KO 改变对应 codon 的 ribosome occupancy，影响 cell cycle/brain-related genes 的翻译。",
            "ESC 不是完整胚胎；但为早期胚胎细胞命运转换提供了 tRNA modification 维度的机制候选。",
            "可优先检测 Mettl1/Wdr4 在 oocyte、2-cell、blastocyst 的表达与 m7G-tRNA 动态。",
            "Source: PMID 29983320 | Mettl1/Wdr4-Mediated m7G tRNA Methylome Is Required for Normal mRNA Translation and Embryonic Stem Cell Self-Renewal and Differentiation.",
            None,
        ),
        Paper(
            "Quantitative analysis of tRNA abundance and modifications by nanopore RNA sequencing.",
            "37024678",
            "Nature Biotechnology",
            "2024",
            "Nano-tRNAseq 同时读出 native tRNA abundance 与 modification dynamics",
            "能否绕过 RT 停顿与短 tRNA reads 丢失，实现 abundance 和 modification 的单实验定量。",
            "优化 nanopore native tRNA library；重新处理 raw current signals；在 yeast 中做 abundance、modification crosstalk 和 oxidative stress 应用。",
            "raw signal re-processing 大幅提高 tRNA reads 回收；可同时估计 tRNA 丰度和修饰变化，并看到 oxidative stress 下 tRNA population 改变。",
            "输入量、测序深度和胚胎样本适配仍需优化；对低输入 embryos 需先做 spike-in 与 pooling。",
            "适合作为胚胎 tRNA abundance + modification 的高通量候选平台。",
            "Source: PMID 37024678 | Quantitative analysis of tRNA abundance and modifications by nanopore RNA sequencing.",
            assets.get("nano_trnaseq"),
        ),
        Paper(
            "High-resolution quantitative profiling of tRNA abundance and modification status in eukaryotes by mim-tRNAseq.",
            "33581077",
            "Molecular Cell",
            "2021",
            "mim-tRNAseq 用 misincorporation signature 定量 eukaryotic tRNA abundance/modification",
            "能否利用修饰诱导错配而非将其视作测序噪声，提升 tRNA full-length sequencing 与修饰推断。",
            "TGIRT template switching library construction；配套 alignment/computational pipeline；在人、果蝇、酵母中验证。",
            "提升 tRNA coverage 和 abundance estimates；发现 human cell lines 之间 isodecoder pool 差异和同一 tRNA 内修饰互作。",
            "需要已知基因组和良好注释；在胚胎低输入条件下要验证反转录偏倚和 batch effects。",
            "适合胚胎阶段样本的 first-pass tRNA abundance/modification profiling。",
            "Source: PMID 33581077 | High-resolution quantitative profiling of tRNA abundance and modification status in eukaryotes by mim-tRNAseq.",
            assets.get("mim_trnaseq"),
        ),
        Paper(
            "Nanopore sequencing of intact aminoacylated tRNAs.",
            "40835813",
            "Nature Communications",
            "2025",
            "aa-tRNA-seq 在单分子层面读取 tRNA charging 与 amino acid identity",
            "能否在不破坏 aminoacyl linkage 的情况下同时解析 tRNA identity、modification 和 charging state。",
            "化学连接 adapter 将 amino acid 夹在 tRNA 与 adapter 之间；nanopore sequencing；机器学习从 signal distortion 识别 amino acid。",
            "可区分 amino acid identity 并估计 charging；在 tRNA modification enzyme loss 背景下发现 hypomodification-associated tRNA instability。",
            "方法新且复杂；胚胎应用需解决低输入、酸性保护、样本损耗和模型训练。",
            "一旦优化，可直接回答 MZT 期间“tRNA abundance 变了还是 charging efficiency 变了”。",
            "Source: PMID 40835813 | Nanopore sequencing of intact aminoacylated tRNAs.",
            assets.get("aa_trnaseq"),
        ),
        Paper(
            "In vivo structure profiling reveals human cytosolic and mitochondrial tRNA structurome and interactome in response to stress.",
            "40447571",
            "Nature Communications",
            "2025",
            "DM-DMS-MaPseq 把 tRNA folding/interactome 带到 in vivo transcriptome-wide 层面",
            "细胞内 tRNA 结构是否与 in vitro folding 不同，并在 oxidative stress 下产生功能性互作改变。",
            "DMS probing + demethylase treatment + MaP sequencing；分析 cytosolic 与 mitochondrial tRNA DMS profiles；比较 arsenite stress。",
            "tRNA in vivo DMS profile 与 in vitro 明显不同，反映 protein/ribosome interactions；arsenite 诱导 cytosolic 和 mitochondrial tRNA 结构/互作变化。",
            "目前在细胞系中完成；低输入胚胎需缩小反应体系并控制 DMS toxicity/time window。",
            "可用于检测胚胎应激或 ZGA 期间 tRNA 是否发生结构和互作重排。",
            "Source: PMID 40447571 | In vivo structure profiling reveals human cytosolic and mitochondrial tRNA structurome and interactome in response to stress.",
            assets.get("dm_dms"),
        ),
        Paper(
            "A ligation-independent sequencing method reveals tRNA-derived RNAs with blocked 3' termini.",
            "39096899",
            "Molecular Cell",
            "2024",
            "LIDAR 捕获传统 ligation-dependent small RNA-seq 漏掉的 blocked-end tDRs",
            "RNA 3′ 端修饰是否导致大量 tRNA-derived RNAs 在常规 small RNA-seq 中被系统性漏检。",
            "quasi-random priming + template switching；不依赖 3′ ligation；应用于 mESC、neural progenitor、mouse tissues 和 sperm。",
            "LIDAR 捕获更丰富 tDRs，尤其是 blocked 3′ termini 的 tDRs；显示传统方法低估了 tDR diversity。",
            "读出覆盖广，但定量解释需注意 priming bias；功能验证仍需单个 tDR perturbation。",
            "对胚胎 tRF 研究很关键：先避免测序方法把最有修饰的 tRF 过滤掉。",
            "Source: PMID 39096899 | A ligation-independent sequencing method reveals tRNA-derived RNAs with blocked 3' termini.",
            None,
        ),
        Paper(
            "tRNA m1A modification regulate HSC maintenance and self-renewal via mTORC1 signaling.",
            "38977676",
            "Nature Communications",
            "2024",
            "TRMT6/TRMT61A-m1A58 通过 TSC1 translation 控制 stem cell quiescence 与 mTORC1",
            "tRNA m1A58 writer 是否通过 codon decoding 改变关键 mRNA 翻译，从而维持 stem cell homeostasis。",
            "Trmt6 conditional KO mouse；FACS HSC phenotyping；scRNA-seq；LC-MS/MS/dot blot/m1A-tRNA-seq；rapamycin rescue。",
            "Trmt6 loss 使 HSC 异常增殖、自我更新下降；mTORC1 pathway 激活；TSC1 translation 受影响；mTOR inhibition 可部分 rescue。",
            "是 HSC 场景，非胚胎；但 stem cell quiescence/activation 的翻译逻辑可迁移到 blastomere fate。",
            "提示胚胎可优先测试 TRMT6/TRMT61A-m1A58 → TSC/mTOR 或 cell-cycle translation 轴。",
            "Source: PMID 38977676 | tRNA m1A modification regulate HSC maintenance and self-renewal via mTORC1 signaling.",
            assets.get("trmt6_hsc"),
        ),
        Paper(
            "A methyltransferase-independent role for METTL1 in tRNA aminoacylation and oncogenic transformation.",
            "39892392",
            "Molecular Cell",
            "2025",
            "METTL1 还能非催化性促进 tRNA aminoacylation 与 protein synthesis",
            "METTL1 的致癌作用是否完全依赖 m7G methyltransferase activity，还是存在 aminoacylation 相关非酶功能。",
            "zebrafish sarcoma model；METTL1 mutants；polysome profiling；IP-MS/Western；tRNA aminoacylation assays。",
            "催化死/磷酸模拟 METTL1 仍促进 oncogenesis；METTL1 结合 multi-tRNA synthetase complex，促进 aminoacylation、polysome formation 和 protein synthesis。",
            "疾病模型与胚胎相距较远，但揭示“tRNA writer 蛋白不只写修饰”的重要陷阱。",
            "胚胎中若扰动 METTL1，需同时区分 methylation-dependent 与 aminoacylation/scaffold-dependent effects。",
            "Source: PMID 39892392 | A methyltransferase-independent role for METTL1 in tRNA aminoacylation and oncogenic transformation.",
            assets.get("mettl1_model"),
        ),
        Paper(
            "Translational regulation by oxidative desulfuration of tRNA modifications.",
            "41807381",
            "Nature Communications",
            "2026",
            "氧化去硫化把 wobble tRNA modification 转化为 stress-sensitive translation regulator",
            "xm5s2U 的 thiocarbonyl group 是否在氧化环境下转化为 xm5h2U，并改变 codon recognition / aminoacylation。",
            "LC-MS/MS 鉴定 human cells 与 mouse tissues 中 xm5h2U；spike-in 验证形成；in vitro translation；aminoacylation assays；cryo-EM。",
            "mcm5h2U 削弱 Lys/Glu/Gln tRNA aminoacylation 与 codon recognition；结构解释 AAA/AAG decoding 改变。",
            "主要是 stress biology；在胚胎中需谨慎控制 oxidative stress 是否为生理范围。",
            "对胚胎培养条件很有启发：氧化状态可能通过 tRNA wobble modification 直接改变翻译。",
            "Source: PMID 41807381 | Translational regulation by oxidative desulfuration of tRNA modifications.",
            assets.get("desulfuration"),
        ),
        Paper(
            "RNA-triggered Cas12a3 cleaves tRNA tails to execute bacterial immunity.",
            "41501459",
            "Nature",
            "2026",
            "Cas12a3 通过 target RNA 触发 tRNA CCA tail cleavage，展示 tRNA depletion 可快速关停翻译",
            "CRISPR-Cas effector 是否可在识别 RNA 后切割非靶 RNA，尤其是保守 tRNA CCA tail。",
            "细胞/生化 cleavage assays；direct RNA sequencing；anti-phage defense；cryo-EM；合成 reporter 验证诊断应用。",
            "target RNA 触发 Cas12a3 切割 diverse tRNA 5′-CCA-3′ tail，引起 growth arrest 和 anti-phage defense；结构揭示 tRNA-loading domain。",
            "细菌免疫系统与胚胎无直接关系；但提供了 tRNA tail 完整性作为翻译开关的概念。",
            "可作为工具启发：在人为系统中精准破坏特定 tRNA tail，观察胚胎翻译和发育窗口敏感性。",
            "Source: PMID 41501459 | RNA-triggered Cas12a3 cleaves tRNA tails to execute bacterial immunity.",
            assets.get("cas12a3"),
        ),
        Paper(
            "AAV-delivered suppressor tRNA overcomes a nonsense mutation in mice.",
            "35322228",
            "Nature",
            "2022",
            "AAV-sup-tRNA 证明 tRNA gene 可作为体内工程干预载体",
            "递送 suppressor tRNA 能否在体内 readthrough premature stop codon，并避免扰乱 global tRNA homeostasis。",
            "rAAV delivery of suppressor tRNA；mouse nonsense mutation model；ribosome profiling；tRNA sequencing；多组织给药优化。",
            "单次给药可长期 rescue disease phenotype；通过 readthrough + NMD inhibition 协同恢复；对正常 stop codon global readthrough 较有限。",
            "治疗场景非胚胎；但证明 tRNA 作为小型遗传元件可被安全递送和功能化。",
            "胚胎研究可借鉴其 tRNA expression cassette 设计，用于 transient 或 tissue-stage specific tRNA perturbation。",
            "Source: PMID 35322228 | AAV-delivered suppressor tRNA overcomes a nonsense mutation in mice.",
            None,
        ),
        Paper(
            "An engineered UGA suppressor tRNA gene for disease-agnostic AAV delivery.",
            "41555020",
            "Nature Biotechnology",
            "2026",
            "工程化 UGA sup-tRNA 拓展 nonsense suppression 的 codon scope",
            "UGA-targeting sup-tRNA 能否解决 AAV packaging/production 难题，并在不同疾病模型中恢复功能。",
            "优化 tRNA gene regulatory elements；rAAV packaging QC；mouse lysosomal storage disease models；组织表达、aminoacylation 与疗效关联分析。",
            "工程化 UGA-sup-tRNA 可高效包装并在两种模型中恢复约 10% 正常酶活；组织间表达和 aminoacylation 差异影响疗效。",
            "胚胎使用不应照搬 AAV；但表达 cassette、charging 和组织差异评价非常可借鉴。",
            "为胚胎 tRNA gain-of-function 或 codon readthrough reporter 构建提供工程化设计原则。",
            "Source: PMID 41555020 | An engineered UGA suppressor tRNA gene for disease-agnostic AAV delivery.",
            assets.get("uga_suppressor"),
        ),
        Paper(
            "RelA-SpoT Homolog toxins pyrophosphorylate the CCA end of tRNA to inhibit protein synthesis.",
            "34174184",
            "Molecular Cell",
            "2021",
            "RSH toxins 通过 tRNA 3′ CCA pyrophosphorylation 抑制 aminoacylation",
            "多类 toxSAS 是否直接把 tRNA 作为底物，从而阻断 protein synthesis。",
            "体外酶学、质谱、translation assays 和 toxin-antitoxin 功能分析；测试 SAH hydrolase reversal。",
            "FaRel2/PhRel/PhRel2/CapRel 将 pyrophosphate 转移到 tRNA 3′ CCA，抑制 aminoacylation 和 RelA sensing；SAH 可逆转。",
            "细菌毒素模型非胚胎；但 tRNA CCA 端修饰作为 charging gate 很有机制价值。",
            "启发胚胎实验关注 CCA tail integrity、charging readout 与快速翻译抑制之间的因果关系。",
            "Source: PMID 34174184 | RelA-SpoT Homolog toxins pyrophosphorylate the CCA end of tRNA to inhibit protein synthesis.",
            None,
        ),
        Paper(
            "Design, construction, and functional characterization of a tRNA neochromosome in yeast.",
            "37944512",
            "Cell",
            "2023",
            "tRNA neochromosome 证明 tRNA gene dosage/location 可被系统工程化",
            "所有核 tRNA genes 被重定位到 de novo chromosome 后，细胞如何维持 tRNA transcription、chromatin 和 genome function。",
            "设计 190-kb tRNA neochromosome；yeast construction；tRNA-seq、transcriptomics、proteomics、nucleosome mapping、replication profiling、FISH、Hi-C。",
            "neochromosome 可承载 275 个 relocated tRNA genes；出现 ploidy adaptation；系统性读出 tRNA gene organization 的影响。",
            "合成酵母模型不直接对应胚胎，但揭示 tRNA gene copy/location 是可工程化变量。",
            "对胚胎研究的启发是区分 tRNA gene transcription、pseudogene activation 与 mature tRNA supply。",
            "Source: PMID 37944512 | Design, construction, and functional characterization of a tRNA neochromosome in yeast.",
            None,
        ),
    ]


def add_method_matrix(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "方法矩阵：早期胚胎 tRNA 湿实验应按“读出层级”组合", "不是单一测序方法能回答全部问题")
    rows = [
        ("丰度 / isodecoder", "ORACLE-tRNAseq, mim-tRNAseq, Nano-tRNAseq", "低输入、spike-in、isodecoder 注释"),
        ("修饰", "tRAM-seq, mim/Nano signatures, LC-MS/MS", "修饰阻断、同分异构体、样本量"),
        ("charging", "acid northern, aa-tRNA-seq", "酸性保护、快速裂解、低输入"),
        ("结构 / 互作", "DM-DMS-MaPseq", "DMS 时间窗、胚胎毒性、线粒体 tRNA"),
        ("tRF / tsRNA", "LIDAR, small RNA-seq, mimic/inhibitor", "3′ blocked ends、修饰、microinjection"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.65), Inches(1.35), Inches(12.0), Inches(4.55)).table
    widths = [Inches(2.2), Inches(4.4), Inches(5.4)]
    for i, width in enumerate(widths):
        table.columns[i].width = width
    for j, htxt in enumerate(["读出层级", "推荐技术", "胚胎样本关键控制"]):
        c = table.cell(0, j)
        c.text = htxt
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(229, 239, 235)
        c.text_frame.paragraphs[0].font.bold = True
        c.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
        c.text_frame.paragraphs[0].font.size = Pt(11)
    for i, row in enumerate(rows, 1):
        for j, txt in enumerate(row):
            c = table.cell(i, j)
            c.text = txt
            for p in c.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(10.5)
                p.font.color.rgb = INK
    add_takeaway(slide, "建议先做低输入 landscape，再用 charging / LC-MS/MS / microinjection 进入因果验证。")
    add_footer(slide, "来源：综合 PMID 42062277, 37024678, 33581077, 40835813, 40447571, 39096899 等")


def add_experiment_plan(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "可执行课题路线：围绕 MZT/ZGA 的 tRNA 多层读出与因果验证", "适用于 mouse 或 zebrafish early embryo；根据样本量调整技术组合")
    steps = [
        ("1. 分期采样", "oocyte / zygote / 2-cell / 4-cell / morula / blastocyst 或 zebrafish MZT time course"),
        ("2. 初筛 landscape", "ORACLE-tRNAseq 或 mim/tRAM-seq：tRNA abundance、isodecoder、修饰 signature"),
        ("3. 深入验证", "LC-MS/MS 验证候选修饰；aa-tRNA-seq/acid northern 验证 charging；LIDAR 捕获 tRF"),
        ("4. 因果扰动", "microinjection tRF mimic/inhibitor、writer knockdown/CRISPRi、candidate tRNA rescue"),
        ("5. 表型读出", "ZGA reporter、Ribo-seq/puromycin incorporation、single embryo RNA-seq、blastocyst rate"),
    ]
    y = Inches(1.35)
    for i, (head, body) in enumerate(steps):
        x = Inches(0.65 + i * 2.45)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.15), Inches(3.85))
        set_fill(box, [RGBColor(233, 244, 240), RGBColor(237, 242, 249), RGBColor(247, 242, 233), RGBColor(241, 238, 249), RGBColor(249, 238, 238)][i])
        set_textbox(box, f"{head}\n\n{body}", size=11, bold=False, color=INK, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(2.07), y + Inches(1.78), Inches(0.34), Inches(0.32))
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(130, 145, 160)
            arr.line.fill.background()
    add_takeaway(slide, "判断机制是否成立的最低闭环：动态相关性 + 候选 tRNA/tRF 扰动 + translation/ZGA/readout rescue。")
    add_footer(slide, "来源：本报告综合实验设计；依据入选文献的方法和早期胚胎应用场景整理")


def add_conclusion(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "总结：tRNA 是早期胚胎翻译控制的可测、可扰动、可工程化层级", "汇报结论")
    add_bullets(
        slide,
        Inches(0.8),
        Inches(1.5),
        Inches(11.7),
        Inches(4.45),
        [
            "早期胚胎证据已经从 small RNA/tsRNA 进入 mature tRNA repertoires、modification 和 translation coordination。",
            "最直接的主线是 ORACLE-tRNAseq / zebrafish time-course 文献；最关键的因果入口是 tRF microinjection、writer perturbation 和 tRNA rescue。",
            "本地 tRNA 湿实验文献提供完整工具箱：abundance、modification、charging、structure/interactome、blocked-end tRF capture。",
            "下一步研究应避免只做关联组学：每一个动态 tRNA/tRF 候选都要配套发育表型、翻译读出和 rescue。",
        ],
        size=17,
    )
    add_axis_map(slide, Inches(3.5), Inches(4.55), Inches(6.2), Inches(1.75))
    add_footer(slide, "来源：本报告综合整理")


def write_notes(papers: list[Paper]) -> None:
    md = [
        "# tRNA 相关湿实验与早期胚胎发育：逐篇精读要点",
        "",
        "本文件与 PPT 同步生成，用于追溯每篇文献的逻辑、方法、结果、讨论和胚胎研究启发。",
        "",
    ]
    for i, p in enumerate(papers, 1):
        md += [
            f"## {i}. {p.title}",
            f"- PMID: {p.pmid}",
            f"- Journal/Year: {p.journal}, {p.year}",
            f"- 汇报定位: {p.role}",
            f"- 逻辑/假说: {p.hypothesis}",
            f"- 方法步骤: {p.methods}",
            f"- 关键结果: {p.results}",
            f"- 讨论与局限: {p.discussion}",
            f"- 对早期胚胎研究的用途: {p.embryo_angle}",
            "",
        ]
    (OUT / "literature_reading_notes_cn.md").write_text("\n".join(md), encoding="utf-8")


def build_deck() -> None:
    ensure_dirs()
    assets = prepare_assets()
    papers = build_papers(assets)

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    add_cover(prs)
    add_matrix_slide(prs)
    add_question_slide(prs)

    add_section_slide(
        prs,
        "第一部分：早期胚胎与生殖发育中的 tRNA/tRF 证据",
        "优先阅读直接涉及 oocyte、MZT/ZGA、zygote、preimplantation embryo 或 ESC commitment 的论文",
        [
            "重点看是否有分期样本、低输入方法、胚胎内 perturbation 或 rescue。",
            "tRF 文献提供最强因果证据；mature tRNA landscape 文献提供最强系统框架。",
            "ESC 文献作为 early lineage commitment 的机制近似模型。",
        ],
        color=GREEN,
    )
    for p in papers[:9]:
        add_paper_slide(prs, p)

    add_section_slide(
        prs,
        "第二部分：可迁移到胚胎样本的 tRNA 湿实验工具箱",
        "这些本地文献主要解决 tRNA 难测、难定量、难区分修饰和 charging 的技术瓶颈",
        [
            "选择方法时先问：样本量、是否需要 single-isodecoder、是否需要修饰或 charging。",
            "结果解释必须配 spike-in、低输入重复和 orthogonal validation。",
            "低输入胚胎建议从 landscape 进入候选验证，不一开始追求全层级。",
        ],
        color=TEAL,
    )
    for p in papers[9:14]:
        add_paper_slide(prs, p)
    add_method_matrix(prs)

    add_section_slide(
        prs,
        "第三部分：tRNA 作为功能扰动对象和工程化调控对象",
        "这些文献不一定发生在胚胎中，但提供了可转化的机制轴和干预范式",
        [
            "修饰酶、CCA tail、aminoacylation 和 suppressor tRNA 都可作为胚胎功能验证入口。",
            "必须区分 tRNA writer 的催化功能、scaffold 功能和 charging 间接效应。",
            "工程化 tRNA 可用于 reporter、rescue 或翻译选择性干预。",
        ],
        color=GOLD,
    )
    for p in papers[14:]:
        add_paper_slide(prs, p)

    add_experiment_plan(prs)
    add_conclusion(prs)

    pptx = OUT / "final_presentation_cn.pptx"
    prs.save(pptx)

    reopened = Presentation(pptx)
    media_count = len(list((OUT / "final_presentation_cn.pptx").parent.glob("assets/figures/*")))
    write_notes(papers)
    manifest = {
        "pptx": str(pptx),
        "slides": len(reopened.slides),
        "paper_slides": len(papers),
        "assets_prepared": assets,
        "screening": {
            "main_added_pmids": ["42062277", "39402326", "38989621", "34797706", "41042673", "26721685", "26721680", "29628141", "29983320"],
            "background_only_pmids": ["35032425", "37173525"],
        },
    }
    (OUT / "asset_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    qa = [
        "# QA report",
        "",
        f"- PPTX creation status: success",
        f"- Output file: `{pptx}`",
        f"- Slide count: {len(reopened.slides)}",
        f"- Literature deep-read slides: {len(papers)}",
        f"- Figure assets prepared: {len(assets)}",
        "- Verification: reopened with python-pptx after saving; assets copied/downloaded/rendered into `output/assets/figures/`.",
        "- Known limitation: no full slide rendering was available in the current workflow; several paywalled/non-PMC papers use rebuilt schematic figures rather than original article figures.",
        "- Source labels: each literature slide includes PMID and full title in the subtitle/footer.",
        "",
    ]
    (OUT / "qa_report.md").write_text("\n".join(qa), encoding="utf-8")
    print(json.dumps({"pptx": str(pptx), "slides": len(reopened.slides), "paper_slides": len(papers), "assets": len(assets)}, ensure_ascii=False))


if __name__ == "__main__":
    build_deck()
