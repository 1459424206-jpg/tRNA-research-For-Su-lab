from __future__ import annotations

import csv
import json
import re
import textwrap
import time
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

import fitz
import requests
from bs4 import BeautifulSoup
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
DEEP_DIR = OUT / "deep_literature_deck"
DEEP_DIR.mkdir(parents=True, exist_ok=True)
PPTX_PATH = OUT / "final_presentation_cn_deep_literature.pptx"
QA_PATH = OUT / "qa_report_deep_literature.md"
AVAIL_PATH = OUT / "pdf_download_availability.csv"
NOTES_PATH = OUT / "deep_literature_summary.json"

SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")

TOPICS = [
    ("biological functions and structural characteristics of tRNA", "tRNA 的生物功能与结构特征", "Biological functions and structural characteristics of tRNA"),
    ("tRNA analysis tools and algorithms", "tRNA 分析工具与算法", "tRNA analysis tools and algorithms"),
    ("tRNA mod-fingerprint database", "tRNA 修饰指纹数据库", "tRNA mod-fingerprint database"),
    ("tRNA-related wet-lab experiments", "tRNA 相关湿实验", "tRNA-related wet-lab experiments"),
]


def safe_name(text: str, limit: int = 95) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")
    return text[:limit] or "item"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def get_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "Chrome/125.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/pdf,*/*;q=0.8",
        }
    )
    return s


SESSION = get_session()


def download_pdf(row: dict[str, str], pdf_dir: Path) -> tuple[Path | None, str, str]:
    pdf_dir.mkdir(parents=True, exist_ok=True)
    pmid = row.get("pmid", "")
    doi = row.get("doi", "")
    title = row.get("title", "")
    target = pdf_dir / f"{pmid}_{safe_name(title)}.pdf"
    if target.exists() and target.stat().st_size > 1000:
        return target, "downloaded", "already_exists"

    candidates: list[tuple[str, str]] = []
    if doi.startswith("10.1038/"):
        slug = doi.split("/")[-1]
        candidates.append(("nature_pdf", f"https://www.nature.com/articles/{slug}.pdf"))
    if doi.startswith("10.1016/"):
        # Elsevier/Cell often blocks unauthenticated PDF access; try the official route
        # but never bypass access control.
        candidates.append(("doi_landing", f"https://doi.org/{doi}"))
    pmcid = row.get("pmcid", "")
    if pmcid:
        candidates.append(("pmc_pdf_endpoint", f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/pdf/"))

    errors = []
    for label, url in candidates:
        try:
            response = SESSION.get(url, timeout=70, allow_redirects=True)
            content_type = response.headers.get("content-type", "").lower()
            if response.status_code == 200 and response.content[:4] == b"%PDF":
                target.write_bytes(response.content)
                return target, "downloaded", f"{label}: {url}"
            if label == "doi_landing":
                # Capture Cell/Molecular Cell PII if visible for transparent diagnostics.
                match = re.search(r"/retrieve/pii/([A-Z0-9]+)", response.url)
                pii = match.group(1) if match else ""
                errors.append(f"{label}: not PDF; landing={response.url}; pii={pii}")
            else:
                errors.append(f"{label}: status={response.status_code}; type={content_type}; url={response.url}")
        except Exception as exc:
            errors.append(f"{label}: {type(exc).__name__}: {exc}")
        time.sleep(0.25)
    return None, "not_publicly_downloadable", " | ".join(errors)


def save_pmc_html(row: dict[str, str], html_dir: Path) -> Path | None:
    pmcid = row.get("pmcid", "")
    if not pmcid:
        return None
    html_dir.mkdir(parents=True, exist_ok=True)
    path = html_dir / f"{row.get('pmid','')}_{pmcid}.html"
    if path.exists() and path.stat().st_size > 50000:
        return path
    try:
        url = f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/"
        response = SESSION.get(url, timeout=50)
        if response.status_code == 200:
            path.write_text(response.text, encoding="utf-8")
            return path
    except Exception:
        return None
    return None


def clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", text)
    return text


def split_sentences(text: str, max_items: int = 5, min_len: int = 35) -> list[str]:
    text = clean_text(text)
    parts = re.split(r"(?<=[.!?])\s+", text)
    out = []
    for part in parts:
        part = part.strip()
        if len(part) >= min_len and not re.search(r"^(copyright|©|publisher|author information)", part, re.I):
            out.append(part)
        if len(out) >= max_items:
            break
    return out


def clip(text: str, length: int = 170) -> str:
    text = clean_text(text)
    return text if len(text) <= length else text[: length - 1].rstrip() + "…"


def extract_pdf_text(pdf_path: Path | None) -> str:
    if not pdf_path or not pdf_path.exists():
        return ""
    try:
        doc = fitz.open(pdf_path)
        pages = []
        for page in doc[: min(len(doc), 12)]:
            pages.append(page.get_text("text"))
        return clean_text("\n".join(pages))
    except Exception:
        return ""


def extract_html_sections(html_path: Path | None) -> dict[str, str]:
    if not html_path or not html_path.exists() or html_path.stat().st_size < 50000:
        return {}
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "aside"]):
        tag.decompose()
    sections: dict[str, list[str]] = defaultdict(list)
    current = "body"
    for node in soup.find_all(["h1", "h2", "h3", "p"]):
        text = clean_text(node.get_text(" "))
        if not text:
            continue
        if node.name in {"h1", "h2", "h3"}:
            low = text.lower()
            if "abstract" in low:
                current = "abstract"
            elif "method" in low or "material" in low:
                current = "methods"
            elif "result" in low:
                current = "results"
            elif "discussion" in low:
                current = "discussion"
            elif "conclusion" in low:
                current = "discussion"
            elif "introduction" in low or "background" in low:
                current = "introduction"
            else:
                current = "body"
        else:
            if len(text) > 35:
                sections[current].append(text)
    return {k: clean_text(" ".join(v[:10])) for k, v in sections.items()}


def extract_pdf_sections(text: str) -> dict[str, str]:
    if not text:
        return {}
    keys = [
        ("abstract", r"\bAbstract\b"),
        ("introduction", r"\bIntroduction\b"),
        ("results", r"\bResults\b"),
        ("discussion", r"\bDiscussion\b|\bConclusion\b"),
        ("methods", r"\bMethods\b|\bMaterials and methods\b|\bMethods summary\b"),
    ]
    positions = []
    for name, pat in keys:
        match = re.search(pat, text, re.I)
        if match:
            positions.append((match.start(), name))
    positions.sort()
    sections = {}
    for idx, (start, name) in enumerate(positions):
        end = positions[idx + 1][0] if idx + 1 < len(positions) else min(len(text), start + 5000)
        chunk = text[start:end]
        chunk = re.sub(r"^(Abstract|Introduction|Results|Discussion|Conclusion|Methods|Materials and methods|Methods summary)\b", "", chunk, flags=re.I)
        sections[name] = clean_text(chunk[:5000])
    return sections


def find_figures(row: dict[str, str], topic_dir: Path, max_figures: int = 3) -> list[Path]:
    figures_dir = topic_dir / "downloaded_resources" / "figures"
    pmcid = row.get("pmcid", "")
    found = []
    if pmcid and figures_dir.exists():
        found.extend(sorted(figures_dir.glob(f"{pmcid}_*.jpg")))
        found.extend(sorted(figures_dir.glob(f"{pmcid}_*.png")))
    return found[:max_figures]


def ensure_pmc_figures(row: dict[str, str], html_path: Path | None, topic_dir: Path, max_figures: int = 3) -> list[Path]:
    existing = find_figures(row, topic_dir, max_figures=max_figures)
    if len(existing) >= max_figures:
        return existing
    pmcid = row.get("pmcid", "")
    if not pmcid or not html_path or not html_path.exists():
        return existing
    try:
        html = html_path.read_text(encoding="utf-8", errors="ignore")
        urls = []
        for match in re.finditer(
            r"https://cdn\.ncbi\.nlm\.nih\.gov/pmc/blobs/[^\"'\s<>]+?(?:Fig|fig)[^\"'\s<>]+?\.(?:jpg|jpeg|png|gif)",
            html,
        ):
            url = match.group(0).replace("&amp;", "&")
            if url not in urls:
                urls.append(url)
        figures_dir = topic_dir / "downloaded_resources" / "figures"
        figures_dir.mkdir(parents=True, exist_ok=True)
        for url in urls:
            if len(existing) >= max_figures:
                break
            suffix = Path(url.split("?")[0]).suffix or ".jpg"
            out = figures_dir / f"{pmcid}_{safe_name(Path(url.split('?')[0]).name)}"
            if not out.suffix:
                out = out.with_suffix(suffix)
            if not out.exists():
                response = SESSION.get(url, timeout=45)
                if response.status_code == 200 and len(response.content) > 2000:
                    out.write_bytes(response.content)
            if out.exists() and out not in existing:
                existing.append(out)
        return existing[:max_figures]
    except Exception:
        return existing


def render_pdf_thumbnail(pdf_path: Path | None, out_path: Path) -> Path | None:
    if not pdf_path or not pdf_path.exists():
        return None
    if out_path.exists():
        return out_path
    try:
        doc = fitz.open(pdf_path)
        page = doc[0]
        pix = page.get_pixmap(matrix=fitz.Matrix(1.1, 1.1), alpha=False)
        pix.save(out_path)
        return out_path
    except Exception:
        return None


def classify_article(row: dict[str, str], text: str, sections: dict[str, str]) -> str:
    title = row.get("title", "").lower()
    abstract = row.get("abstract", "").lower()
    hay = f"{title} {abstract} {text[:2000].lower()}"
    if any(x in hay for x in ["review", "perspective", "renovatio", "overview"]):
        return "review"
    if "methods" in sections or "results" in sections or any(x in hay for x in ["we show", "we demonstrate", "assay", "sequencing", "mice", "cells"]):
        return "experimental"
    return "experimental"


def infer_focus(topic_cn: str, row: dict[str, str]) -> str:
    title = row.get("title", "")
    if "结构" in topic_cn:
        return f"与本子项目的关系：该文可用于解释 tRNA 结构、修饰或互作如何转化为翻译/免疫/疾病功能。"
    if "算法" in topic_cn:
        return f"与本子项目的关系：该文可用于提取 tRNA 定量、注释、测序信号处理或预测算法的设计原则。"
    if "数据库" in topic_cn:
        return f"与本子项目的关系：该文可转化为修饰位点、实验条件、检测技术和证据等级字段。"
    return f"与本子项目的关系：该文可用于设计样本处理、检测平台、关键对照和结果验证流程。"


def build_article_summary(topic_cn: str, row: dict[str, str], pdf_path: Path | None, html_path: Path | None, figures: list[Path]) -> dict[str, Any]:
    pdf_text = extract_pdf_text(pdf_path)
    html_sections = extract_html_sections(html_path)
    pdf_sections = extract_pdf_sections(pdf_text)
    sections = {**pdf_sections, **{k: v for k, v in html_sections.items() if v}}
    abstract = row.get("abstract", "") or sections.get("abstract", "") or pdf_text[:1200]
    article_type = classify_article(row, pdf_text, sections)
    intro_points = split_sentences(sections.get("introduction", "") or abstract, max_items=4)
    method_points = split_sentences(sections.get("methods", ""), max_items=4)
    result_points = split_sentences(sections.get("results", "") or abstract, max_items=5)
    discussion_points = split_sentences(sections.get("discussion", "") or abstract, max_items=4)

    if not intro_points:
        intro_points = split_sentences(abstract, max_items=4)
    if not result_points:
        result_points = split_sentences(abstract, max_items=5)
    if not method_points and article_type == "experimental":
        method_points = [
            "全文方法段未能自动稳定抽取；本页基于摘要、题名、图像和可下载全文信息概括实验/分析路线。",
        ]
    if not discussion_points:
        discussion_points = result_points[-3:] if result_points else intro_points[-3:]

    return {
        "topic_cn": topic_cn,
        "type": article_type,
        "meta": row,
        "pdf": str(pdf_path) if pdf_path else "",
        "html": str(html_path) if html_path else "",
        "figures": [str(f) for f in figures],
        "question": [clip(x, 145) for x in intro_points[:3]],
        "methods": [clip(x, 145) for x in method_points[:4]],
        "results": [clip(x, 145) for x in result_points[:5]],
        "discussion": [clip(x, 145) for x in discussion_points[:4]],
        "focus": infer_focus(topic_cn, row),
    }


def color(rgb):
    return RGBColor(*rgb)


def add_text(slide, x, y, w, h, text, size=14, bold=False, gray=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color((80, 80, 80) if gray else (0, 0, 0))
    return box


def add_bullets(slide, x, y, w, h, bullets, size=12, max_bullets=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.03)
    for i, b in enumerate(bullets[:max_bullets]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.space_after = Pt(4)
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size)
            run.font.color.rgb = color((0, 0, 0))
    return box


def add_title(slide, title, subtitle=""):
    add_text(slide, 0.45, 0.22, 12.3, 0.43, title, size=20, bold=True)
    if subtitle:
        add_text(slide, 0.47, 0.68, 12.1, 0.25, subtitle, size=8, gray=True)
    line = slide.shapes.add_shape(1, Inches(0.45), Inches(0.98), Inches(12.35), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = color((0, 0, 0))
    line.line.color.rgb = color((0, 0, 0))


def fit_image(slide, path: Path, x, y, w, h):
    try:
        with Image.open(path) as img:
            iw, ih = img.size
        ratio = iw / ih
        box_ratio = w / h
        if ratio > box_ratio:
            nw, nh = w, w / ratio
        else:
            nh, nw = h, h * ratio
        px, py = x + (w - nw) / 2, y + (h - nh) / 2
        slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(nw), height=Inches(nh))
        return True
    except Exception:
        return False


def source_line(s: dict[str, Any]) -> str:
    m = s["meta"]
    pmcid = m.get("pmcid") or "NA"
    return f"Source: {m.get('journal','')} {m.get('year','')}; PMID {m.get('pmid','')}; PMCID {pmcid}; DOI {m.get('doi','') or 'NA'}"


def add_article_slides(prs: Presentation, s: dict[str, Any], index_label: str, thumb_dir: Path):
    m = s["meta"]
    title = m.get("title", "")
    subtitle = f"{index_label} | {m.get('journal','')} ({m.get('year','')}) | {s['type']} | DOI {m.get('doi','') or 'NA'}"
    fig_paths = [Path(p) for p in s.get("figures", [])]
    pdf_path = Path(s["pdf"]) if s.get("pdf") else None
    thumb = render_pdf_thumbnail(pdf_path, thumb_dir / f"{m.get('pmid','')}_cover.png") if pdf_path else None

    # Slide 1: compressed paper card.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, f"{index_label} 文献概览：{clip(title, 75)}", subtitle)
    add_text(slide, 0.55, 1.25, 7.2, 0.7, "研究问题 / 背景", size=15, bold=True)
    add_bullets(slide, 0.55, 1.75, 7.15, 2.4, s["question"], size=12, max_bullets=4)
    add_text(slide, 0.55, 4.45, 7.15, 0.6, s["focus"], size=13, bold=True)
    visual = fig_paths[0] if fig_paths else thumb
    if visual:
        fit_image(slide, visual, 8.0, 1.25, 4.65, 4.7)
    add_text(slide, 0.55, 6.78, 12.1, 0.25, source_line(s), size=7, gray=True)

    # Slide 2: methods/results for experimental, major arguments for review.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if s["type"] == "review":
        add_title(slide, f"{index_label} 综述主要论点：{clip(title, 70)}", subtitle)
        add_text(slide, 0.55, 1.25, 6.2, 0.4, "主要论点", size=15, bold=True)
        add_bullets(slide, 0.55, 1.75, 6.2, 4.35, s["results"] or s["question"], size=12, max_bullets=5)
        visual = fig_paths[1] if len(fig_paths) > 1 else (fig_paths[0] if fig_paths else thumb)
        if visual:
            fit_image(slide, visual, 7.0, 1.25, 5.6, 4.9)
    else:
        add_title(slide, f"{index_label} 材料方法与结果：{clip(title, 70)}", subtitle)
        add_text(slide, 0.55, 1.25, 4.0, 0.4, "材料/方法/分析路线", size=14, bold=True)
        add_bullets(slide, 0.55, 1.68, 4.1, 4.45, s["methods"], size=11, max_bullets=4)
        add_text(slide, 4.95, 1.25, 3.25, 0.4, "核心结果", size=14, bold=True)
        add_bullets(slide, 4.95, 1.68, 3.35, 4.45, s["results"], size=10, max_bullets=5)
        visual = fig_paths[1] if len(fig_paths) > 1 else (fig_paths[0] if fig_paths else thumb)
        if visual:
            fit_image(slide, visual, 8.55, 1.25, 4.05, 4.9)
    add_text(slide, 0.55, 6.78, 12.1, 0.25, source_line(s), size=7, gray=True)

    # Slide 3: discussion/synthesis.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    label = "讨论与局限" if s["type"] != "review" else "综合观点与空白"
    add_title(slide, f"{index_label} {label}：{clip(title, 72)}", subtitle)
    add_text(slide, 0.55, 1.25, 6.15, 0.4, label, size=15, bold=True)
    add_bullets(slide, 0.55, 1.75, 6.2, 4.1, s["discussion"], size=12, max_bullets=4)
    add_text(slide, 0.55, 5.95, 6.2, 0.55, "可用于本项目的提炼", size=14, bold=True)
    add_bullets(slide, 0.55, 6.25, 6.2, 0.5, [s["focus"]], size=10, max_bullets=1)
    visual = fig_paths[2] if len(fig_paths) > 2 else (fig_paths[-1] if fig_paths else thumb)
    if visual:
        fit_image(slide, visual, 7.05, 1.25, 5.55, 5.15)
    add_text(slide, 0.55, 6.78, 12.1, 0.25, source_line(s), size=7, gray=True)


def build_deck(summaries_by_topic: list[tuple[str, str, str, list[dict[str, Any]]]]):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    thumb_dir = DEEP_DIR / "pdf_thumbnails"
    thumb_dir.mkdir(exist_ok=True)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, 0.65, 1.0, 12.0, 0.7, "tRNA research 文献逐篇精读版", size=28, bold=True)
    add_text(slide, 0.68, 1.82, 11.6, 0.42, "每篇文献 3 页：概览、方法/结果或综述论点、讨论与项目转化", size=15)
    add_text(slide, 0.68, 2.35, 11.6, 0.3, "白底黑字；图像来自公开 PMC/Nature PDF/PMC figure；中文主讲", size=10, gray=True)
    add_text(slide, 0.68, 6.35, 11.6, 0.3, "说明：非公开 PDF 不绕过版权访问控制；对应页面使用 DOI/PMC HTML/摘要进行内容抽取。", size=10, gray=True)

    total = sum(len(x[3]) for x in summaries_by_topic)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "新版制作逻辑", "Deck logic")
    add_bullets(
        slide,
        0.75,
        1.4,
        11.8,
        4.8,
        [
            f"覆盖 4 个子项目、{total} 篇入选文献；每篇固定 3 页，高度浓缩但保留方法、结果、讨论或综述论点。",
            "实验文章：第 1 页讲研究问题，第 2 页讲材料/方法和核心结果，第 3 页讲讨论、局限与本项目可转化点。",
            "综述文章：第 1 页讲背景与范围，第 2 页列主要论点，第 3 页列综合观点、争议和未来空白。",
            "图像优先级：PMC 原始图 > PDF 首页缩略图 > 题录信息；所有图像页均保留来源标签。",
        ],
        size=17,
    )

    for topic_idx, (folder, topic_cn, topic_en, summaries) in enumerate(summaries_by_topic, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(slide, 0.65, 1.25, 12.0, 0.7, f"Part {topic_idx}. {topic_cn}", size=27, bold=True)
        add_text(slide, 0.68, 2.0, 12.0, 0.35, topic_en, size=14)
        add_text(slide, 0.68, 2.52, 12.0, 0.35, f"本部分包含 {len(summaries)} 篇文献，每篇 3 页。", size=13, gray=True)
        rows = [f"{i+1}. {clip(s['meta'].get('title',''), 95)}" for i, s in enumerate(summaries[:15])]
        add_bullets(slide, 0.75, 3.15, 11.6, 3.3, rows, size=9, max_bullets=15)
        for i, summary in enumerate(summaries, 1):
            add_article_slides(prs, summary, f"{topic_idx}.{i}", thumb_dir)

    prs.save(PPTX_PATH)


def build_contact_sheet(pptx_path: Path) -> tuple[Path | None, Path | None, int]:
    preview = OUT / "preview_deep_literature"
    preview.mkdir(exist_ok=True)
    pdf = preview / (pptx_path.stem + ".pdf")
    if SOFFICE.exists():
        import subprocess

        subprocess.run(
            [str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(preview), str(pptx_path)],
            cwd=str(ROOT),
            check=False,
            timeout=240,
        )
    if not pdf.exists():
        return None, None, 0
    png_dir = preview / "slides_png"
    png_dir.mkdir(exist_ok=True)
    doc = fitz.open(pdf)
    pngs = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(0.65, 0.65), alpha=False)
        p = png_dir / f"slide_{i+1:03d}.png"
        pix.save(p)
        pngs.append(p)
    cols = 5
    tw, th, label_h = 240, 135, 20
    rows = (len(pngs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw, rows * (th + label_h)), "white")
    draw = ImageDraw.Draw(sheet)
    for idx, p in enumerate(pngs):
        img = Image.open(p).convert("RGB")
        img.thumbnail((tw, th))
        x = (idx % cols) * tw + (tw - img.width) // 2
        y = (idx // cols) * (th + label_h) + label_h
        sheet.paste(img, (x, y))
        draw.text(((idx % cols) * tw + 6, (idx // cols) * (th + label_h) + 3), str(idx + 1), fill="black")
    sheet_path = preview / "contact_sheet.png"
    sheet.save(sheet_path)
    return pdf, sheet_path, len(doc)


def main():
    availability_rows = []
    summaries_by_topic = []
    all_summaries = []
    for folder, topic_cn, topic_en in TOPICS:
        topic_dir = ROOT / folder
        data_dir = topic_dir / "downloaded_resources"
        pdf_dir = data_dir / "pdfs"
        html_dir = data_dir / "pmc_html"
        rows = read_csv(data_dir / "pubmed_literature_selected.csv")
        summaries = []
        for row in rows:
            pdf_path, status, detail = download_pdf(row, pdf_dir)
            html_path = save_pmc_html(row, html_dir)
            figures = ensure_pmc_figures(row, html_path, topic_dir)
            summary = build_article_summary(topic_cn, row, pdf_path, html_path, figures)
            summaries.append(summary)
            all_summaries.append(summary)
            availability_rows.append(
                {
                    "topic": topic_cn,
                    "pmid": row.get("pmid", ""),
                    "pmcid": row.get("pmcid", ""),
                    "doi": row.get("doi", ""),
                    "title": row.get("title", ""),
                    "pdf_status": status,
                    "pdf_path": str(pdf_path) if pdf_path else "",
                    "html_path": str(html_path) if html_path else "",
                    "detail": detail,
                }
            )
            print(f"{topic_cn} | {row.get('pmid')} | {status}")
        summaries_by_topic.append((folder, topic_cn, topic_en, summaries))

    write_csv(AVAIL_PATH, availability_rows)
    NOTES_PATH.write_text(json.dumps(all_summaries, ensure_ascii=False, indent=2), encoding="utf-8")
    build_deck(summaries_by_topic)
    pdf, sheet, page_count = build_contact_sheet(PPTX_PATH)

    with zipfile.ZipFile(PPTX_PATH) as zf:
        media_count = len([x for x in zf.namelist() if x.startswith("ppt/media/")])
    downloaded = sum(1 for r in availability_rows if r["pdf_status"] == "downloaded")
    not_public = len(availability_rows) - downloaded
    QA_PATH.write_text(
        "\n".join(
            [
                "# Deep Literature PPT QA",
                "",
                f"- PPTX: {PPTX_PATH}",
                f"- Articles covered: {len(availability_rows)}",
                f"- Slide count expected/rendered: {len(Presentation(PPTX_PATH).slides)} / {page_count or 'not rendered'}",
                f"- Public PDFs downloaded: {downloaded}",
                f"- PDFs not publicly downloadable: {not_public}",
                f"- PDF availability CSV: {AVAIL_PATH}",
                f"- Extracted summary JSON: {NOTES_PATH}",
                f"- Embedded media files: {media_count}",
                f"- LibreOffice PDF preview: {pdf or 'not generated'}",
                f"- Contact sheet: {sheet or 'not generated'}",
                "- Note: Cell/Molecular Cell publisher PDFs that returned 403 were not bypassed; content was summarized from PubMed metadata, PMC HTML when available, and available figures.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
