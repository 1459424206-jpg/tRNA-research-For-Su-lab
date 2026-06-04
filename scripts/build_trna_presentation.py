from __future__ import annotations

import csv
import json
import math
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
PPTX_PATH = OUTPUT / "final_presentation_cn.pptx"
REPORT_PATH = OUTPUT / "trna_four_project_bilingual_report.md"
QA_PATH = OUTPUT / "qa_report.md"


TOPICS = [
    {
        "folder": "biological functions and structural characteristics of tRNA",
        "cn": "tRNA 的生物功能与结构特征",
        "en": "Biological functions and structural characteristics of tRNA",
        "short": "生物功能与结构",
        "theme": [
            "tRNA 不只是“接头分子”，而是连接遗传密码、翻译速度、细胞应激、免疫防御和疾病治疗的调控节点。",
            "结构层面从三叶草二级结构、L 型三级结构、修饰核苷到氨酰化/核糖体复合物共同决定读取精度。",
            "近期高影响力研究集中在 suppressor tRNA、tRNA 尾部切割、tRNA 片段和 tRNA 结构互作组。",
        ],
        "en_detail": [
            "tRNA functions as more than an adaptor: it links genetic decoding, translation kinetics, stress response, immunity, and therapeutic engineering.",
            "Its cloverleaf secondary structure, L-shaped tertiary fold, modified nucleosides, and aminoacylation state jointly determine decoding fidelity.",
            "Recent high-impact work emphasizes suppressor tRNAs, tRNA-tail cleavage, tRNA-derived fragments, and in vivo tRNA structurome/interactome maps.",
        ],
        "figure_hint": "PMC12851939_41586_2025_9852_Fig1_HTML.jpg",
    },
    {
        "folder": "tRNA analysis tools and algorithms",
        "cn": "tRNA 分析工具与算法",
        "en": "tRNA analysis tools and algorithms",
        "short": "工具与算法",
        "theme": [
            "算法核心难点来自 tRNA 的短序列、多拷贝、强修饰、结构相似性以及测序逆转录停顿/错配。",
            "工具链通常包括基因识别、参考库构建、读段比对、修饰信号建模、表达定量和可视化解释。",
            "tRNAscan-SE、mim-tRNAseq、nanopore direct tRNA sequencing 等代表了从注释到定量修饰的关键路线。",
        ],
        "en_detail": [
            "Algorithmic challenges arise from short length, multi-copy genes, heavy modification, structural similarity, and modification-induced RT stops or mismatches.",
            "A practical toolchain covers gene detection, reference construction, read alignment, modification-signal modeling, abundance quantification, and visualization.",
            "tRNAscan-SE, mim-tRNAseq, and direct nanopore tRNA sequencing represent major routes from annotation to quantitative modification profiling.",
        ],
        "figure_hint": "PMC10791586_41587_2023_1743_Fig1_HTML.jpg",
    },
    {
        "folder": "tRNA mod-fingerprint database",
        "cn": "tRNA 修饰指纹数据库",
        "en": "tRNA mod-fingerprint database",
        "short": "修饰指纹数据库",
        "theme": [
            "tRNA 修饰指纹把“哪一种 tRNA、哪个位点、哪类修饰、在哪个条件下变化”组织为可查询的证据结构。",
            "数据库建设需要统一 tRNA 坐标、修饰命名、物种来源、实验技术、置信度和可视化字段。",
            "MODOMICS、RNAcentral 及测序/质谱分析项目可作为设计本地数据库的外部锚点。",
        ],
        "en_detail": [
            "A tRNA modification fingerprint organizes which tRNA, which position, which modification, and which condition into a queryable evidence model.",
            "Database design requires harmonized tRNA coordinates, modification nomenclature, species metadata, experimental technology, confidence scores, and visualization fields.",
            "MODOMICS, RNAcentral, and sequencing/mass-spectrometry projects can anchor a local database schema.",
        ],
        "figure_hint": "PMC12976133_41467_2026_70126_Fig1_HTML.jpg",
    },
    {
        "folder": "tRNA-related wet-lab experiments",
        "cn": "tRNA 相关湿实验",
        "en": "tRNA-related wet-lab experiments",
        "short": "湿实验路线",
        "theme": [
            "湿实验围绕样本质量、tRNA 富集、修饰保真、氨酰化状态保护和测序/质谱平台兼容性展开。",
            "常见实验包括 Northern blot、氨酰化测定、LC-MS/MS、去修饰/酶处理、small RNA-seq 与 nanopore direct RNA sequencing。",
            "关键质控包括 spike-in、去氨酰化对照、酶切效率、RT-stop/mismatch 校准以及 tRNA 片段与成熟 tRNA 的区分。",
        ],
        "en_detail": [
            "Wet-lab work depends on sample integrity, tRNA enrichment, modification preservation, aminoacylation protection, and platform compatibility.",
            "Common assays include Northern blotting, aminoacylation assays, LC-MS/MS, demodification/enzyme treatment, small RNA-seq, and direct nanopore RNA sequencing.",
            "Key controls include spike-ins, deacylation controls, enzyme-efficiency checks, RT-stop/mismatch calibration, and discrimination between tRNA fragments and mature tRNAs.",
        ],
        "figure_hint": "PMC12368100_41467_2025_62545_Fig1_HTML.jpg",
    },
]


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists() or path.stat().st_size == 0:
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_topic(topic: dict) -> dict:
    folder = ROOT / topic["folder"]
    data = folder / "downloaded_resources"
    papers = read_csv(data / "pubmed_literature_selected.csv")
    repos = read_csv(data / "github_repositories_selected.csv")
    manifest_path = data / "figure_asset_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else []
    figure_files = list((data / "figures").glob("*.jpg")) + list((data / "figures").glob("*.png"))
    chosen = None
    for fig in figure_files:
        if topic["figure_hint"] in fig.name:
            chosen = fig
            break
    if chosen is None and figure_files:
        chosen = figure_files[0]
    return {**topic, "folder_path": folder, "papers": papers, "repos": repos, "manifest": manifest, "figure": chosen}


def paper_source_for_figure(topic_data: dict, fig_path: Path | None) -> str:
    if not fig_path:
        return "Source: no article figure available"
    pmcid = fig_path.name.split("_", 1)[0]
    for paper in topic_data["papers"]:
        if paper.get("pmcid") == pmcid:
            return f"Source: {paper.get('title', '')}; {paper.get('journal', '')}, {paper.get('year', '')}; PMID {paper.get('pmid', '')}; PMCID {pmcid}"
    return f"Source: PMC article figure, {pmcid}"


def build_report(topics: list[dict]) -> None:
    lines = [
        "# tRNA research 四个子项目中英文资料报告",
        "",
        "检索日期：2026-05-23（Asia/Shanghai）。",
        "",
        "筛选说明：PubMed 文献以 2020-2026 年为时间窗，优先选择 Nature、Science、Cell、Nature Biotechnology、Nature Methods、Molecular Cell、Nucleic Acids Research、Nature Communications 等高影响力期刊；若某主题不足 15 篇，则按题名/摘要相关性补齐。这里的“影响因子最高”采用期刊层级代理分，不等同于 Clarivate 官方 JIF。",
        "",
        "Selection note: PubMed records were searched from 2020-2026, prioritizing high-impact journals such as Nature, Science, Cell, Nature Biotechnology, Nature Methods, Molecular Cell, Nucleic Acids Research, and Nature Communications. If fewer than 15 records were available for a topic, relevance-based fallback records were added. Journal impact was approximated by a journal-tier proxy, not official Clarivate JIF values.",
        "",
    ]
    for topic in topics:
        lines.extend(
            [
                f"## {topic['cn']}",
                f"**English:** {topic['en']}",
                "",
                "### 中文详细阐述",
            ]
        )
        for bullet in topic["theme"]:
            lines.append(f"- {bullet}")
        lines.extend(["", "### Detailed English Explanation"])
        for bullet in topic["en_detail"]:
            lines.append(f"- {bullet}")
        lines.extend(["", "### 最新高影响力文献 / Recent high-impact literature"])
        for i, paper in enumerate(topic["papers"], 1):
            lines.append(
                f"{i}. {paper.get('title','')} | {paper.get('journal','')} ({paper.get('year','')}) | PMID {paper.get('pmid','')} | DOI {paper.get('doi','') or 'NA'}"
            )
        lines.extend(["", "### GitHub/数据库资源 / GitHub and database resources"])
        if topic["repos"]:
            for repo in topic["repos"]:
                lines.append(
                    f"- {repo.get('name','')} | stars={repo.get('stars','0')} | updated={repo.get('updated_at','')} | {repo.get('url','')}"
                )
        else:
            lines.append("- GitHub API rate/keyword filtering did not return a high-confidence repository for this topic; use PubMed and domain databases as primary sources.")
        lines.extend(["", f"### 图像来源 / Figure source", f"- {paper_source_for_figure(topic, topic['figure'])}", ""])
    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def set_text(frame, text: str, size: int = 16, bold: bool = False, color=(0, 0, 0), align=None):
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)


def add_textbox(slide, x, y, w, h, text, size=16, bold=False, color=(0, 0, 0), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(0.04)
    box.text_frame.margin_right = Inches(0.04)
    box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    set_text(box.text_frame, text, size=size, bold=bold, color=color, align=align)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(6)
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)
    return box


def add_rule(slide, y):
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(y), Inches(12.25), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 0, 0)
    line.line.color.rgb = RGBColor(0, 0, 0)


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.28, 11.9, 0.45, title, size=24, bold=True)
    if subtitle:
        add_textbox(slide, 0.58, 0.76, 11.7, 0.26, subtitle, size=9, color=(70, 70, 70))
    add_rule(slide, 1.06)


def fit_image(slide, image_path: Path, x, y, w, h):
    with Image.open(image_path) as img:
        iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        new_w = w
        new_h = w / img_ratio
    else:
        new_h = h
        new_w = h * img_ratio
    px = x + (w - new_w) / 2
    py = y + (h - new_h) / 2
    return slide.shapes.add_picture(str(image_path), Inches(px), Inches(py), width=Inches(new_w), height=Inches(new_h))


def add_table_like(slide, rows, x, y, w, h, col_widths, font_size=8):
    row_h = h / max(len(rows), 1)
    for r_idx, row in enumerate(rows):
        cy = y + row_h * r_idx
        for c_idx, cell in enumerate(row):
            cx = x + sum(col_widths[:c_idx]) * w
            cw = col_widths[c_idx] * w
            shape = slide.shapes.add_shape(1, Inches(cx), Inches(cy), Inches(cw), Inches(row_h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
            shape.line.color.rgb = RGBColor(180, 180, 180)
            tf = shape.text_frame
            tf.margin_left = Inches(0.03)
            tf.margin_right = Inches(0.03)
            tf.margin_top = Inches(0.02)
            tf.word_wrap = True
            set_text(tf, cell, size=font_size, bold=(r_idx == 0))


def add_process_slide(prs, topics):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "四个文件夹形成一条 tRNA 研究链路", "From folders to an integrated tRNA research workflow")
    steps = [
        ("结构/功能", "定义生物问题"),
        ("工具/算法", "把测序与注释转成可比较数据"),
        ("修饰数据库", "沉淀位点、条件和证据等级"),
        ("湿实验", "验证机制并校准数据偏差"),
    ]
    x0 = 0.75
    for i, (head, body) in enumerate(steps):
        x = x0 + i * 3.05
        shape = slide.shapes.add_shape(1, Inches(x), Inches(2.15), Inches(2.45), Inches(1.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(0, 0, 0)
        set_text(shape.text_frame, f"{i+1}. {head}\n{body}", size=15, bold=True)
        if i < len(steps) - 1:
            add_textbox(slide, x + 2.52, 2.72, 0.4, 0.3, ">", size=22, bold=True, align=PP_ALIGN.CENTER)
    add_bullets(
        slide,
        0.8,
        4.55,
        11.7,
        1.3,
        [
            "建议把四个子项目视为一个闭环：文献提出假设，工具产生特征，数据库固化证据，湿实验验证因果。",
            "PPT 中所有原始图均来自 PMC 开放页面或保存的文章图像文件，并在图下标出来源。",
        ],
        size=15,
    )


def build_ppt(topics: list[dict]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, 0.7, 1.0, 11.8, 0.7, "tRNA research 四个子项目资料汇报", size=30, bold=True)
    add_textbox(slide, 0.72, 1.82, 11.3, 0.32, "基于 PubMed 与 GitHub 的资料筛选、文献高亮、图像证据与研究路线整合", size=14)
    add_textbox(slide, 0.72, 2.26, 11.3, 0.25, "白底黑字版 | 2026-05-23 | 中文主讲", size=10, color=(70, 70, 70))
    add_rule(slide, 5.95)
    add_textbox(slide, 0.72, 6.15, 11.8, 0.55, "四个文件夹：结构功能、工具算法、修饰指纹数据库、湿实验", size=16, bold=True)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "资料筛选策略", "PubMed literature plus GitHub resources")
    add_bullets(
        slide,
        0.75,
        1.55,
        11.6,
        4.8,
        [
            "PubMed：每个主题保留 15 篇 2020-2026 年候选文献，优先高影响力期刊；不足时按题名/摘要相关性补齐。",
            "GitHub：按文件夹主题关键词检索仓库，保存 stars、语言、更新时间、许可证与链接；去除了 transformation/transformer 等拼写噪音。",
            "图像：优先抓取开放 PMC 文章页中的原始 figure 图像；每张图均在 PPT 里标注 PMID/PMCID/期刊来源。",
            "限制：未调用 Clarivate JCR 官方数据库，因此“影响因子最高”在本资料包中体现为高影响力期刊代理排序。",
        ],
        size=17,
    )

    add_process_slide(prs, topics)

    for topic in topics:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, topic["cn"], topic["en"])
        add_textbox(slide, 0.75, 1.35, 3.0, 0.35, "中文阐述", size=16, bold=True)
        add_bullets(slide, 0.75, 1.78, 5.65, 2.65, topic["theme"], size=14)
        add_textbox(slide, 6.9, 1.35, 3.4, 0.35, "English explanation", size=14, bold=True)
        add_bullets(slide, 6.9, 1.78, 5.65, 2.65, topic["en_detail"], size=11)
        add_textbox(slide, 0.75, 5.15, 11.6, 0.6, "本页目的：先把文件夹命名转成研究问题，再把文献、代码资源和实验/数据库路线接到同一个框架中。", size=14, bold=True)

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, f"{topic['short']}：最新高影响力文献清单", "Top recent high-impact PubMed records")
        rows = [["#", "题名", "期刊/年份", "PMID/DOI"]]
        for i, paper in enumerate(topic["papers"][:10], 1):
            rows.append(
                [
                    str(i),
                    paper.get("title", "")[:90],
                    f"{paper.get('journal','')[:28]} / {paper.get('year','')}",
                    f"{paper.get('pmid','')}\n{paper.get('doi','')[:30]}",
                ]
            )
        add_table_like(slide, rows, 0.55, 1.35, 12.25, 5.45, [0.05, 0.56, 0.22, 0.17], font_size=7)
        add_textbox(slide, 0.6, 6.95, 12.0, 0.22, "完整 15 篇及摘要字段见各子文件夹 downloaded_resources/pubmed_literature_selected.csv", size=8, color=(90, 90, 90))

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, f"{topic['short']}：文章图像证据", "Representative figure from screened literature")
        if topic["figure"] and topic["figure"].exists():
            fit_image(slide, topic["figure"], 0.65, 1.35, 8.4, 4.85)
            add_bullets(
                slide,
                9.35,
                1.55,
                3.3,
                3.8,
                [
                    "这张图作为本主题的视觉锚点，用于解释研究设计、关键机制或技术路线。",
                    "汇报时建议只讲图中与本文件夹主题直接相关的 1-2 个结论，避免逐 panel 读图。",
                    "原始科学数据未被重绘；这里只做缩放嵌入。",
                ],
                size=12,
            )
            add_textbox(slide, 0.68, 6.32, 11.9, 0.5, paper_source_for_figure(topic, topic["figure"]), size=7, color=(80, 80, 80))
        else:
            add_textbox(slide, 0.8, 2.3, 11.5, 1.0, "该主题未成功抓取开放文章图像；已保留文献与仓库清单，可后续手动补充出版商图或作者授权图。", size=16)

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, f"{topic['short']}：GitHub 与可复用资源", "GitHub repositories and reusable assets")
        if topic["repos"]:
            rows = [["仓库", "stars", "语言", "用途判断"]]
            for repo in topic["repos"][:8]:
                rows.append(
                    [
                        repo.get("name", ""),
                        repo.get("stars", ""),
                        repo.get("language", "") or "NA",
                        (repo.get("description", "") or "无描述")[:82],
                    ]
                )
            add_table_like(slide, rows, 0.65, 1.45, 12.0, 3.8, [0.28, 0.08, 0.12, 0.52], font_size=8)
        else:
            add_textbox(slide, 0.8, 1.8, 11.5, 0.8, "本主题没有高置信 GitHub 仓库；建议优先使用 PubMed 文献与领域数据库，并在后续人工补充作者代码。", size=16)
        add_bullets(
            slide,
            0.8,
            5.55,
            11.5,
            1.0,
            [
                "使用建议：先复现 README 中的最小示例，再把输入/输出字段映射到本项目的 tRNA 坐标、修饰位点或实验批次。",
                "保存路径：每个子文件夹的 downloaded_resources/ 下包含 CSV、JSON、图像 manifest 与可追溯链接。",
            ],
            size=13,
        )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "综合结论：四个子项目应合并为证据闭环", "Integrated conclusion")
    add_bullets(
        slide,
        0.8,
        1.55,
        11.5,
        4.6,
        [
            "结构功能主题提供生物学假设：哪些 tRNA、修饰或片段可能改变翻译、免疫或疾病表型。",
            "工具算法主题把高通量数据转成可比较特征：丰度、错配、RT-stop、修饰概率、结构互作与片段谱。",
            "修饰指纹数据库主题负责证据沉淀：统一位点坐标、文献来源、实验方法、置信度和查询接口。",
            "湿实验主题负责验证：通过 LC-MS/MS、Northern、氨酰化测定和定向测序校准算法输出。",
        ],
        size=17,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "后续建议", "Suggested next steps")
    add_bullets(
        slide,
        0.8,
        1.55,
        11.5,
        4.6,
        [
            "如果要继续深化，建议先确定一个核心物种/疾病/实验条件，否则四个主题会很容易扩散。",
            "数据库字段优先级：tRNA ID、isoacceptor/isodecoder、修饰位点、修饰类型、检测技术、PMID、样本条件、置信度。",
            "实验优先级：先做可重复的 tRNA abundance/modification profiling，再针对关键位点做验证实验。",
            "文献更新：建议每月按相同脚本刷新 PubMed 与 GitHub，避免综述和工具版本过期。",
        ],
        size=17,
    )

    prs.save(PPTX_PATH)


def build_qa(topics: list[dict]) -> None:
    with zipfile.ZipFile(PPTX_PATH) as zf:
        media = [name for name in zf.namelist() if name.startswith("ppt/media/")]
    prs = Presentation(PPTX_PATH)
    lines = [
        "# QA Report",
        "",
        f"- PPTX created: {PPTX_PATH}",
        f"- Slide count: {len(prs.slides)}",
        f"- Embedded media files: {len(media)}",
        f"- Bilingual report: {REPORT_PATH}",
        "- Verification: reopened with python-pptx and inspected PPTX package media entries.",
        "- Style: white background, black text, no decorative gradients or stock images.",
        "- Known limitation: journal impact ranking uses a high-impact journal proxy rather than official Clarivate JIF values.",
    ]
    for topic in topics:
        lines.append(f"- {topic['cn']}: papers={len(topic['papers'])}, repos={len(topic['repos'])}, figure={topic['figure'].name if topic['figure'] else 'NA'}")
    QA_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    topics = [load_topic(t) for t in TOPICS]
    build_report(topics)
    build_ppt(topics)
    build_qa(topics)
    print(f"Wrote {PPTX_PATH}")
    print(f"Wrote {REPORT_PATH}")
    print(f"Wrote {QA_PATH}")


if __name__ == "__main__":
    main()
