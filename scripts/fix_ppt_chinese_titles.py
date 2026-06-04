from __future__ import annotations

import json
import re
import time
from pathlib import Path

from deep_translator import GoogleTranslator
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from translate_ppt_to_chinese import normalize_abbreviations, postprocess_translation, set_para_text, should_translate


ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "output" / "final_presentation_cn_deep_literature_zh_only.pptx"
CACHE = ROOT / "output" / "translation_cache_zh.json"
REPORT = ROOT / "output" / "translation_title_fix_report.json"

JOURNAL_ZH = {
    "Nature biotechnology": "《自然·生物技术》",
    "Nature Biotechnology": "《自然·生物技术》",
    "Nature communications": "《自然·通讯》",
    "Nature Communications": "《自然·通讯》",
    "Nature genetics": "《自然·遗传学》",
    "Nature Genetics": "《自然·遗传学》",
    "Nature structural & molecular biology": "《自然·结构与分子生物学》",
    "Nature Structural & Molecular Biology": "《自然·结构与分子生物学》",
    "Molecular cell": "《分子细胞》",
    "Molecular Cell": "《分子细胞》",
    "Cell": "《细胞》",
    "Nature": "《自然》",
}


def iter_text_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_shapes(shape.shapes)
        elif hasattr(shape, "text_frame"):
            yield shape


def translate_fragment(fragment: str, translator: GoogleTranslator, cache: dict[str, str]) -> str:
    fragment = normalize_abbreviations(fragment).strip()
    if not fragment:
        return fragment
    cache_key = "FRAG::" + fragment
    if cache_key in cache:
        return cache[cache_key]
    try:
        translated = translator.translate(fragment)
        translated = postprocess_translation(fragment, translated)
    except Exception:
        translated = fragment
    cache[cache_key] = translated
    time.sleep(0.08)
    return translated


def local_metadata_zh(text: str) -> str:
    text = normalize_abbreviations(text)
    text = text.replace("experimental", "实验研究").replace("review", "综述")
    text = text.replace("Source:", "来源：")
    text = text.replace("PMCID NA", "PMCID 无").replace("DOI NA", "DOI 无")
    for en, zh in sorted(JOURNAL_ZH.items(), key=lambda kv: len(kv[0]), reverse=True):
        text = re.sub(rf"\b{re.escape(en)}\b", zh, text)
    text = text.replace("；PMCID NA", "")
    return text


def fix_text(text: str, translator: GoogleTranslator, cache: dict[str, str]) -> tuple[str, bool]:
    original = text
    text = local_metadata_zh(text)

    # Mixed Chinese heading followed by an English paper title.
    m = re.match(r"^(\d+(?:\.\d+)?\s+[^：]{2,20}：)(.+)$", text)
    if m:
        prefix, rest = m.groups()
        if len(re.findall(r"[A-Za-z]", rest)) >= 8:
            rest_zh = translate_fragment(rest, translator, cache)
            text = prefix + rest_zh

    # Section bibliography entries like "4. AAV-delivered ...".
    m = re.match(r"^(\d+\.\s+)([A-Za-z].+)$", text)
    if m:
        prefix, rest = m.groups()
        text = prefix + translate_fragment(rest, translator, cache)

    # Any remaining mostly-English text that is not a metadata/source line.
    is_metadata = bool(re.search(r"(PMID|PMCID|DOI|10\.\d{4,9}/)", text))
    if not is_metadata and should_translate(text):
        text = translate_fragment(text, translator, cache)

    return text, text != original


def main() -> None:
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    prs = Presentation(PPT)
    translator = GoogleTranslator(source="auto", target="zh-CN")
    changed = 0
    for slide in prs.slides:
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if not text:
                    continue
                fixed, did = fix_text(text, translator, cache)
                if did:
                    set_para_text(para, fixed)
                    changed += 1
    prs.save(PPT)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    check = Presentation(PPT)
    remaining = []
    for si, slide in enumerate(check.slides, 1):
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if not text:
                    continue
                if should_translate(text) and not re.search(r"(PMID|PMCID|DOI|10\.\d{4,9}/)", text):
                    remaining.append({"slide": si, "text": text})
    REPORT.write_text(
        json.dumps(
            {
                "ppt": str(PPT),
                "changed_paragraphs": changed,
                "remaining_non_metadata_english_like_count": len(remaining),
                "remaining_sample": remaining[:80],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Changed paragraphs: {changed}")
    print(f"Remaining non-metadata English-like: {len(remaining)}")


if __name__ == "__main__":
    main()
