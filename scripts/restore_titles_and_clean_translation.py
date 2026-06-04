from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from translate_ppt_chinese_body_keep_names import (
    DST,
    REPORT,
    SRC,
    is_metadata_or_source,
    iter_text_shapes,
    localize_metadata,
    looks_like_paper_title_line,
)
from translate_ppt_to_chinese import normalize_abbreviations, set_para_text


ROOT = Path(__file__).resolve().parents[1]
CLEAN_REPORT = ROOT / "output" / "translation_final_clean_report.json"


def collect_paragraphs(prs: Presentation):
    out = []
    for slide in prs.slides:
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                out.append(para)
    return out


def is_mixed_title(cur: str, orig: str) -> bool:
    cur = normalize_abbreviations(cur)
    orig = normalize_abbreviations(orig)
    if looks_like_paper_title_line(orig):
        return True
    if re.match(r"^\d+\.\d+\s+(文献概览|材料方法与结果|讨论与局限|综述主要论点|综合观点与空白)：", cur):
        return True
    return False


def restore_title(cur: str, orig: str) -> str:
    cur = normalize_abbreviations(cur)
    orig = normalize_abbreviations(orig)
    # For slide titles, preserve the current Chinese section prefix, but restore
    # the English literature title from the original after the colon.
    cm = re.match(r"^(\d+\.\d+\s+[^：]+：).+$", cur)
    om = re.match(r"^\d+\.\d+\s+[^：]+：(.+)$", orig)
    if cm and om:
        return cm.group(1) + om.group(1)
    return orig


def clean_text(cur: str, orig: str) -> str:
    cur = normalize_abbreviations(cur)
    orig = normalize_abbreviations(orig)
    if "Error 500" in cur or "Server Error" in cur:
        return "该段网页/PDF 自动抽取失败；请以原文 PDF、DOI 或 PubMed 记录为准。"
    if "__TERM" in cur:
        return restore_title(cur, orig) if is_mixed_title(cur, orig) else orig
    if is_mixed_title(cur, orig):
        return restore_title(cur, orig)
    if is_metadata_or_source(orig):
        return localize_metadata(orig)
    return cur


def main() -> None:
    src = Presentation(SRC)
    dst = Presentation(DST)
    src_paras = collect_paragraphs(src)
    dst_paras = collect_paragraphs(dst)
    changed = 0
    for src_para, dst_para in zip(src_paras, dst_paras):
        orig = "".join(run.text for run in src_para.runs).strip()
        cur = "".join(run.text for run in dst_para.runs).strip()
        if not cur:
            continue
        fixed = clean_text(cur, orig)
        if fixed != cur:
            set_para_text(dst_para, fixed)
            changed += 1
    dst.save(DST)

    check = Presentation(DST)
    bad = []
    title_translated = []
    for si, slide in enumerate(check.slides, 1):
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if "Error 500" in text or "__TERM" in text:
                    bad.append({"slide": si, "text": text})
                if re.match(r"^\d+\.\d+\s+[^：]+：", text) and len(re.findall(r"[A-Za-z]", text)) < 8:
                    title_translated.append({"slide": si, "text": text})
    CLEAN_REPORT.write_text(
        json.dumps(
            {
                "ppt": str(DST),
                "changed_paragraphs": changed,
                "bad_placeholder_or_error_count": len(bad),
                "bad_sample": bad[:40],
                "possibly_translated_title_count": len(title_translated),
                "possibly_translated_title_sample": title_translated[:40],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Cleaned paragraphs: {changed}")
    print(f"Remaining placeholders/errors: {len(bad)}")
    print(f"Possibly translated titles: {len(title_translated)}")


if __name__ == "__main__":
    main()
