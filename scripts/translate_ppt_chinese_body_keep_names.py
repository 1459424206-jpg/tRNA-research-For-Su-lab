from __future__ import annotations

import json
import re
import subprocess
import time
from pathlib import Path

import fitz
from deep_translator import GoogleTranslator
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from translate_ppt_to_chinese import normalize_abbreviations, postprocess_translation, set_para_text


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "output" / "final_presentation_cn_deep_literature.pptx"
DST = ROOT / "output" / "final_presentation_cn_deep_literature_zh_body_keep_names.pptx"
CACHE = ROOT / "output" / "translation_cache_zh_keep_names.json"
REPORT = ROOT / "output" / "translation_zh_body_keep_names_report.json"
PREVIEW = ROOT / "output" / "preview_zh_body_keep_names"
SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")


LOCAL_REPLACEMENTS = {
    "Deck logic": "PPT 制作逻辑",
    "Biological functions and structural characteristics of tRNA": "tRNA 的生物功能与结构特征",
    "tRNA analysis tools and algorithms": "tRNA 分析工具与算法",
    "tRNA mod-fingerprint database": "tRNA 修饰指纹数据库",
    "tRNA-related wet-lab experiments": "tRNA 相关湿实验",
    "experimental": "实验研究",
    "review": "综述",
    "Source:": "来源：",
    "no article figure available": "无可用文章图像",
}


def iter_text_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_shapes(shape.shapes)
        elif hasattr(shape, "text_frame"):
            yield shape


def looks_like_paper_title_line(text: str) -> bool:
    t = normalize_abbreviations(text).strip()
    if re.match(r"^\d+\.\s+[A-Z0-9].+[.!?]$", t):
        return True
    if re.match(r"^\d+\.\d+\s+(文献概览|材料方法与结果|讨论与局限|综述主要论点)：", t):
        return True
    if re.match(r"^第\s*\d+\s*部分", t):
        return False
    return False


def is_metadata_or_source(text: str) -> bool:
    return bool(re.search(r"\b(PMID|PMCID|DOI)\b|10\.\d{4,9}/|\|\s*(Nature|Cell|Molecular|Nucleic)", text))


def localize_metadata(text: str) -> str:
    out = normalize_abbreviations(text)
    for old, new in LOCAL_REPLACEMENTS.items():
        out = re.sub(re.escape(old), new, out, flags=re.I if old in {"experimental", "review"} else 0)
    out = out.replace("PMCID NA", "PMCID 无").replace("DOI NA", "DOI 无")
    out = out.replace("；PMCID NA", "")
    return out


def should_translate_body(text: str) -> bool:
    t = normalize_abbreviations(text).strip()
    if not t:
        return False
    if looks_like_paper_title_line(t) or is_metadata_or_source(t):
        return False
    letters = len(re.findall(r"[A-Za-z]", t))
    cjk = len(re.findall(r"[\u4e00-\u9fff]", t))
    return letters >= 14 and letters > cjk


def protect_terms(text: str) -> tuple[str, dict[str, str]]:
    # Protect likely gene/protein names, RNA species, database IDs, plasmid names,
    # and journal/source identifiers. This keeps scientific names readable.
    patterns = [
        r"\b(?:tRNA|mRNA|rRNA|ncRNA|RNA|DNA|AAV|UGA|CRISPR-Cas|LC-MS/MS|LC-MS|RT|PMID|PMCID|DOI|PMC)\b",
        r"\b[A-Z][A-Z0-9-]{1,}(?:\b|(?=[,.;:)]))",
        r"\b[A-Z][A-Za-z]*\d+[A-Za-z0-9-]*\b",
        r"\b(?:m1A|m3C|m6A|m7G|N7-methylguanosine)\b",
        r"\b(?:Cas12a3|PARIS|METTL1|METTL6|SerRS|AARS1|NAT10|ALKB-1|PYROXD1|Trl1|EF-Tu|Ocr|sfGFP)\b",
        r"\b10\.\d{4,9}/[A-Za-z0-9._;()/:+-]+\b",
    ]
    protected: dict[str, str] = {}
    out = text
    tokens: list[str] = []
    for pat in patterns:
        tokens.extend(re.findall(pat, out))
    # Longest first avoids replacing substrings inside longer names.
    for i, token in enumerate(sorted(set(tokens), key=len, reverse=True)):
        key = f"__TERM{i}__"
        protected[key] = token
        out = out.replace(token, key)
    return out, protected


def unprotect_terms(text: str, protected: dict[str, str]) -> str:
    out = text
    for key, token in protected.items():
        out = out.replace(key, token)
        out = out.replace(key.lower(), token)
        out = out.replace(key.replace("_", " "), token)
    return out


def translate_body(text: str, translator: GoogleTranslator, cache: dict[str, str]) -> str:
    text = normalize_abbreviations(text)
    if text in cache:
        return cache[text]
    protected_text, protected = protect_terms(text)
    try:
        translated = translator.translate(protected_text)
        translated = unprotect_terms(translated, protected)
        translated = postprocess_translation(text, translated)
    except Exception:
        translated = text
    cache[text] = translated
    time.sleep(0.08)
    return translated


def transform_text(text: str, translator: GoogleTranslator, cache: dict[str, str]) -> tuple[str, bool]:
    original = text
    text = localize_metadata(text)
    if looks_like_paper_title_line(text):
        # Keep the paper title itself as the original English title.
        text = normalize_abbreviations(text)
    elif should_translate_body(text):
        text = translate_body(text, translator, cache)
    return text, text != original


def render_preview() -> tuple[str, str, int] | tuple[str, str, str]:
    PREVIEW.mkdir(parents=True, exist_ok=True)
    if not SOFFICE.exists():
        return "", "", "LibreOffice not found"
    subprocess.run(
        [str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(PREVIEW), str(DST)],
        cwd=str(ROOT),
        check=False,
        timeout=240,
    )
    pdf = PREVIEW / f"{DST.stem}.pdf"
    if not pdf.exists():
        return "", "", "PDF not generated"
    doc = fitz.open(pdf)
    png_dir = PREVIEW / "slides_png"
    png_dir.mkdir(exist_ok=True)
    pngs = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(0.55, 0.55), alpha=False)
        p = png_dir / f"slide_{i+1:03d}.png"
        pix.save(p)
        pngs.append(p)
    cols, tw, th, label_h = 5, 220, 124, 18
    rows = (len(pngs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw, rows * (th + label_h)), "white")
    draw = ImageDraw.Draw(sheet)
    for idx, p in enumerate(pngs):
        img = Image.open(p).convert("RGB")
        img.thumbnail((tw, th))
        x = (idx % cols) * tw + (tw - img.width) // 2
        y = (idx // cols) * (th + label_h) + label_h
        sheet.paste(img, (x, y))
        draw.text(((idx % cols) * tw + 5, (idx // cols) * (th + label_h) + 2), str(idx + 1), fill="black")
    sheet_path = PREVIEW / "contact_sheet.png"
    sheet.save(sheet_path)
    return str(pdf), str(sheet_path), len(doc)


def main() -> None:
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    prs = Presentation(SRC)
    translator = GoogleTranslator(source="auto", target="zh-CN")
    changed = 0
    translated = 0
    skipped_titles = 0
    for slide in prs.slides:
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if not text:
                    continue
                if looks_like_paper_title_line(text):
                    skipped_titles += 1
                fixed, did = transform_text(text, translator, cache)
                if did:
                    set_para_text(para, fixed)
                    changed += 1
                    if should_translate_body(text):
                        translated += 1
    prs.save(DST)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    pdf, sheet, pages = render_preview()

    # QA: remaining non-metadata English-like text is expected to be titles/names.
    check = Presentation(DST)
    remaining = []
    for si, slide in enumerate(check.slides, 1):
        for shape in iter_text_shapes(slide.shapes):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if should_translate_body(text):
                    remaining.append({"slide": si, "text": text})
    REPORT.write_text(
        json.dumps(
            {
                "source": str(SRC),
                "output": str(DST),
                "changed_paragraphs": changed,
                "translated_body_paragraphs": translated,
                "paper_title_lines_preserved": skipped_titles,
                "remaining_body_english_like_count": len(remaining),
                "remaining_body_english_like_sample": remaining[:80],
                "preview_pdf": pdf,
                "contact_sheet": sheet,
                "rendered_pages": pages,
                "rule": "Journal names, protein names, gene names, and literature titles are preserved in English; explanatory body text is translated into Chinese.",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Wrote {DST}")
    print(f"Changed paragraphs: {changed}; body translated: {translated}; title lines preserved: {skipped_titles}")
    print(f"Rendered pages: {pages}; contact sheet: {sheet}")


if __name__ == "__main__":
    main()
