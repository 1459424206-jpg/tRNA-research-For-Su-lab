from pathlib import Path
import json
import math
import shutil
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


BASE = Path(r"D:\Project\tRNA research\tRNA analysis tools and algorithms")
RES = BASE / "downloaded_resources"
OUT = BASE / "output_tRNA_tools_ppt"
ASSETS = OUT / "assets" / "figures"
ASSETS.mkdir(parents=True, exist_ok=True)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


local_lit = {p["pmid"]: p for p in load_json(RES / "pubmed_literature_selected.json")}

added_lit = {
    "34417604": {
        "pmid": "34417604",
        "title": "tRNAscan-SE 2.0: improved detection and functional classification of transfer RNA genes",
        "journal": "Nucleic Acids Research",
        "year": "2021",
        "doi": "10.1093/nar/gkab688",
        "role": "新增：tRNA 基因识别与功能分类的底层工具。",
    },
    "36869120": {
        "pmid": "36869120",
        "title": "A standardized ontology for naming tRNA-derived RNAs based on molecular origin",
        "journal": "Nature Methods",
        "year": "2023",
        "doi": "10.1038/s41592-023-01801-8",
        "role": "新增：tRNA-derived RNA 的命名、注释与跨研究比较基础。",
    },
    "39952700": {
        "pmid": "39952700",
        "title": "Analyzing, visualizing, and annotating tRNA-derived RNAs using tRAX and tDRnamer",
        "journal": "Current Protocols",
        "year": "2025",
        "doi": "",
        "role": "新增：tDR 小 RNA 数据分析协议与工具化入口。",
    },
}


def paper(pmid):
    return local_lit.get(pmid) or added_lit[pmid]


PAPER_CARDS = [
    {
        "pmid": "34417604",
        "status": "主线新增",
        "theme": "基因识别",
        "logic": "先用高速候选扫描缩小搜索空间，再用协方差模型和特征分类确认 tRNA 基因、伪基因与功能类别。",
        "hypothesis": "tRNA 的保守二级结构和身份元件足以支持跨基因组、高灵敏度的自动识别与分类。",
        "methods": ["候选区域扫描", "Infernal/协方差模型评分", "isotype 与 anticodon 分类", "伪基因和线粒体 tRNA 过滤"],
        "results": "提升复杂基因组中 tRNA 与 tRNA-like 元件的识别能力，是后续 tRNA-seq 注释与参考库构建的上游基础。",
        "discussion": "局限在于表达状态、修饰和样本特异性无法由基因组序列直接推断；后续应与测序证据和 pangenome/reference graph 结合。",
        "visual": "pipeline",
    },
    {
        "pmid": "32796835",
        "status": "主线保留",
        "theme": "丰度定量",
        "logic": "QuantM-tRNA-seq 通过成熟 tRNA 建库、环化/接头策略与全长读段来量化组织间 tRNA 表达差异。",
        "hypothesis": "成熟 tRNA 的丰度与修饰差异可通过专门建库和比对策略从常规测序偏差中恢复。",
        "methods": ["tRNA 末端处理与接头连接", "RT 与 PCR 建库", "tRNA 专用参考比对", "组织/等解码 tRNA 差异比较"],
        "results": "小鼠组织间 isodecoder 表达和修饰痕迹差异显著，但 anticodon pool 受到缓冲。",
        "discussion": "短读长仍受多重比对、RT stop 与高相似序列限制；需要 spike-in、UMI 与概率分配模型提升绝对定量。",
        "visual": "PMC7428014",
    },
    {
        "pmid": "33581077",
        "status": "主线保留",
        "theme": "修饰诱导错配",
        "logic": "mim-tRNAseq 将修饰导致的 misincorporation 从噪声转化为定量特征，同时估计丰度和修饰状态。",
        "hypothesis": "不同修饰在逆转录时形成可重复的错配/停顿指纹，能用于高分辨率地追踪 tRNA 修饰状态。",
        "methods": ["全长 cDNA 建库", "错配/RT-stop 统计", "isodecoder 层级归并", "跨物种可迁移分析工具包"],
        "results": "在酵母、果蝇和人细胞中量化 tRNA abundance 与 modification status，发现细胞系特异 isodecoder pool 和位点间修饰联动。",
        "discussion": "优点是短读长平台可用；缺陷是对 RT 酶、参考注释和修饰指纹库依赖强，未知修饰难以唯一归因。",
        "visual": "pipeline",
    },
    {
        "pmid": "37024678",
        "status": "主线保留",
        "theme": "纳米孔全长读段",
        "logic": "Nano-tRNAseq 直接读取 native tRNA，并通过 raw signal 重处理恢复被默认流程丢弃的短 tRNA 读段。",
        "hypothesis": "原始电流信号中保留了 tRNA 丰度和修饰状态；定制处理能减少长度和修饰造成的系统偏差。",
        "methods": ["native tRNA 纳米孔测序", "raw signal reprocessing", "tRNA mapping", "错配/信号异常推断修饰"],
        "results": "重处理使可用 tRNA reads 增加约 12 倍，并捕获氧化应激下 tRNA 丰度与修饰联动。",
        "discussion": "纳米孔优势是长读段和单分子层级；短 tRNA、接头效率、basecaller 偏差和修饰训练集仍是主要瓶颈。",
        "visual": "PMC10791586",
    },
    {
        "pmid": "39865096",
        "status": "主线保留",
        "theme": "RT readthrough 修饰图谱",
        "logic": "Induro-tRNAseq 选择更强 readthrough 的 group-II intron RT，将 RT stop 降低但保留 misincorporation 信号。",
        "hypothesis": "改变 RT 酶的停顿行为可以把“读不穿”的修饰位点转化为可比较的错配信号。",
        "methods": ["Induro RT 时间梯度", "Induro 与对照 RT 平行比较", "错配频率矩阵", "组织/细胞间修饰差异分析"],
        "results": "在多个人细胞系和小鼠组织中绘制 tRNA 修饰，显示遗传密码读取相关修饰更稳定，其他位点更动态。",
        "discussion": "强项是全基因组修饰 profiling；短板是修饰类别归因仍需标准品/酶缺失验证，适合发展联合修饰 caller。",
        "visual": "PMC11770116",
    },
    {
        "pmid": "40835813",
        "status": "主线保留",
        "theme": "氨酰化状态",
        "logic": "aa-tRNA-seq 用化学连接保留 aminoacylated tRNA，再用纳米孔信号和机器学习识别单分子氨基酸身份。",
        "hypothesis": "氨基酸夹在 tRNA 与接头之间会产生可学习的电流扰动，可在单分子层面解析 charging fidelity。",
        "methods": ["aminoacyl-tRNA 化学保护/连接", "nanopore sequencing", "特征提取", "机器学习 amino acid classifier"],
        "results": "能识别 tRNA 所携带氨基酸，并用于分析修饰酶缺失导致的 hypomodification 和 tRNA 稳定性变化。",
        "discussion": "方法打开 charging 分析入口；局限是模型训练、氨基酸类别泛化、化学处理效率和复杂样本校准。",
        "visual": "PMC12368100",
    },
    {
        "pmid": "40447571",
        "status": "主线保留",
        "theme": "结构与互作",
        "logic": "DM-DMS-MaPseq 将体内 DMS 化学探针、demethylase 处理和 mutational profiling 结合，解析 tRNA structurome/interactome。",
        "hypothesis": "体内 DMS 反应性不仅反映结构可及性，也反映蛋白/核糖体互作遮蔽，可用于追踪应激调控。",
        "methods": ["in vivo DMS probing", "DM 去甲基化处理", "MaP-seq 突变读出", "胞质/线粒体 tRNA 结构比较"],
        "results": "体内 tRNA 结构总体稳定，但与体外谱差异显著；砷酸盐应激改变胞质和线粒体 tRNA 互作。",
        "discussion": "结构读出与互作读出耦合是优势也是混杂来源；未来需联合 Ribo-seq/IP 或扰动实验解卷积。",
        "visual": "PMC12125218",
    },
    {
        "pmid": "42062277",
        "status": "主线保留",
        "theme": "低输入 tRNA-seq",
        "logic": "ORACLE-tRNAseq 面向极少量卵母细胞/早期胚胎，优化低输入捕获并与多组学和 Ribo-seq 联合分析。",
        "hypothesis": "早期胚胎 tRNA pool 的动态变化可与 ZGA、染色质重塑和翻译效率相互协调。",
        "methods": ["低输入建库", "tRNA repertoire 定量", "pseudogene activation 分析", "Ribo-seq/组蛋白修饰整合"],
        "results": "从 oocyte 到 blastocyst 绘制 tRNA landscape，发现 4-cell 阶段 tRNA pseudogene 上调和 translation machinery 建立。",
        "discussion": "低输入是突出贡献；算法挑战在于 dropout、批次效应和跨组学时间点对齐。",
        "visual": "pipeline",
    },
    {
        "pmid": "39096899",
        "status": "主线保留",
        "theme": "tDR 捕获",
        "logic": "LIDAR 通过 quasi-random priming 与 template switching 绕开 3' 端连接依赖，捕获有 blocked 3' termini 的 tDR。",
        "hypothesis": "传统 ligation-based small RNA-seq 系统性漏检末端被修饰或封闭的 RNA，小 RNA 组成被低估。",
        "methods": ["ligation-independent 建库", "quasi-random RT primer", "template switching", "tDR 与 miRNA/read coverage 对比"],
        "results": "LIDAR 检出更多 tRNA-derived RNA，尤其是传统方法难以捕获的 blocked 3' end tDR。",
        "discussion": "适合发现层面；但 tDR 起源解析仍需要 tDRnamer/tRAX、tRNA reference 和多重比对模型支持。",
        "visual": "PMC11455606",
    },
    {
        "pmid": "36869120",
        "status": "主线新增",
        "theme": "tDR 命名与注释",
        "logic": "tDRnamer 把 tRNA-derived RNA 的命名锚定到分子来源、起止位置和 tRNA 层级，减少跨研究不可比。",
        "hypothesis": "统一 ontology 是 tDR 定量、差异分析和数据库整合的前提，而不是报告阶段的格式问题。",
        "methods": ["tRNA 来源层级定义", "片段坐标规则", "多等位/多位点歧义处理", "标准化名称输出"],
        "results": "为 tDR/tRF 结果提供机器可读、可比较的命名体系，能作为 LIDAR、small RNA-seq 和 tRAX 的下游桥梁。",
        "discussion": "命名不能替代准确 mapping；未来应把命名不确定性传递到差异分析和功能富集。",
        "visual": "pipeline",
    },
    {
        "pmid": "39952700",
        "status": "主线新增",
        "theme": "tRAX/tDR 分析协议",
        "logic": "tRAX/tDRnamer 将 small RNA reads 的 tRNA 来源识别、注释、可视化和命名整合为可复现流程。",
        "hypothesis": "tDR 分析需要同时解决 read trimming、非 tRNA 污染、tRNA 参考归属和标准命名。",
        "methods": ["small RNA QC/trimming", "多级参考比对", "tRNA/tDR 分类", "coverage、长度、片段类别可视化"],
        "results": "使 tDR 数据从单纯 counts 表转为可解释的来源、位置、类型和样本间变化。",
        "discussion": "短读长与重复 tRNA 家族仍限制唯一定位；应引入概率归属、UMI 与 full-length evidence 校正。",
        "visual": "pipeline",
    },
    {
        "pmid": "35032425",
        "status": "背景保留",
        "theme": "修饰动态综述",
        "logic": "综述 tRNA modification dynamics 从单生物体扩展到微生物群落 metaepitranscriptomics。",
        "hypothesis": "tRNA 修饰不是静态装饰，而是环境响应、翻译调控和 tRF 生成的动态调节层。",
        "methods": ["汇总测序方法", "动态修饰功能框架", "microbiome/metaepitranscriptomics 概念", "方法学瓶颈归纳"],
        "results": "强调修饰测序需要同时处理物种组成、tRNA 同源性、修饰读出和群落差异。",
        "discussion": "作为工具路线图很有价值；但不提供新算法，需要转化为 benchmark 和标准化 workflow。",
        "visual": "pipeline",
    },
    {
        "pmid": "37802077",
        "status": "背景保留",
        "theme": "tDR 生物学综述",
        "logic": "tRNA fragmentation 被定义为受调控的再利用过程，为 tDR 分析提供生物问题边界。",
        "hypothesis": "tsRNA/tDR 不是随机降解片段，而是具有来源、结构和互作特征的功能分子。",
        "methods": ["整合 biogenesis", "结构与互作机制", "疾病/病毒相关功能", "AI 与结构预测方向"],
        "results": "提出未来需要全长序列、修饰、结构和互作的综合方法，支撑 tDR 功能解释。",
        "discussion": "适合作为问题背景；不宜作为工具性能证据，主线中应降权。",
        "visual": "PMC10841463",
    },
    {
        "pmid": "37173525",
        "status": "背景保留",
        "theme": "定量方法评论",
        "logic": "从方法评论角度强调 tRNA abundance 测序的主要技术壁垒和新方法价值。",
        "hypothesis": "tRNA 结构、修饰和序列相似性使常规 RNA-seq 不能直接替代 tRNA-seq。",
        "methods": ["比较建库偏差", "讨论修饰阻断", "解释定量瓶颈", "指出 benchmark 需求"],
        "results": "为 QuantM/mim 等方法提供背景定位，说明为什么定量 tRNA 需要专门工具。",
        "discussion": "没有新数据或算法，因此作为导读/过渡文献，不占主线证据页。",
        "visual": "pipeline",
    },
    {
        "pmid": "35322228",
        "status": "删出主线，保留附录",
        "theme": "治疗应用",
        "logic": "AAV-delivered suppressor tRNA 主要回答基因治疗可行性，而不是分析工具本身。",
        "hypothesis": "工程化 sup-tRNA 可在 premature stop codon 处恢复蛋白表达，同时不过度扰动全局终止密码子读穿。",
        "methods": ["sup-tRNA 设计", "AAV 递送", "ribosome profiling", "tRNA-seq 检查内源稳态"],
        "results": "小鼠模型中长期恢复功能，全球正常 stop codon readthrough 有限。",
        "discussion": "分析启发在于 tRNA-seq/ribo-seq 可作为安全性 readout；但不应作为工具算法主线。",
        "visual": "PMC9446716",
    },
    {
        "pmid": "37944512",
        "status": "删出主线，保留附录",
        "theme": "合成基因组",
        "logic": "tRNA neochromosome 是合成生物学工程，使用 tRNA-seq、多组学和 Hi-C 评估构建体功能。",
        "hypothesis": "将全部核编码 tRNA 转移到设计染色体上仍可维持酵母生长并用于扰动 tRNA gene organization。",
        "methods": ["neochromosome 设计", "SCRaMbLE", "tRNA-seq/transcriptomics/proteomics", "FISH/Hi-C/replication profiling"],
        "results": "构建体可存活但出现倍性变化等适应性压力，提示 tRNA 基因组织具有系统效应。",
        "discussion": "可作为 tRNA reference 和功能验证案例；但不是通用分析工具论文。",
        "visual": "pipeline",
    },
    {
        "pmid": "35654044",
        "status": "删出主线，保留附录",
        "theme": "tRF 功能机制",
        "logic": "研究 5'-tRFCys 促进转移的机制，数据分析侧重 tRF 差异表达和靶蛋白/代谢通路验证。",
        "hypothesis": "特定 tRF 可通过促进 Nucleolin oligomerization 稳定代谢 mRNA，从而增强癌细胞转移。",
        "methods": ["small RNA profiling", "差异 tRF 筛选", "RBP 互作鉴定", "动物/细胞功能验证"],
        "results": "5'-tRFCys 上调并驱动促转移代谢网络。",
        "discussion": "说明 tDR 差异分析需要可靠命名和来源解析；但本文贡献是生物机制而非工具。",
        "visual": "PMC9444141",
    },
    {
        "pmid": "41501459",
        "status": "删出主线，保留启发",
        "theme": "CRISPR-tRNA 识别机制",
        "logic": "Cas12a3 被靶 RNA 激活后切割 tRNA 3' CCA tail，证明 tRNA tail 可以成为可编程识别/报告位点。",
        "hypothesis": "部分 type V CRISPR effector 通过 target RNA 触发非靶 tRNA 裂解实现抗噬菌体防御。",
        "methods": ["cell-based/biochemical assay", "direct RNA sequencing", "cryo-EM", "synthetic tRNA-tail reporter"],
        "results": "发现 tRNA-loading domain 定位 tRNA tail 至 RuvC active site，并扩展 CRISPR RNA 检测能力。",
        "discussion": "对 tRNA tail 识别和纳米孔/direct RNA readout 有启发；但不是 tRNA 分析算法论文。",
        "visual": "PMC12851939",
    },
]


SECTION_SOURCES = {
    "all": [
        "34417604", "32796835", "33581077", "37024678", "39865096", "40835813",
        "40447571", "42062277", "39096899", "36869120", "39952700",
    ],
    "mapping": ["34417604", "32796835", "33581077", "36869120", "39952700"],
    "modification": ["33581077", "37024678", "39865096", "40447571", "40835813"],
}


def source_label(pmids):
    labels = []
    for pmid in pmids:
        p = paper(pmid)
        labels.append(f"PMID {pmid}: {p['title']}")
    return " | ".join(labels)


FIG_FILES = sorted((RES / "figures").glob("*.jpg"))


def fig_for_key(key):
    if key == "pipeline":
        return None
    candidates = [p for p in FIG_FILES if key in p.name]
    if not candidates:
        return None
    candidates = sorted(candidates, key=lambda p: ("Fig1" not in p.name, p.name))
    dest = ASSETS / candidates[0].name
    if not dest.exists():
        shutil.copy2(candidates[0], dest)
    return dest


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

COLORS = {
    "ink": RGBColor(31, 38, 46),
    "muted": RGBColor(92, 102, 112),
    "line": RGBColor(214, 219, 224),
    "blue": RGBColor(45, 93, 123),
    "green": RGBColor(88, 124, 93),
    "red": RGBColor(153, 73, 67),
    "gold": RGBColor(171, 129, 63),
    "bg": RGBColor(248, 249, 247),
    "panel": RGBColor(255, 255, 255),
    "pale_blue": RGBColor(231, 239, 244),
    "pale_green": RGBColor(232, 241, 234),
    "pale_gold": RGBColor(246, 240, 228),
    "pale_red": RGBColor(246, 233, 231),
}


def add_bg(slide):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS["bg"]
    shp.line.fill.background()


def tf_text(tf, text, font_size=16, bold=False, color="ink", align=None):
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = COLORS[color] if isinstance(color, str) else color
    if align:
        p.alignment = align
    return p


def add_text(slide, x, y, w, h, text, size=16, bold=False, color="ink", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    tf_text(box.text_frame, text, size, bold, color, align)
    return box


def add_title(slide, title, kicker=None):
    if kicker:
        add_text(slide, 0.55, 0.22, 8.5, 0.25, kicker, 8.5, False, "muted")
    add_text(slide, 0.55, 0.43 if kicker else 0.27, 11.3, 0.55, title, 23, True, "ink")
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.04), Inches(12.2), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()


def add_source(slide, text):
    if len(text) > 210:
        text = text[:207] + "..."
    add_text(slide, 0.55, 7.08, 12.25, 0.28, "Source: " + text, 7.3, False, "muted")


def add_takeaway(slide, text, y=6.48):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(y), Inches(12.2), Inches(0.42))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS["pale_blue"]
    shp.line.color.rgb = COLORS["line"]
    box = slide.shapes.add_textbox(Inches(0.74), Inches(y + 0.08), Inches(11.8), Inches(0.24))
    tf_text(box.text_frame, "Take-home: " + text, 10.5, True, "blue")


def add_bullets(slide, x, y, w, h, bullets, size=13, color="ink"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS[color]
        p.space_after = Pt(4)
    return box


def add_label(slide, x, y, text, color="blue", fill="pale_blue"):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.15), Inches(0.25))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[fill]
    shp.line.fill.background()
    box = slide.shapes.add_textbox(Inches(x + 0.06), Inches(y + 0.045), Inches(1.03), Inches(0.14))
    tf_text(box.text_frame, text, 7.2, True, color, PP_ALIGN.CENTER)


def add_card(slide, x, y, w, h, title, body, fill="panel", title_color="blue"):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[fill]
    shp.line.color.rgb = COLORS["line"]
    add_text(slide, x + 0.18, y + 0.14, w - 0.36, 0.26, title, 10.5, True, title_color)
    if isinstance(body, list):
        add_bullets(slide, x + 0.18, y + 0.48, w - 0.36, h - 0.55, body, 9.2)
    else:
        add_text(slide, x + 0.18, y + 0.48, w - 0.36, h - 0.55, body, 9.4)


def add_image(slide, img_path, x, y, w, h):
    if not img_path or not Path(img_path).exists():
        return False
    try:
        slide.shapes.add_picture(str(img_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        return True
    except Exception:
        return False


def add_pipeline(slide, x, y, w, steps, color="blue"):
    n = len(steps)
    gap = 0.12
    bw = (w - gap * (n - 1)) / n
    for i, step in enumerate(steps):
        bx = x + i * (bw + gap)
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(bw), Inches(0.55))
        shp.fill.solid()
        shp.fill.fore_color.rgb = COLORS["pale_blue"] if color == "blue" else COLORS["pale_green"]
        shp.line.color.rgb = COLORS[color]
        box = slide.shapes.add_textbox(Inches(bx + 0.06), Inches(y + 0.12), Inches(bw - 0.12), Inches(0.28))
        tf_text(box.text_frame, step, 8.4, True, color, PP_ALIGN.CENTER)
        if i < n - 1:
            add_text(slide, bx + bw + 0.015, y + 0.16, 0.09, 0.2, ">", 10, True, "muted")


def add_matrix(slide, x, y, rows, cols, data, col_widths=None, row_h=0.44, size=8.6):
    if col_widths is None:
        col_widths = [1 / cols] * cols
    total_w = sum(col_widths)
    for r in range(rows):
        cx = x
        for c in range(cols):
            w = col_widths[c]
            fill = "pale_blue" if r == 0 else ("panel" if r % 2 else "bg")
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(y + r * row_h), Inches(w), Inches(row_h))
            shp.fill.solid()
            shp.fill.fore_color.rgb = COLORS[fill]
            shp.line.color.rgb = COLORS["line"]
            txt = data[r][c] if c < len(data[r]) else ""
            box = slide.shapes.add_textbox(Inches(cx + 0.04), Inches(y + r * row_h + 0.055), Inches(w - 0.08), Inches(row_h - 0.08))
            tf_text(box.text_frame, txt, size if r else size + 0.2, r == 0, "ink")
            cx += w


def title_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_text(slide, 0.75, 0.95, 11.7, 1.05, "tRNA 分析工具与算法：\n从基因识别到修饰、结构、氨酰化和 tDR 注释", 26, True, "ink")
    add_text(slide, 0.78, 2.45, 10.8, 0.52, "正式汇报版 | 文献精读 + 工具架构 + 缺陷与改动方向", 15, False, "blue")
    add_pipeline(slide, 0.78, 3.35, 10.9, ["Gene detection", "Abundance", "Modification", "Structure", "Charging", "tDR annotation"], "blue")
    add_text(slide, 0.78, 4.45, 11.5, 0.7, "核心问题：tRNA 的短长度、高相似性、强二级结构和密集修饰，使常规 RNA-seq/小 RNA-seq/纳米孔流程都必须重新设计。", 16, True, "ink")
    add_text(slide, 0.78, 6.72, 11.5, 0.28, "来源：本地文献包 + 新增 tRNAscan-SE 2.0、tDRnamer/tRAX 方法源；每页底部标注 PMID 与标题。", 8.4, False, "muted")


def screening_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "文献重新筛选：把“tRNA 相关”收束为“tRNA 分析工具”")
    rows = [
        ["分类", "纳入主线", "作用"],
        ["核心方法/算法", "tRNAscan-SE 2.0, QuantM, mim, Nano, Induro, aa-tRNA-seq, DM-DMS-MaPseq, ORACLE, LIDAR, tDRnamer/tRAX", "用于底层架构、识别逻辑、差异分析和改动方向"],
        ["背景/综述", "tRNA modification dynamics; tRNA renovatio; Quantifying tRNA abundance by sequencing", "定义技术瓶颈和未来 benchmark 需求"],
        ["删出主线", "AAV suppressor tRNA; tRNA neochromosome; 5'-tRFCys metastasis; Cas12a3 tRNA-tail cleavage", "偏治疗、合成生物学或机制；保留精读卡片和方法启发"],
    ]
    add_matrix(slide, 0.65, 1.45, len(rows), 3, rows, [1.4, 7.0, 3.35], row_h=0.78, size=8.2)
    add_takeaway(slide, "最终主线以“识别对象 -> 建库/读出 -> 比对/分类 -> 差异分析 -> 可改动方向”为骨架。")
    add_source(slide, source_label(["34417604", "32796835", "33581077", "37024678", "39865096", "40447571", "40835813", "36869120"]))


def landscape_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "tRNA 分析不是一个任务，而是一组互相耦合的推断问题")
    add_card(slide, 0.65, 1.35, 2.15, 4.45, "对象层", ["tRNA gene / pseudogene", "mature tRNA / isodecoder", "modification site", "structure/interactome", "aminoacylation", "tDR/tsRNA"], "pale_blue")
    add_card(slide, 3.05, 1.35, 2.15, 4.45, "读出层", ["genome sequence", "short-read tRNA-seq", "RT mismatch/stop", "nanopore signal", "chemical probing", "small RNA-seq"], "pale_green", "green")
    add_card(slide, 5.45, 1.35, 2.15, 4.45, "算法层", ["covariance model", "multi-mapping", "error/signature matrix", "ML signal classifier", "differential model", "ontology naming"], "pale_gold", "gold")
    add_card(slide, 7.85, 1.35, 2.15, 4.45, "统计层", ["spike-in/UMI", "negative binomial", "beta-binomial", "mixed model", "batch correction", "uncertainty propagation"], "panel")
    add_card(slide, 10.25, 1.35, 2.15, 4.45, "解释层", ["codon demand", "stress response", "modification crosstalk", "translation efficiency", "tDR function", "clinical/engineering use"], "pale_red", "red")
    add_takeaway(slide, "工具缺陷通常不是单点问题，而是从建库偏差一路传递到统计和生物解释。")
    add_source(slide, source_label(SECTION_SOURCES["all"]))


def pain_points_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "tRNA 工具的共同底层难点")
    add_card(slide, 0.75, 1.3, 3.75, 1.35, "短且高度相似", "isodecoder/isotype 之间差异小，多重比对导致 gene-level abundance 不稳定。", "panel")
    add_card(slide, 4.8, 1.3, 3.75, 1.35, "修饰密集", "RT stop、misincorporation 和 nanopore signal shift 既是偏差，也是可利用的信号。", "panel")
    add_card(slide, 8.85, 1.3, 3.75, 1.35, "结构强", "二级结构影响接头连接、RT 穿越、化学探针可及性和读段完整性。", "panel")
    add_card(slide, 0.75, 3.1, 3.75, 1.35, "末端状态复杂", "CCA、氨酰化、blocked 3' termini 和 tDR 切割位点改变建库可见性。", "panel")
    add_card(slide, 4.8, 3.1, 3.75, 1.35, "参考体系不唯一", "基因组 tRNA、成熟 tRNA、线粒体 tRNA、pseudogene 和 tDR ontology 需要统一坐标。", "panel")
    add_card(slide, 8.85, 3.1, 3.75, 1.35, "差异分析层级多", "abundance、modification fraction、structure reactivity、charging ratio 与 fragment class 不是同一种统计量。", "panel")
    add_takeaway(slide, "一个稳健 pipeline 必须把“无法唯一判断”的不确定性保留下来，而不是过早硬分配。")
    add_source(slide, source_label(["32796835", "33581077", "37024678", "39865096", "36869120", "39952700"]))


def paper_card_slide(card, idx):
    p = paper(card["pmid"])
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, f"精读 {idx:02d}｜{card['theme']}：{p['title'][:62]}", f"{card['status']} | {p.get('journal','')} {p.get('year','')} | PMID {card['pmid']}")
    add_label(slide, 0.68, 1.22, card["status"], "blue" if "主线" in card["status"] else "red", "pale_blue" if "主线" in card["status"] else "pale_red")
    img = fig_for_key(card["visual"])
    image_ok = add_image(slide, img, 0.68, 1.58, 5.15, 3.05)
    if not image_ok:
        steps = card["methods"][:4]
        add_pipeline(slide, 0.82, 2.3, 4.65, steps, "blue")
        add_text(slide, 1.05, 3.1, 4.25, 0.9, "方法示意图（根据文献流程整理）", 15, True, "muted", PP_ALIGN.CENTER)
    add_card(slide, 6.08, 1.25, 3.05, 1.3, "文章逻辑", card["logic"], "panel")
    add_card(slide, 9.35, 1.25, 3.05, 1.3, "核心假说", card["hypothesis"], "panel")
    add_card(slide, 6.08, 2.82, 3.05, 1.55, "方法步骤", card["methods"], "panel")
    add_card(slide, 9.35, 2.82, 3.05, 1.55, "结果与讨论", [card["results"], card["discussion"]], "panel")
    add_text(slide, 0.7, 4.82, 5.15, 0.58, "汇报定位：" + card["theme"], 13.5, True, "blue")
    add_bullets(slide, 0.86, 5.32, 11.4, 0.78, [
        "识别逻辑：" + card["logic"],
        "可改动方向：" + card["discussion"],
    ], 9.4, "ink")
    add_takeaway(slide, card["theme"] + " 的关键不是单一软件命令，而是把生物化学偏差转化为可统计建模的信号。")
    add_source(slide, f"PMID {card['pmid']}: {p['title']}")


def architecture_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "通用底层架构：reference-first 与 signal-first 两条路线应汇合")
    add_text(slide, 0.7, 1.28, 2.7, 0.35, "Reference-first", 15, True, "blue")
    add_pipeline(slide, 0.7, 1.78, 5.9, ["Genome", "tRNAscan-SE", "mature ref", "alignment", "abundance/tDR"], "blue")
    add_text(slide, 0.7, 2.75, 5.8, 0.7, "适合 tRNA gene catalog、短读长定量和 tDR 注释；主要风险是多拷贝、多等解码 tRNA 和 pseudogene 的归属。", 12.5)
    add_text(slide, 0.7, 4.0, 2.7, 0.35, "Signal-first", 15, True, "green")
    add_pipeline(slide, 0.7, 4.5, 5.9, ["Raw reads", "RT/signature", "event signal", "classifier", "mod/charging"], "green")
    add_text(slide, 0.7, 5.47, 5.8, 0.7, "适合修饰、结构和氨酰化状态；主要风险是训练集依赖、酶/平台批次和信号可解释性。", 12.5)
    add_card(slide, 7.0, 1.45, 5.35, 4.55, "推荐的统一层", [
        "同一份 mature+tDR reference graph：保留 gene/isodecoder/isotype 层级。",
        "读段级不确定性：multi-mapping 不硬分配，传递到 differential model。",
        "信号级证据矩阵：RT-stop、mismatch、DMS reactivity、nanopore event 作为独立 evidence channels。",
        "输出层同时给 counts、fraction、confidence 和 ontology-compatible names。"
    ], "pale_gold", "gold")
    add_takeaway(slide, "未来工具应从“单方法 pipeline”走向“多证据联合推断框架”。")
    add_source(slide, source_label(["34417604", "33581077", "37024678", "39865096", "40447571", "40835813", "36869120"]))


def deep_dive_slides():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖一：tRNAscan-SE 2.0 的基因识别架构")
    add_pipeline(slide, 0.72, 1.42, 11.6, ["Genome windows", "fast candidate scan", "CM validation", "isotype model", "pseudogene filter", "functional class"], "blue")
    add_card(slide, 0.78, 2.48, 3.6, 2.7, "底层算法", ["协方差模型同时编码序列保守性和二级结构。", "先粗筛候选，再用结构模型精筛，降低全基因组搜索成本。", "anticodon、identity element 和结构完整性共同支持功能分类。"], "panel")
    add_card(slide, 4.88, 2.48, 3.6, 2.7, "识别逻辑", ["真 tRNA：结构完整、score 高、anticodon/CCA 等特征一致。", "伪基因：局部像 tRNA，但结构破损或身份元件异常。", "线粒体 tRNA：结构更非典型，需要专门模型。"], "panel")
    add_card(slide, 8.98, 2.48, 3.6, 2.7, "改动方向", ["用样本 tRNA-seq 证据辅助过滤 pseudogene。", "输出可直接用于 mature tRNA/tDR 的 reference graph。", "把保守性、表达证据和多拷贝不确定性一起报告。"], "panel")
    add_takeaway(slide, "tRNAscan-SE 解决的是“可能是哪一类 tRNA 基因”，不是“在样本中表达了多少”。")
    add_source(slide, source_label(["34417604"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖二：短读长 tRNA-seq 的定量算法")
    add_pipeline(slide, 0.72, 1.42, 11.6, ["adapter/CCA handling", "RT stop-aware reads", "mature tRNA reference", "multi-map assignment", "normalization", "DE testing"], "blue")
    add_card(slide, 0.78, 2.48, 3.6, 2.7, "QuantM 的核心", ["优化成熟 tRNA 建库并评估组织间 tRNA 丰度。", "从 reads 层追踪 RT fall-off 和修饰相关偏差。", "比较 gene/isodecoder/anticodon 不同层级的表达。"], "panel")
    add_card(slide, 4.88, 2.48, 3.6, 2.7, "mim 的核心", ["把修饰诱导错配作为 site-level feature。", "同时估计 abundance 与 modification status。", "用分析工具包减少手工注释和错配统计。"], "panel")
    add_card(slide, 8.98, 2.48, 3.6, 2.7, "差异分析要点", ["abundance 用 NB/GLM，但必须处理多重比对。", "modification fraction 用比例模型，而非原始 read counts。", "最好保留 isotype、isodecoder、gene 三层输出。"], "panel")
    add_takeaway(slide, "短读长工具的关键是不要把 RT 停顿和多重比对造成的偏差误当作生物差异。")
    add_source(slide, source_label(["32796835", "33581077", "37173525"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖三：mim 与 Induro 如何把 RT 异常变成修饰信号")
    add_pipeline(slide, 0.72, 1.42, 11.6, ["modified tRNA", "RT enzyme behavior", "stop/readthrough", "misincorporation matrix", "site signature", "modification inference"], "green")
    add_card(slide, 0.78, 2.48, 3.6, 2.9, "mim-tRNAseq", ["观察 RT 过程中稳定出现的 misincorporation。", "适合细胞系/物种间的相对修饰状态比较。", "依赖已知修饰指纹和高质量参考。"], "panel")
    add_card(slide, 4.88, 2.48, 3.6, 2.9, "Induro-tRNAseq", ["使用 readthrough 更强的 Induro RT。", "通过时间梯度与对照 RT 比较区分停顿和错配。", "适合从全基因组角度看 coordinated changes。"], "panel")
    add_card(slide, 8.98, 2.48, 3.6, 2.9, "联合 caller", ["把 stop、mismatch、coverage drop 分开建模。", "用 knockout/标准品训练 modification probability。", "输出低置信位点，避免过度解释。"], "panel")
    add_takeaway(slide, "RT 异常既是技术障碍，也是 tRNA 修饰测序最重要的可观测变量。")
    add_source(slide, source_label(["33581077", "39865096"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖四：纳米孔 tRNA 工具的 raw-signal 路线")
    add_pipeline(slide, 0.72, 1.42, 11.6, ["native/ligated tRNA", "nanopore current", "read recovery", "event features", "mapping", "mod/charging inference"], "green")
    add_card(slide, 0.78, 2.48, 3.6, 2.9, "Nano-tRNAseq", ["核心改动不是单纯测序，而是重处理 raw signal。", "恢复默认流程会丢掉的短 tRNA reads。", "同时观察 abundance 与 modification dynamics。"], "panel")
    add_card(slide, 4.88, 2.48, 3.6, 2.9, "aa-tRNA-seq", ["化学连接把 amino acid identity 转为可读信号扰动。", "用机器学习识别单分子携带的氨基酸。", "适合 charging fidelity 和 hypomodification 后果。"], "panel")
    add_card(slide, 8.98, 2.48, 3.6, 2.9, "算法瓶颈", ["短 RNA basecalling 训练域不足。", "修饰与氨酰化信号可能互相混杂。", "需要 read-level calibration、置信区间和独立验证。"], "panel")
    add_takeaway(slide, "纳米孔路线的优势在单分子整合读出；风险在于 signal model 的可解释性和训练集泛化。")
    add_source(slide, source_label(["37024678", "40835813"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖五：结构/互作分析不能只看二级结构预测")
    add_pipeline(slide, 0.72, 1.42, 11.6, ["in vivo DMS", "demethylase treatment", "MaP mutations", "reactivity profile", "stress comparison", "interaction inference"], "blue")
    add_card(slide, 0.78, 2.48, 3.6, 2.9, "DM-DMS-MaPseq 的价值", ["在人细胞内同时观察胞质和线粒体 tRNA。", "DMS 反应性反映局部可及性。", "体内与体外差异揭示蛋白/核糖体保护。"], "panel")
    add_card(slide, 4.88, 2.48, 3.6, 2.9, "差异分析逻辑", ["position-level reactivity 是主要单位。", "stress/control 需配对比较。", "必须区分结构改变和互作遮蔽改变。"], "panel")
    add_card(slide, 8.98, 2.48, 3.6, 2.9, "改动方向", ["加入 Ribo-seq、RIP/CLIP 或蛋白扰动做解卷积。", "建立 tRNA family-aware reactivity baseline。", "把结构变化与 abundance/modification 联合建模。"], "panel")
    add_takeaway(slide, "tRNA structurome 的读出本质是“结构 + 互作 + 条件”的合成信号。")
    add_source(slide, source_label(["40447571"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具深挖六：低输入与 tDR 分析代表两个极端场景")
    add_card(slide, 0.8, 1.32, 5.45, 3.95, "ORACLE-tRNAseq：极低输入", ["样本量小，核心挑战是建库效率、dropout 和批次效应。", "需要把 tRNA repertoire 与 ZGA、H3K4me3、Ribo-seq 对齐。", "差异分析应优先稳健性和时间序列一致性，而非单点显著性。"], "panel")
    add_card(slide, 6.9, 1.32, 5.45, 3.95, "LIDAR/tRAX/tDRnamer：末端与来源复杂", ["blocked 3' termini 会让传统 ligation-based 方法漏检 tDR。", "small RNA reads 很短，来源 tRNA 经常多解。", "标准命名必须绑定 origin uncertainty，不能只给漂亮名字。"], "panel")
    add_pipeline(slide, 1.0, 5.66, 11.0, ["low-input capture", "or", "blocked-end capture", "reference-aware mapping", "standard ontology", "differential interpretation"], "green")
    add_takeaway(slide, "低输入和 tDR 场景提醒我们：建库可见性与来源归属会决定后续所有统计结论。")
    add_source(slide, source_label(["42062277", "39096899", "36869120", "39952700"]))


def mapping_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "识别逻辑：从 tRNA gene 到 tDR origin 不能只靠最佳比对")
    rows = [
        ["层级", "主要输入", "核心判据", "常见误差", "建议改动"],
        ["gene", "genome sequence", "二级结构 + CM score + anticodon", "pseudogene/tRNA-like 元件", "加入表达证据和保守性"],
        ["mature tRNA", "tRNA-seq reads", "CCA、intron removal、isodecoder", "多重比对、RT stop", "概率分配 + UMI/spike-in"],
        ["modification", "mismatch/stop/signal", "位点模式与对照差异", "未知修饰混同", "联合酶/标准品/knockout evidence"],
        ["tDR", "small RNA reads", "fragment start/end + source tRNA", "同源 tRNA 来源不唯一", "tDRnamer + uncertainty-aware counts"],
    ]
    add_matrix(slide, 0.55, 1.35, len(rows), 5, rows, [1.0, 2.15, 3.05, 2.65, 3.25], row_h=0.64, size=8.1)
    add_takeaway(slide, "tRNA 工具的“识别”本质上是层级归因，最佳实践是输出层级置信度而非单一标签。")
    add_source(slide, source_label(SECTION_SOURCES["mapping"]))


def differential_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "差异分析：不同读出需要不同统计模型")
    rows = [
        ["读出", "推荐比较单位", "统计模型", "必须控制"],
        ["abundance", "gene/isodecoder/isotype/anticodon", "negative binomial 或 compositional model", "library size、spike-in、multi-mapping"],
        ["modification fraction", "site × tRNA family", "beta-binomial/logistic mixed model", "coverage、RT enzyme、batch、neighbor sequence"],
        ["nanopore signal", "event/site-level feature", "supervised ML + calibration curve", "basecaller、motor、adapter、training domain"],
        ["DMS reactivity", "position-level reactivity", "paired differential reactivity model", "in vitro baseline、protein protection、structure context"],
        ["charging", "amino acid identity/ratio", "classifier probability + hierarchical model", "chemical ligation efficiency、class imbalance"],
        ["tDR", "fragment class/name", "small RNA DE + origin uncertainty", "length bias、3' block、multi-origin reads"],
    ]
    add_matrix(slide, 0.5, 1.22, len(rows), 4, rows, [2.0, 3.25, 3.35, 3.35], row_h=0.58, size=7.8)
    add_takeaway(slide, "把所有信号都塞进同一个 counts 表会掩盖机制；应按读出类型分层建模，再综合解释。")
    add_source(slide, source_label(["32796835", "33581077", "37024678", "39865096", "40447571", "40835813", "39096899"]))


def comparison_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "主线工具横向比较")
    rows = [
        ["工具/方法", "主要对象", "读出优势", "关键算法", "主要缺陷"],
        ["tRNAscan-SE 2.0", "tRNA genes", "跨基因组 gene catalog", "CM/结构模型 + 分类", "不解决表达/修饰"],
        ["QuantM", "mature tRNA abundance", "组织/样本定量", "专用参考比对", "短读长与 RT bias"],
        ["mim-tRNAseq", "abundance + modification status", "错配可转化为修饰指纹", "misincorporation profiling", "未知修饰归因难"],
        ["Nano-tRNAseq", "native tRNA", "单分子、全长、修饰联动", "raw signal reprocessing", "basecaller/短 RNA 偏差"],
        ["Induro-tRNAseq", "genome-wide modifications", "降低 RT stop 保留错配", "enzyme-contrast signature", "需要验证集"],
        ["DM-DMS-MaPseq", "structure/interactome", "体内结构与互作", "DMS-MaP reactivity", "结构/互作混杂"],
        ["aa-tRNA-seq", "aminoacylation", "单分子 amino acid identity", "nanopore ML classifier", "训练泛化与化学效率"],
        ["LIDAR/tDRnamer/tRAX", "tDR/tsRNA", "blocked-end capture + 标准命名", "origin naming/visualization", "来源不确定性"],
    ]
    add_matrix(slide, 0.42, 1.08, len(rows), 5, rows, [1.75, 2.15, 2.65, 2.7, 2.45], row_h=0.55, size=7.1)
    add_takeaway(slide, "没有一个工具覆盖所有层级；汇报和项目设计应按研究问题组合工具，而不是寻找唯一最佳工具。")
    add_source(slide, source_label(SECTION_SOURCES["all"]))


def defects_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "工具缺陷与可改动方向")
    rows = [
        ["模块", "现有缺陷", "可落地改动方向"],
        ["参考构建", "基因组 tRNA、成熟 tRNA、线粒体 tRNA 和 tDR 坐标割裂", "构建 mature/reference graph；保留 gene-isodecoder-isotype 层级"],
        ["比对计数", "重复序列和短 reads 导致硬分配偏差", "EM/Bayesian multi-mapping；UMI 去重；spike-in 绝对定量"],
        ["修饰识别", "RT stop、mismatch 和 nanopore signal 分散在不同工具", "联合 evidence matrix；用 knockout/标准品训练 site-level caller"],
        ["差异分析", "常把 abundance、modification、structure 统一当 counts", "按读出建模：NB、beta-binomial、mixed model、classifier probability"],
        ["tDR 注释", "fragment name 标准化不足，功能分析跨研究不可比", "强制输出 tDRnamer-compatible name 和 origin uncertainty"],
        ["软件工程", "流程依赖手工参数和私有 reference", "Nextflow/Snakemake + container + benchmark dataset + report template"],
    ]
    add_matrix(slide, 0.52, 1.15, len(rows), 3, rows, [1.65, 5.05, 5.45], row_h=0.72, size=8.4)
    add_takeaway(slide, "最值得做的改动不是重写一个新工具，而是把参考、比对、修饰证据和统计层统一成可复现工作流。")
    add_source(slide, source_label(SECTION_SOURCES["all"]))


def improvement_slides():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "建议方向一：统一 reference graph 与不确定性传递")
    add_pipeline(slide, 0.8, 1.55, 11.6, ["tRNAscan-SE gene catalog", "mature tRNA transcriptome", "CCA/intron/edit rules", "tDR coordinate system", "probabilistic read assignment"], "blue")
    add_card(slide, 0.9, 2.75, 3.5, 2.6, "为什么需要", ["tRNA 多拷贝导致 gene-level 唯一定位经常不成立。", "不同工具的 reference 坐标不统一，结果难以合并。", "硬分配会把上游不确定性伪装成下游显著性。"], "panel")
    add_card(slide, 4.85, 2.75, 3.5, 2.6, "实现方式", ["用 tRNAscan-SE 构建候选 gene 层。", "生成 mature transcript 与 tDR fragment graph。", "用 EM/Bayesian 模型输出层级 counts 和 posterior。"], "panel")
    add_card(slide, 8.8, 2.75, 3.5, 2.6, "预期收益", ["跨样本差异更稳健。", "tDR 命名和 abundance 可追溯。", "更容易接入 spike-in、UMI 和多组学证据。"], "panel")
    add_takeaway(slide, "先统一“对象是什么”，再讨论“差异是否显著”。")
    add_source(slide, source_label(["34417604", "32796835", "36869120", "39952700"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "建议方向二：修饰识别应从单一 signature 走向联合 caller")
    add_pipeline(slide, 0.8, 1.55, 11.6, ["RT-stop", "Mismatch", "Induro/readthrough", "Nanopore signal", "DMS reactivity", "Joint modification caller"], "green")
    add_card(slide, 0.9, 2.75, 3.55, 2.72, "输入证据", ["mim/Induro 的错配矩阵。", "Nano-tRNAseq 的 event/signal 偏移。", "DM-DMS 的反应性变化。", "修饰酶敲除或标准品作为标签。"], "panel")
    add_card(slide, 4.85, 2.75, 3.55, 2.72, "建模策略", ["site × tRNA × sample 三维张量。", "beta-binomial 处理比例信号。", "ML/校准曲线处理纳米孔特征。", "输出 modification probability 而非硬标签。"], "panel")
    add_card(slide, 8.8, 2.75, 3.55, 2.72, "风险控制", ["区分覆盖不足与真实缺失。", "控制 RT 酶、批次和相邻序列效应。", "跨平台验证高置信位点。"], "panel")
    add_takeaway(slide, "修饰 caller 的目标应是可校准概率，而不是把所有异常读段归因给某一种修饰。")
    add_source(slide, source_label(["33581077", "37024678", "39865096", "40447571"]))

    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "建议方向三：面向汇报和项目复用的标准 pipeline")
    add_card(slide, 0.8, 1.35, 2.65, 3.9, "输入标准", ["FASTQ/FAST5/POD5", "sample sheet", "species/reference version", "spike-in/UMI 设置", "method-specific metadata"], "pale_blue")
    add_card(slide, 3.8, 1.35, 2.65, 3.9, "流程标准", ["QC + trimming", "reference graph", "read assignment", "signal/signature extraction", "differential module", "HTML/PPT report"], "pale_green", "green")
    add_card(slide, 6.8, 1.35, 2.65, 3.9, "输出标准", ["layered count table", "modification probability", "structure reactivity", "charging ratio", "tDRnamer names", "uncertainty columns"], "pale_gold", "gold")
    add_card(slide, 9.8, 1.35, 2.65, 3.9, "验证标准", ["synthetic controls", "enzyme knockouts", "cross-platform concordance", "read-level audit plots", "benchmark dataset"], "pale_red", "red")
    add_takeaway(slide, "真正可推广的工具，需要把算法结果、来源证据和不确定性一起交付。")
    add_source(slide, source_label(SECTION_SOURCES["all"]))


def closing_slide():
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_title(slide, "结论：tRNA 分析工具的下一步是“多证据、层级化、可复现”")
    add_card(slide, 0.95, 1.45, 3.55, 3.85, "1. 分析对象要分层", ["gene、mature tRNA、modification、structure、charging、tDR 不能混为同一层。", "每个层级都需要明确 reference 与 uncertainty。"], "panel")
    add_card(slide, 4.9, 1.45, 3.55, 3.85, "2. 偏差也是信号", ["RT-stop、mismatch、纳米孔电流扰动和 DMS reactivity 是方法学核心。", "关键是校准，而不是简单过滤。"], "panel")
    add_card(slide, 8.85, 1.45, 3.55, 3.85, "3. 差异分析要匹配读出", ["counts、fractions、reactivity、classifier probabilities 应使用不同模型。", "跨方法综合解释应保留不确定性。"], "panel")
    add_takeaway(slide, "以 tRNA 为对象的计算工具，核心竞争力会从“能跑通”转向“能解释、能校准、能复用”。")
    add_source(slide, source_label(SECTION_SOURCES["all"]))


title_slide()
screening_slide()
landscape_slide()
pain_points_slide()
for i, card in enumerate(PAPER_CARDS, 1):
    paper_card_slide(card, i)
deep_dive_slides()
architecture_slide()
mapping_slide()
differential_slide()
comparison_slide()
defects_slide()
improvement_slides()
closing_slide()


def set_doc_props(prs):
    props = prs.core_properties
    props.title = "tRNA 分析工具与算法汇报"
    props.subject = "tRNA analysis tools and algorithms"
    props.author = "Codex"
    props.comments = "Generated from local literature package and selected added method sources."


set_doc_props(prs)
pptx_path = OUT / "tRNA_analysis_tools_algorithms_report_cn.pptx"
prs.save(pptx_path)

notes_path = OUT / "paper_reading_matrix.md"
with open(notes_path, "w", encoding="utf-8-sig") as f:
    f.write("# tRNA 分析工具文献精读矩阵\n\n")
    for i, c in enumerate(PAPER_CARDS, 1):
        p = paper(c["pmid"])
        f.write(f"## {i}. PMID {c['pmid']} | {p['title']}\n")
        f.write(f"- Status: {c['status']}\n")
        f.write(f"- Theme: {c['theme']}\n")
        f.write(f"- Article logic: {c['logic']}\n")
        f.write(f"- Hypothesis: {c['hypothesis']}\n")
        f.write(f"- Methods: {'; '.join(c['methods'])}\n")
        f.write(f"- Results: {c['results']}\n")
        f.write(f"- Discussion/Direction: {c['discussion']}\n\n")

manifest_path = OUT / "asset_manifest.md"
with open(manifest_path, "w", encoding="utf-8-sig") as f:
    f.write("# Figure Asset Manifest\n\n")
    for img in sorted(ASSETS.glob("*")):
        f.write(f"- {img.name}: copied from local downloaded_resources/figures and used as source visual where matching PMID/PMCID was available.\n")

slide_count = len(prs.slides)
media_count = 0
with zipfile.ZipFile(pptx_path, "r") as zf:
    media_count = len([n for n in zf.namelist() if n.startswith("ppt/media/")])
    has_presentation = "ppt/presentation.xml" in zf.namelist()

qa_path = OUT / "qa_report.md"
with open(qa_path, "w", encoding="utf-8-sig") as f:
    f.write("# QA Report\n\n")
    f.write(f"- PPTX: `{pptx_path}`\n")
    f.write(f"- Creation status: success\n")
    f.write(f"- Slide count: {slide_count}\n")
    f.write(f"- Embedded media files: {media_count}\n")
    f.write(f"- Package validation: {'ppt/presentation.xml present' if has_presentation else 'presentation.xml missing'}\n")
    f.write("- Source labels: every slide includes a PMID/title label or integrated source list.\n")
    f.write("- Visual policy: source figures were inserted when local assets were available; otherwise slides use native PPT workflow schematics.\n")
    f.write("- Known limitations: some local PMC HTML files were reCAPTCHA/abstract-only, so those cards rely on local PubMed metadata, abstracts, available PDFs, and method-level synthesis rather than complete local full text.\n")

print(pptx_path)
print(qa_path)
print(notes_path)
print(slide_count, media_count)
