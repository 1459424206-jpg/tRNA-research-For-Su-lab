# -*- coding: utf-8 -*-
from pathlib import Path
import textwrap
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor


BASE = Path(r"D:\Project\tRNA research\tRNA mod-fingerprint database\downloaded_resources")
FIG = BASE / "figures"
FIG_ALL = BASE / "figures_all"
OUT = BASE / "analysis_outputs" / "paper2ppt_output" / "formal_report_deck"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

PPTX = OUT / "tRNA_modification_fingerprint_database_formal_report.pptx"
QA = OUT / "qa_report.md"
MANIFEST = OUT / "asset_manifest.md"


def font(size, bold=False):
    for name in (["msyhbd.ttc", "simhei.ttf", "arial.ttf"] if bold else ["msyh.ttc", "simhei.ttf", "arial.ttf"]):
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            pass
    return ImageFont.load_default()


F_TITLE = font(34, True)
F_HEAD = font(25, True)
F_TEXT = font(21)
F_SMALL = font(16)


def wrap_text(text, width):
    lines = []
    for seg in str(text).split("\n"):
        if len(seg) <= width:
            lines.append(seg)
        else:
            lines.extend(textwrap.wrap(seg, width=width, break_long_words=True, replace_whitespace=False))
    return lines


def make_flow(path, title, steps, footer=None):
    w, h = 1600, 900
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w - 1, h - 1], outline="black", width=3)
    d.text((55, 35), title, fill="black", font=F_TITLE)
    top = 145
    box_h = 90
    gap = 34
    x = 100
    bw = w - 200
    for i, step in enumerate(steps):
        y = top + i * (box_h + gap)
        d.rounded_rectangle([x, y, x + bw, y + box_h], radius=7, outline="black", width=3, fill="white")
        yy = y + 14
        for line in wrap_text(step, 52)[:2]:
            d.text((x + 26, yy), line, fill="black", font=F_TEXT)
            yy += 29
        if i < len(steps) - 1:
            cx = w // 2
            d.line([cx, y + box_h + 7, cx, y + box_h + gap - 14], fill="black", width=4)
            d.polygon([(cx - 10, y + box_h + gap - 14), (cx + 10, y + box_h + gap - 14), (cx, y + box_h + gap + 1)], fill="black")
    if footer:
        d.text((55, h - 58), footer, fill="black", font=F_SMALL)
    im.save(path, quality=95)
    return path


def make_matrix(path, title, headers, rows, col_widths=None, foot=None):
    w, h = 1800, 1000
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w - 1, h - 1], outline="black", width=3)
    d.text((45, 30), title, fill="black", font=F_TITLE)
    x0, y0 = 45, 118
    table_w = w - 90
    if col_widths is None:
        col_widths = [1] * len(headers)
    total = sum(col_widths)
    widths = [int(table_w * c / total) for c in col_widths]
    xs = [x0]
    for cw in widths:
        xs.append(xs[-1] + cw)
    row_h = int((h - y0 - 85) / (len(rows) + 1))
    for j, head in enumerate(headers):
        d.rectangle([xs[j], y0, xs[j + 1], y0 + row_h], outline="black", width=2, fill="white")
        d.text((xs[j] + 10, y0 + 14), head, fill="black", font=F_HEAD)
    for i, row in enumerate(rows):
        y = y0 + (i + 1) * row_h
        for j, cell in enumerate(row):
            d.rectangle([xs[j], y, xs[j + 1], y + row_h], outline="black", width=2, fill="white")
            yy = y + 9
            width_chars = max(12, int(widths[j] / 25))
            for line in wrap_text(cell, width_chars)[:5]:
                d.text((xs[j] + 10, yy), line, fill="black", font=F_SMALL)
                yy += 22
    if foot:
        d.text((45, h - 55), foot, fill="black", font=F_SMALL)
    im.save(path, quality=95)
    return path


def make_position_map(path):
    w, h = 1600, 900
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w - 1, h - 1], outline="black", width=3)
    d.text((55, 35), "tRNA修饰位置与主要功能指向", fill="black", font=F_TITLE)
    cx, cy = 650, 470
    d.line([cx, cy, cx, 220], fill="black", width=5)
    d.line([cx, cy, cx, 725], fill="black", width=5)
    d.arc([cx - 210, cy - 160, cx - 40, cy + 10], 90, 360, fill="black", width=5)
    d.arc([cx + 40, cy - 160, cx + 210, cy + 10], 180, 450, fill="black", width=5)
    d.arc([cx - 125, cy + 40, cx + 125, cy + 290], 180, 360, fill="black", width=5)
    labels = [
        ("U34 wobble", cx - 210, cy + 250, "mcm5s2U, Q, I, k2C\n解码效率/特异性"),
        ("A37邻位", cx - 15, cy + 260, "t6A, i6A, yW\nreading frame稳定"),
        ("C32/38/39", cx + 230, cy + 65, "m3C, Ψ\n结构选择/翻译精度"),
        ("G46/T-loop", cx + 230, cy - 105, "m7G46, T, Ψ\n稳定性/RTD"),
        ("D-loop", cx - 330, cy - 105, "D, m1A, m2G\n折叠与识别"),
        ("3' CCA", cx + 35, cy - 280, "氨酰化\nCCA完整性"),
    ]
    for head, x, y, body in labels:
        d.rectangle([x, y, x + 310, y + 102], outline="black", width=2, fill="white")
        d.text((x + 10, y + 8), head, fill="black", font=F_HEAD)
        yy = y + 43
        for line in body.split("\n"):
            d.text((x + 10, yy), line, fill="black", font=F_SMALL)
            yy += 23
    d.text((55, h - 62), "同一修饰的检测信号和表型必须绑定位置、tRNA家族和结构区域。", fill="black", font=F_TEXT)
    im.save(path, quality=95)
    return path


VISUALS = {
    "title": make_flow(ASSETS / "00_title.png", "tRNA修饰指纹库构建框架", [
        "修饰对象：tRNA上的化学修饰、位点、结构区域和动态状态",
        "检测信号：质谱、逆转录错配/停顿、化学切割、纳米孔错误/终止",
        "工具实现：先验参考、修饰字典、经验错误谱、化学规则和统计模型",
        "数据库目标：连接修饰、方法、识别率/通读率、功能表型和来源证据",
    ]),
    "definition": make_matrix(ASSETS / "01_definition.png", "指纹库在工具中的作用方式", ["作用方式", "是否是先验序列列表", "典型工具", "数据库含义"], [
        ["参考序列库", "是。提供tRNA序列、mature CCA、基因组坐标", "SPORTS1.1、tRNA reference construction、Galaxy workflow", "解决reads映射和位点坐标"],
        ["修饰字典/先验位点表", "部分是。来自MODOMICS或已知修饰位点", "MODOMICS_Decoder、tRNA004、mim-tRNAseq", "提供修饰名、符号、位置和真值锚点"],
        ["经验信号库", "不是单纯序列列表。记录修饰引起的错误谱/终止谱", "tRNA004、Nanopore 2-Read", "用于比较识别率、通读率和上下文差异"],
        ["化学规则库", "不是。由反应规则定义可检测位点", "TRAC-Seq", "m7G经NaBH4/aniline产生切割信号"],
        ["功能证据库", "不是。连接扰动、翻译和表型", "Ribo-seq、氨酰化、感染/衰老模型", "解释修饰为什么重要"],
    ], [1.2, 1.7, 1.6, 2.2]),
    "counts": make_matrix(ASSETS / "02_counts.png", "修饰数量与覆盖范围", ["层级", "数量/范围", "汇报口径"], [
        ["RNA修饰总库", "MODOMICS 2023 update：所有RNA类别已记录 >170 种修饰", "作为全量命名和化学字典"],
        ["tRNA修饰子集", "文献常用表述：tRNA含有 >100 种修饰，单条tRNA可含多个修饰位点", "作为数据库长期覆盖目标"],
        ["实验可识别集合", "tRNA004以43种已知tRNA修饰作为纳米孔benchmark", "作为工具性能评估集合"],
        ["第一版优先集合", "m7G、m1A、m3C、Ψ、I/Q、mcm5s2U、mcm5h2U、ac4C、k2C、t6A/i6A/yW", "优先覆盖强证据和强表型修饰"],
    ], [1.4, 2.7, 2.4]),
    "families": make_matrix(ASSETS / "03_families.png", "主要修饰家族与可观测信号", ["家族", "代表修饰", "典型信号", "表型方向"], [
        ["甲基化", "m1A, m3C, m5C, m7G, Gm/Cm/Um", "RT错配/停顿、TRAC切割、MS质量差异", "tRNA稳定、翻译效率、应激"],
        ["异构/还原", "Ψ, D", "MS或特异测序；普通错配信号较弱", "结构稳定、折叠、疾病相关"],
        ["U34复杂修饰", "mcm5U, mcm5s2U, Q, I, k2C", "codon识别差异、纳米孔错误谱", "解码效率和特异性"],
        ["A37超修饰", "t6A, i6A, ms2i6A, yW", "结构识别、MS、翻译精度", "reading frame稳定"],
        ["动态化学转换", "mcm5s2U→mcm5h2U", "氧化条件下MS/功能变化", "应激下翻译调节"],
        ["乙酰化", "ac4C", "ac4C富集/测序/扰动证据", "染色质相关tRNA和肿瘤表型"],
    ], [1.2, 1.7, 2.5, 2.2]),
    "position": make_position_map(ASSETS / "04_position.png"),
    "metrics": make_matrix(ASSETS / "05_metrics.png", "识别率与通读率需要按方法解释", ["方法", "指标", "主要影响因素"], [
        ["mim-tRNAseq", "mismatch rate, RT stop, coverage, 3'-CCA completeness", "修饰化学结构、RT酶、coverage、remap参数"],
        ["TRAC-Seq", "cleavage score, score ratio", "m7G还原/切割效率、背景切割、tRNA丰度"],
        ["Nanopore", "basecalling error, indel, termination, readthrough", "chemistry、basecaller、序列上下文、修饰体积"],
        ["LC-MS/MS", "nucleoside abundance, fragment mass, modification frequency", "消化方式、片段覆盖、定量标准品"],
        ["PANDORA/LIDAR", "capture gain, terminal state, fragment class", "末端化学、ligation bias、blocked termini"],
    ], [1.4, 2.9, 2.8]),
    "schema": make_flow(ASSETS / "99_schema.png", "tRNA修饰指纹库最小数据模型", [
        "Reference layer：species / tRNA gene / anticodon / mature CCA / genomic coordinate",
        "Modification layer：MODOMICS symbol / short name / chemical family / writer-eraser",
        "Evidence layer：method / pretreatment / signal / coverage / threshold / source",
        "Performance layer：识别率 / 通读率 / 错误谱 / 终止谱 / 化学切割分数",
        "Function layer：decoding / aminoacylation / stability / stress / disease phenotype",
    ]),
}


TOOLS = {
    "modomics": make_flow(ASSETS / "tool_modomics.png", "MODOMICS_Decoder：修饰字典导入", [
        "输入：MODOMICS unicode RNA sequence",
        "解析：symbol → short name，例如 7=m7G，?=m5C，Ѣ=m1A",
        "输出：position-symbol-modification table",
        "用途：提供先验修饰名称、符号和位置，不直接调用修饰",
    ]),
    "reference": make_flow(ASSETS / "tool_reference.png", "tRNA reference construction：参考序列库", [
        "输入：genomic tRNA predictions",
        "处理：去重复、合并、生成优化参考",
        "验证：reads mapping",
        "用途：提供坐标系统和多拷贝归属规则",
    ]),
    "sports": make_flow(ASSETS / "tool_sports.png", "SPORTS1.1：tsRNA/tRF注释", [
        "输入：small RNA或PANDORA-seq reads",
        "先验：genome/rRNA/tRNA/miRNA/piRNA/Rfam reference",
        "输出：fragment class、reads、length、annotation、mismatch summary",
        "用途：识别tsRNA来源和片段层表型",
    ]),
    "mim": make_flow(ASSETS / "tool_mim.png", "mim-tRNAseq：RT错配型修饰证据", [
        "先验：tRNA clusters和modification index",
        "过程：GSNAP SNP-tolerant alignment和deconvolution",
        "输出：coverage、expression、misincorporation、3'-CCA、SLAC",
        "用途：建立方法特异的RT错配/通读指纹",
    ]),
    "trac": make_flow(ASSETS / "tool_trac.png", "TRAC-Seq：化学切割规则", [
        "预处理：AlkB/D135S降低RT阻碍",
        "化学规则：NaBH4还原m7G，aniline切割",
        "输出：cleavage score和m7G abundance",
        "用途：m7G46的位点特异证据",
    ]),
    "trna004": make_flow(ASSETS / "tool_trna004.png", "tRNA004：纳米孔经验信号库", [
        "先验：MODOMICS已知修饰位点作为truth anchors",
        "比较：RNA002与RNA004 chemistry",
        "输出：43种修饰的basecalling error profiles",
        "用途：评估纳米孔识别率、错误率和上下文差异",
    ]),
    "tworead": make_flow(ASSETS / "tool_2read.png", "Nanopore 2-Read：通读和终止指纹", [
        "输入：ONT direct RNA reads",
        "检测：basecalling error、indel、polymerase read-through、termination",
        "输出：per-read/per-site candidate modification evidence",
        "用途：记录通读率和终止位点差异",
    ]),
    "galaxy": make_flow(ASSETS / "tool_galaxy.png", "Galaxy tRNAmod workflow：错配证据整理", [
        "先验：CCA mature tRNA reference和clustered reference",
        "策略：multimapper phasing保留一致错配模式",
        "输出：SAM/BAM/VCF-like evidence",
        "用途：规范化传统NGS错配证据",
    ]),
}


def src_lit(title, pmid):
    return f"文献：{title} PMID: {pmid}."


def src_tool(name, url):
    return f"GitHub工具：{name}. URL: {url}"


SLIDES = [
    ("tRNA修饰指纹库构建：修饰谱、检测信号与工具实现", "综合来源：MODOMICS 2023 update；本地文献库；GitHub工具仓库。", VISUALS["title"],
     ["汇报目标是说明tRNA修饰如何转化为可计算、可比较、可追溯的指纹证据。", "指纹库不等同于单一序列列表；不同工具使用的“指纹”形式不同。", "本PPT按修饰类型、检测信号、工具实现、表型证据和数据库结构展开。"]),
    ("指纹库在工具中的作用：先验列表只是其中一种形式", "综合来源：MODOMICS 2023 update；mim-tRNAseq、tRNA004、TRAC-Seq、SPORTS1.1、Galaxy tRNAmod等工具说明。", VISUALS["definition"],
     ["参考序列库和修饰字典属于先验信息。", "纳米孔错误谱、RT错配谱和化学切割规则属于方法特异指纹。", "功能表型证据用于解释修饰信号的生物学意义。"]),
    ("tRNA修饰覆盖范围：数量、子集与实验可识别集合", "文献/数据库：MODOMICS 2023 update, Nucleic Acids Research 2024; GitHub工具：rnabioco/tRNA004. URL: https://github.com/rnabioco/tRNA004", VISUALS["counts"],
     ["全量RNA修饰字典和tRNA修饰子集需要分开维护。", "实验可识别集合取决于测序方法、化学预处理和平台。", "第一版数据库应优先覆盖强证据、强表型、强工具支持的修饰。"]),
    ("主要修饰家族及其可观测信号", "综合来源：MODOMICS 2023 update；本地文献库中m7G、m1A、m3C、Ψ、U34修饰相关研究。", VISUALS["families"],
     ["修饰的化学类型决定其可能产生的RT、MS和纳米孔信号。", "U34和A37修饰对翻译解码最敏感。", "同一化学家族内部仍需按位点和tRNA上下文细分。"]),
    ("修饰位置决定检测信号和功能解释", "综合来源：tRNA结构与本地文献库中METTL6、PUS3、METTL1、lysidine、U34脱硫研究。", VISUALS["position"],
     ["U34主要影响codon recognition和通读/识别效率。", "G46、C32、38/39等位点更常连接结构稳定、writer选择性和RTD。", "数据库中的修饰记录必须绑定结构区域和标准位置。"]),
    ("识别率和通读率差异需要按方法存储", "综合来源：mim-tRNAseq、TRAC-Seq、tRNA004、Nanopore 2-Read、PANDORA/LIDAR相关资料。", VISUALS["metrics"],
     ["不同工具的识别率不是同一个物理量。", "RT通读率受修饰、酶、预处理和结构影响；Nanopore通读/终止受chemistry和basecaller影响。", "数据库应记录method、parameter、coverage和reference_version。"]),

    ("METTL1/m7G46：m7G-tRNA下降可导致翻译衰老", src_lit("Perturbation of METTL1-mediated tRNA N7-methylguanosine modification induces senescence and aging.", "38977661"), FIG_ALL / "38977661_41467_2024_49796_Fig4_HTML.jpg",
     ["文章逻辑：衰老中METTL1-WDR4下降，导致tRNA m7G46减少。", "方法步骤：TRAC-seq通过NaBH4/aniline切割m7G位点，计算cleavage score。", "主要结果：识别并验证一组METTL1调控的m7G-tRNAs。", "对指纹库的贡献：m7G46可采用cleavage_score、tRNA_abundance和METTL1-WDR4状态作为核心字段。"]),
    ("METTL1/m7G46：核糖体停顿连接SASP衰老表型", src_lit("Perturbation of METTL1-mediated tRNA N7-methylguanosine modification induces senescence and aging.", "38977661"), FIG_ALL / "38977661_41467_2024_49796_Fig5_HTML.jpg",
     ["假说验证：m7G-tRNA下降会造成特定codon翻译效率降低。", "方法步骤：puromycin、polysome profiling、Ribo-seq和RNC-qPCR。", "主要结果：m7G-tRNA解码codon的A-site ribosome occupancy升高。", "数据库字段：codon_occupancy、translation_efficiency、RTD、SASP/senescence phenotype。"]),
    ("m6A与mcm5s2U：tRNA修饰可读取mRNA修饰压力", src_lit("tRNA modifications tune m6A-dependent mRNA decay.", "40311619"), FIG_ALL / "40311619_nihms-2074120-f0005.jpg",
     ["文章逻辑：CDS中的m6A会降低相应codon的翻译效率并促进mRNA decay。", "核心假说：tRNA anticodon loop中的mcm5s2U可抵消m6A codon造成的解码障碍。", "主要结果：mcm5s2U促进m6A-modified codons的有效解码。", "数据库字段：mRNA_mod_context、decoded_codon、tRNA_U34_state、decay_effect。"]),
    ("m6A与mcm5s2U：修饰互作影响核糖体解码和癌症通路", src_lit("tRNA modifications tune m6A-dependent mRNA decay.", "40311619"), FIG_ALL / "40311619_nihms-2074120-f0007.jpg",
     ["方法步骤：翻译报告、Ribo-seq、通路扰动和癌症相关分析。", "主要结果：mRNA m6A和tRNA mcm5s2U共同调节ribosomal decoding。", "讨论要点：tRNA修饰指纹应允许记录跨RNA类别的相互作用。", "数据库字段：pan-epitranscriptomic_interaction、translation_decay_axis、tumor_association。"]),
    ("氧化脱硫：U34修饰可在应激条件下转换状态", src_lit("Translational regulation by oxidative desulfuration of tRNA modifications.", "41807381"), FIG / "PMC12976133_41467_2026_70126_Fig1_HTML.jpg",
     ["文章逻辑：xm5s2U34增强wobble解码，但硫代基易受氧化影响。", "核心假说：氧化应激可将xm5s2U34转为xm5h2U34。", "方法步骤：LC-MS、CID和co-injection验证脱硫产物。", "数据库字段：precursor_modification、converted_modification、redox_state、MS_fragment。"]),
    ("氧化脱硫：修饰转换降低氨酰化和A-site识别", src_lit("Translational regulation by oxidative desulfuration of tRNA modifications.", "41807381"), FIG / "PMC12976133_41467_2026_70126_Fig2_HTML.jpg",
     ["方法步骤：spike-in控制、体外翻译、氨酰化、A-site binding和cryo-EM。", "主要结果：mcm5h2U降低tRNALys/Gln/Glu氨酰化并削弱codon recognition。", "表型含义：氧化应激可通过U34化学转换调节codon-specific translation。", "数据库字段：readthrough_rate、recognition_rate、aminoacylation_effect、stress_condition。"]),
    ("冠状病毒：病毒codon usage依赖宿主tRNA修饰景观", src_lit("Coronaviruses reprogram the tRNA epitranscriptome to favor viral protein expression.", "41714626"), FIG / "PMC13031925_41467_2026_69700_Fig1_HTML.jpg",
     ["文章逻辑：冠状病毒基因组富含A/U-ending suboptimal codons。", "核心假说：感染诱导宿主tRNA修饰重编程以提高病毒蛋白翻译。", "方法步骤：RSCU分析、LC-MS/MS、mim-tRNAseq和酶扰动。", "数据库字段：condition、codon_usage_pressure、required_tRNA_modification、viral_translation_output。"]),
    ("冠状病毒：感染改变anticodon loop修饰", src_lit("Coronaviruses reprogram the tRNA epitranscriptome to favor viral protein expression.", "41714626"), FIG / "PMC13031925_41467_2026_69700_Fig3_HTML.jpg",
     ["主要结果：SARS-CoV-2和HCoV-OC43感染改变I、Q、mcm5U/mcm5s2U、m5C/f5C等修饰。", "功能验证：ELP3、QTRT1、NSUN2等酶通路扰动影响病毒NP蛋白表达。", "讨论要点：指纹库需要支持condition-specific modification profile。", "数据库字段：infection_time、enzyme_expression、modification_fold_change、viral_protein_readout。"]),
    ("PUS3/Ψ38-39：writer选择性决定同类修饰的位点特异性", src_lit("The molecular basis of tRNA selectivity by human pseudouridine synthase 3.", "38996458"), FIG_ALL / "38996458_gr2.jpg",
     ["文章逻辑：Ψ很常见，但PUS3只修饰特定tRNA的38/39位点。", "方法步骤：PUS3 apo和tRNA-bound cryo-EM、突变验证、Pseudo-seq。", "主要结果：PUS3二聚体定位tRNA并把目标U放入活性位点。", "数据库字段：writer、target_position、substrate_selectivity、structural_evidence。"]),
    ("METTL6/m3C32：SerRS作为底物选择因子", src_lit("Structural basis of tRNA recognition by the m3C RNA methyltransferase METTL6 in complex with SerRS seryl-tRNA synthetase.", "38918637"), FIG / "PMC11479938_41594_2024_1341_Fig1_HTML.jpg",
     ["文章逻辑：m3C32影响翻译精度，但METTL6选择tRNASer的机制不清楚。", "方法步骤：cryo-EM、体外甲基化、mutagenesis和MS。", "主要结果：SerRS增强METTL6活性并参与tRNASer底物选择。", "数据库字段：cofactor_or_selector、associated_A37_modification、structure_state、m3C32_evidence。"]),
    ("lysidine/k2C34：单个wobble修饰决定AUA识别并避开AUG", src_lit("Structures of the ribosome bound to EF-Tu-isoleucine tRNA elucidate the mechanism of AUG avoidance.", "38538914"), FIG_ALL / "38538914_nihms-2030058-f0002.jpg",
     ["文章逻辑：tRNAIle CAU必须识别AUA但避免误读AUG。", "核心假说：C34的lysidine改变wobble位点碱基几何。", "方法步骤：70S ribosome-EF-Tu-Ile-tRNA cryo-EM。", "数据库字段：codon_specificity、avoidance_codon、structural_decoding_evidence、wobble_modification。"]),
    ("tsRNA-Glu-CTC：修饰型tRNA片段具有更强表型效应", src_lit("A cholesterol-responsive hepatic tRNA-derived small RNA regulates cholesterol homeostasis and atherosclerosis development.", "41398161"), FIG_ALL / "41398161_41467_2025_67387_Fig9_HTML.jpg",
     ["文章逻辑：tsRNA可能是带修饰的功能性小RNA，而非单纯降解产物。", "方法步骤：PANDORA-seq识别tsRNA，MLC-seq解析内源tsRNA修饰。", "主要结果：修饰型tsRNA-Glu-CTC比未修饰合成体更强地影响胆固醇和肝脂质。", "数据库字段：fragment_start/end、source_tRNA、modification_pattern、lipid_phenotype。"]),
    ("ALKB-1/m1A：eraser失活引起翻译和线粒体表型", src_lit("ALKB-1-dependent tRNA methylation is required for efficient paternal mitochondrial elimination.", "41611679"), FIG_ALL / "41611679_41467_2026_68813_Fig2_HTML.jpg",
     ["文章逻辑：tRNA methylation是否影响paternal mitochondrial elimination尚不明确。", "方法步骤：ALKB-1 D247A/RNAi、LC-MS、m1A-seq、RNC-seq和proteomics。", "主要结果：tRNA m1A升高，线粒体蛋白稳态和ROS通路改变。", "数据库字段：eraser_state、m1A_level、translation_imbalance、PME_phenotype。"]),
    ("vRNAP-associated tRNA：缺失修饰也可能构成功能指纹", src_lit("tRNA as an assembly chaperone for a macromolecular transcription-processing complex.", "40908366"), FIG_ALL / "40908366_41594_2025_1653_Fig6_HTML.jpg",
     ["文章逻辑：poxvirus vRNAP装配需要特定tRNAGln/Arg作为assembly chaperone。", "方法步骤：装配重构、gel shift、LC-MS/MS和cryo-EM。", "主要结果：特定修饰图案及缺失mcm5s2U34参与tRNA选择。", "数据库字段：absence_of_modification、non_translational_function、viral_RNP_assembly。"]),
    ("NAT10/ac4C：染色质相关tRNA连接转移表型", src_lit("NAT10 promotes cancer metastasis by modulating p300/CBP activity through chromatin-associated tRNA.", "41344331"), FIG_ALL / "41344331_nihms-2122374-f0003.jpg",
     ["文章逻辑：NAT10作为ac4C writer，其促转移机制需要解释。", "方法步骤：NAT10扰动、RNA ac4C、chromatin-associated tRNA和p300/CBP分析。", "主要结果：染色质相关tRNA ac4C变化影响p300/CBP和enhancer组织。", "数据库字段：subcellular_context、chromatin_associated_tRNA、ac4C、metastasis_phenotype。"]),
    ("LIDAR：末端修饰会改变小RNA捕获率", src_lit("A ligation-independent sequencing method reveals tRNA-derived RNAs with blocked 3' termini.", "39096899"), make_flow(ASSETS / "paper_lidar.png", "LIDAR检测blocked 3' termini tDRs", ["quasi-random priming", "template switching", "ligation-independent library construction", "capture RNAs with blocked 3' ends", "store terminal_state and capture_bias"]),
     ["文章逻辑：ligation-dependent small RNA-seq会漏检3'端阻断的RNA。", "方法步骤：quasi-random priming结合template switching。", "主要结果：检测到更多tRNA-derived RNAs，包括blocked 3' termini物种。", "数据库字段：terminal_state、library_method、capture_bias、tDR_class。"]),

    ("工具：MODOMICS_Decoder提供先验修饰字典", src_tool("monimaanam/Modomics_Decoder", "https://github.com/monimaanam/Modomics_Decoder"), TOOLS["modomics"],
     ["作用方式：先验修饰符号和短名列表。", "它不是caller，不评估识别率或通读率。", "适合用于数据库初始化：symbol、short_name、position。"]),
    ("工具：tRNA_reference_construction提供参考序列库", src_tool("FlorianPichot/tRNA_reference_construction", "https://github.com/FlorianPichot/tRNA_reference_construction"), TOOLS["reference"],
     ["作用方式：先验tRNA reference和坐标体系。", "解决多拷贝tRNA、mature CCA和reference version问题。", "它不调用修饰，但决定修饰信号归属。"]),
    ("工具：SPORTS1.1提供tsRNA/tRF注释入口", src_tool("junchaoshi/sports1.1", "https://github.com/junchaoshi/sports1.1"), TOOLS["sports"],
     ["作用方式：依赖先验参考库进行小RNA映射和分类。", "输出片段来源、长度、reads和annotation。", "mismatch summary可作为候选修饰线索，但不能直接等同于修饰率。"]),
    ("工具：mim-tRNAseq提供RT错配型修饰证据", src_tool("nedialkova-lab/mim-tRNAseq", "https://github.com/nedialkova-lab/mim-tRNAseq"), TOOLS["mim"],
     ["作用方式：先验modification index + 实验错配信号。", "输出misincorporation、coverage、expression、3'-CCA和SLAC。", "识别率取决于RT通读、coverage、修饰类型和比对参数。"]),
    ("工具：TRAC-Seq提供m7G化学切割规则", src_tool("rnabioinfor/TRAC-Seq", "https://github.com/rnabioinfor/TRAC-Seq"), TOOLS["trac"],
     ["作用方式：不是普通序列列表，而是m7G特异化学反应规则。", "核心指标是cleavage score和score ratio。", "适合记录m7G46位点级证据。"]),
    ("工具：tRNA004提供纳米孔经验错误谱", src_tool("rnabioco/tRNA004", "https://github.com/rnabioco/tRNA004"), TOOLS["trna004"],
     ["作用方式：已知修饰位点作为truth anchors，建立basecalling error profiles。", "它用于解释RNA004/RNA002下不同修饰的识别率差异。", "适合作为纳米孔指纹库的benchmark层。"]),
    ("工具：Nanopore 2-Read强调通读率和终止位点", src_tool("AteeshaNegi/nanopore-2Read-trna-pipeline", "https://github.com/AteeshaNegi/nanopore-2Read-trna-pipeline"), TOOLS["tworead"],
     ["作用方式：经验信号检测，不是单纯先验序列列表。", "关注basecalling error、indel、read-through和termination events。", "适合记录per-read/per-site异质性。"]),
    ("工具：Galaxy tRNAmod workflow整理传统NGS错配证据", src_tool("jfallmann/tRNA_ngs_mod_map_call_galaxy", "https://github.com/jfallmann/tRNA_ngs_mod_map_call_galaxy"), TOOLS["galaxy"],
     ["作用方式：先验参考库 + 多重比对错配模式过滤。", "强调tRNA多拷贝造成的mapping歧义。", "适合抽象为multi_mapping_policy和mismatch_evidence字段。"]),
    ("数据库结构：修饰、工具信号和表型分层保存", "综合来源：上述文献和GitHub工具。", VISUALS["schema"],
     ["Reference layer解决位置和归属。", "Modification layer解决修饰名、化学类型和writer/eraser。", "Evidence与Performance layer解决不同工具的识别率、通读率和错误谱。", "Function layer保存翻译、稳定性、应激和疾病表型。"]),
    ("第一版数据库优先收录对象", "综合来源：PMID 38977661、40311619、41807381、41714626、38918637、38996458、38538914及相关工具仓库。", make_matrix(ASSETS / "priority.png", "第一版优先收录修饰与理由", ["修饰/位点", "核心文献或工具", "优先原因"], [
        ["m7G46", "METTL1/TRAC-seq", "化学切割信号明确，连接衰老和翻译停顿"],
        ["mcm5s2U34/mcm5h2U34", "m6A decay、氧化脱硫、tRNA004", "连接解码效率、应激和纳米孔错误谱"],
        ["m1A", "ALKB-1/PME", "writer/eraser状态和线粒体表型清楚"],
        ["m3C32", "METTL6-SerRS", "结构选择性和底物选择机制清楚"],
        ["Ψ38/39", "PUS3", "同类修饰的位点选择性代表"],
        ["Q/I/m5C/f5C/k2C", "冠状病毒、lysidine、Nanopore benchmark", "决定codon preference和识别率差异"],
    ], [1.8, 2.6, 3.2]),
     ["第一版不追求一次覆盖全部>100种tRNA修饰。", "优先选择有位点证据、工具信号和功能表型的修饰。", "后续通过MODOMICS字典扩展全量修饰清单。"]),
    ("结论：指纹库应是证据图谱，而不是静态修饰表", "综合来源：本PPT全部文献和工具。", make_flow(ASSETS / "final.png", "汇报结论", [
        "先验序列列表只解决参考和坐标问题",
        "修饰字典解决命名和化学类别问题",
        "实验工具提供方法特异的识别率、通读率和错误谱",
        "功能实验解释修饰信号对应的翻译和表型后果",
        "最终数据库需要保留来源、图号、参数和证据等级",
    ]),
     ["tRNA修饰指纹库的核心价值是跨方法整合。", "同一修饰在不同工具中的信号不可直接等价，必须方法分层。", "建议下一步把现有文献和工具输出整理为schema原型。"]),
]


def add_text(slide, x, y, w, h, text, size=18, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)
    return shape


def add_bullets(slide, x, y, w, h, bullets, size=12.2):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.space_after = Pt(5)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0, 0, 0)
    return shape


def fit_pic(slide, image, x, y, w, h):
    image = Path(image)
    with Image.open(image) as im:
        iw, ih = im.size
    box_ratio = w / h
    im_ratio = iw / ih
    if im_ratio > box_ratio:
        dw = w
        dh = w / im_ratio
    else:
        dh = h
        dw = h * im_ratio
    px = x + (w - dw) / 2
    py = y + (h - dh) / 2
    pic = slide.shapes.add_picture(str(image), Inches(px), Inches(py), width=Inches(dw), height=Inches(dh))
    pic.line.color.rgb = RGBColor(0, 0, 0)
    pic.line.width = Pt(0.75)
    return pic


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, (title, source, image, bullets) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        add_text(slide, 0.32, 0.14, 12.7, 0.40, title, 19.5, True)
        add_text(slide, 0.36, 0.58, 12.35, 0.36, source, 7.8)
        fit_pic(slide, image, 0.45, 1.03, 7.25, 5.90)
        add_bullets(slide, 7.95, 1.05, 4.95, 5.70, bullets, 12.2)
        add_text(slide, 12.45, 7.15, 0.55, 0.18, str(i), 8, align=PP_ALIGN.RIGHT)

    prs.save(PPTX)

    reopened = Presentation(str(PPTX))
    bad = []
    missing_pics = []
    for i, slide in enumerate(reopened.slides, 1):
        if not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes):
            missing_pics.append(i)
        for sh in slide.shapes:
            if hasattr(sh, "text") and "????" in sh.text:
                bad.append(i)
                break
    with zipfile.ZipFile(PPTX) as zf:
        media = [n for n in zf.namelist() if n.startswith("ppt/media/")]
        slide_xml = [n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    MANIFEST.write_text(
        "# Asset Manifest\n\n"
        + f"PPTX: `{PPTX}`\n\n"
        + "\n".join(f"- Slide {i}: `{image}` | {source}" for i, (_, source, image, _) in enumerate(SLIDES, 1))
        + "\n",
        encoding="utf-8",
    )
    result = "PASS" if len(reopened.slides) == len(SLIDES) and not bad and not missing_pics else "CHECK"
    QA.write_text(
        "# QA Report\n\n"
        + f"- PPTX: `{PPTX}`\n"
        + f"- Slides: {len(reopened.slides)} / expected {len(SLIDES)}\n"
        + f"- Embedded media files: {len(media)}\n"
        + f"- Slide XML files: {len(slide_xml)}\n"
        + f"- Slides without picture objects: {missing_pics or 'none'}\n"
        + f"- Slides containing literal `????`: {bad or 'none'}\n"
        + "- Style: formal report tone, white background, black text, one visual per slide, source line on every slide.\n"
        + f"\nValidation result: {result}\n",
        encoding="utf-8",
    )
    print(f"WROTE {PPTX}")
    print(f"WROTE {QA}")
    print(f"slides={len(reopened.slides)} media={len(media)} missing={missing_pics} bad={bad}")


if __name__ == "__main__":
    build()
