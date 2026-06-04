# -*- coding: utf-8 -*-
from pathlib import Path
import textwrap
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


BASE = Path(r"D:\Project\tRNA research\tRNA mod-fingerprint database\downloaded_resources")
FIG_DIR = BASE / "figures"
OUT_DIR = BASE / "analysis_outputs" / "paper2ppt_output"
ASSET_DIR = OUT_DIR / "assets_utf8"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

PPTX_PATH = OUT_DIR / "trna_mod_fingerprint_database_reading_deck_utf8_fixed.pptx"
QA_PATH = OUT_DIR / "pptx_qa_report_utf8_fixed.md"
MANIFEST_PATH = OUT_DIR / "asset_manifest_utf8_fixed.md"


def load_font(size):
    for font_name in ["msyh.ttc", "simhei.ttf", "arial.ttf"]:
        try:
            return ImageFont.truetype(font_name, size)
        except Exception:
            pass
    return ImageFont.load_default()


FONT_TITLE = load_font(36)
FONT_TEXT = load_font(25)
FONT_SMALL = load_font(18)


def wrap_zh(text, width=34):
    lines = []
    for seg in text.split("\n"):
        if len(seg) <= width:
            lines.append(seg)
        else:
            lines.extend(textwrap.wrap(seg, width=width, break_long_words=True, replace_whitespace=False))
    return lines


def diagram(path, title, steps, note=None):
    width, height = 1400, 850
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, width - 1, height - 1], outline="black", width=3)
    draw.text((48, 32), title, fill="black", font=FONT_TITLE)

    top = 150
    box_h = 85
    gap = 34
    x = 90
    box_w = width - 180
    for i, step in enumerate(steps):
        y = top + i * (box_h + gap)
        draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=7, outline="black", width=3, fill="white")
        yy = y + 13
        for line in wrap_zh(step, 42)[:2]:
            draw.text((x + 28, yy), line, fill="black", font=FONT_TEXT)
            yy += 31
        if i < len(steps) - 1:
            cx = width // 2
            ay = y + box_h + 7
            draw.line([cx, ay, cx, ay + gap - 16], fill="black", width=4)
            draw.polygon([(cx - 10, ay + gap - 16), (cx + 10, ay + gap - 16), (cx, ay + gap - 2)], fill="black")

    if note:
        draw.text((50, height - 70), note, fill="black", font=FONT_SMALL)
    img.save(path, quality=95)
    return path


PROJECT_DIAGRAMS = {
    "sports": diagram(
        ASSET_DIR / "github_sports11.png",
        "SPORTS1.1：small RNA / tsRNA 注释流程",
        [
            "FASTQ / FASTA / SRA reads",
            "adapter trimming + length filtering",
            "Bowtie mapping 到 genome、rRNA、tRNA、miRNA、piRNA、Rfam",
            "tRNA-derived fragments：5 end / 3 end / internal + reads",
            "输出 annotation table、length plots、mismatch summary",
        ],
        "用于 tsRNA/tRF 来源识别；mismatch 只作为修饰候选线索。",
    ),
    "mim": diagram(
        ASSET_DIR / "github_mimtrnaseq.png",
        "mim-tRNAseq：misincorporation 修饰指纹",
        [
            "trimmed tRNA-seq FASTQ",
            "cluster tRNAs + build modification index",
            "SNP-tolerant GSNAP alignment + deconvolution",
            "coverage、expression、3'-CCA completeness",
            "modification-induced mismatch signatures + SLAC crosstalk",
        ],
        "最适合纳入 RT/misincorporation 型 tRNA 修饰指纹。",
    ),
    "trna004": diagram(
        ASSET_DIR / "github_trna004.png",
        "tRNA004：Nanopore RNA004 benchmark",
        [
            "direct RNA nanopore tRNA reads",
            "RNA002 vs RNA004 chemistry comparison",
            "basecalling errors 映射到 MODOMICS 位点和 tRNA 结构",
            "43 类 RNA modifications across sequence contexts",
            "输出 modification-specific nanopore error profiles",
        ],
        "适合作为 nanopore 修饰指纹参考集和模型评估集。",
    ),
    "trac": diagram(
        ASSET_DIR / "github_tracseq.png",
        "TRAC-Seq：m7G cleavage-based profiling",
        [
            "AlkB / D135S pretreatment 降低 RT blockers",
            "NaBH4 reduction of m7G",
            "aniline cleavage at reduced m7G sites",
            "small RNA library sequencing",
            "输出 G46 cleavage score + m7G-tRNA abundance",
        ],
        "GitHub 文档极少，方法细节主要来自论文。",
    ),
    "galaxy": diagram(
        ASSET_DIR / "github_galaxy_trnamod.png",
        "Galaxy tRNA modification mapping workflow",
        [
            "tRNAscan references + CCA mature tRNA FASTA",
            "read preprocessing and mapping",
            "multimapper phasing：保留一致 misincorporation pattern",
            "GATK / samtools / Picard wrapper 处理 BAM 和 variants",
            "输出 tRNAmod-compatible SAM/BAM/VCF-like evidence",
        ],
        "可作为传统 NGS 修饰映射 workflow 参考；依赖较旧。",
    ),
    "modomics": diagram(
        ASSET_DIR / "github_modomics_decoder.png",
        "Modomics_Decoder：MODOMICS 符号解析",
        [
            "输入 MODOMICS unicode RNA sequence",
            "symbol dictionary：如 7=m7G，?=m5C，Ѣ=m1A，Щ=m3C",
            "enumerate sequence positions",
            "mark modified bases and short names",
            "输出 position-symbol-modification table",
        ],
        "适合做参考修饰表导入器，需要补充物种和坐标标准。",
    ),
    "tworead": diagram(
        ASSET_DIR / "github_2read.png",
        "Nanopore 2-Read tRNA modification pipeline",
        [
            "POD5 / BAM / FASTQ from ONT direct RNA",
            "Dorado basecalling + BWA-MEM alignment",
            "per-read / per-site mismatch and indel analysis",
            "polymerase read-through and termination pattern detection",
            "输出 candidate modification calls with confidence signals",
        ],
        "与 tRNA004 互补：更强调 termination/read-through 指纹。",
    ),
    "reference": diagram(
        ASSET_DIR / "github_reference_construction.png",
        "tRNA reference construction",
        [
            "输入 genomic tRNA predictions",
            "remove duplicates / merge related tRNA records",
            "build optimized mature and genomic references",
            "validate by trimming、Bowtie2 mapping、samtools summaries",
            "输出 database-ready tRNA reference versions",
        ],
        "解决多拷贝 tRNA reference 和 reads 归属问题。",
    ),
}


SLIDES = [
    {
        "title": "1. METTL6-SerRS 结构解释 m3C32 指纹",
        "source": "PMID 38918637 | Fig. 1 | Nature Structural & Molecular Biology 2024",
        "image": FIG_DIR / "PMC11479938_41594_2024_1341_Fig1_HTML.jpg",
        "bullets": [
            "图中把 METTL6、SerRS、tRNASer 放在同一 cryo-EM 复合物中，说明 m3C32 不是孤立位点。",
            "SerRS 作为底物选择因子增强 METTL6 对 tRNASer 的甲基化。",
            "指纹库应记录：m3C32、METTL6、SerRS 依赖性、A37 预修饰和 variable arm 识别。",
            "证据类型：结构 + 体外甲基化 + MS/突变验证。",
        ],
    },
    {
        "title": "2. METTL1-WDR4/m7G46 下降连接衰老指纹",
        "source": "PMID 38977661 | Fig. 1 | Nature Communications 2024",
        "image": FIG_DIR / "PMC11231295_41467_2024_49796_Fig1_HTML.jpg",
        "bullets": [
            "图中显示 senescence / aging 过程中 METTL1-WDR4 和 RNA/tRNA m7G 水平下降。",
            "m7G46 影响 tRNA 稳定性、RTD、核糖体停顿和翻译效率。",
            "数据库字段应包括 cleavage score、tRNA abundance、METTL1/WDR4 状态、衰老表型。",
            "TRAC-seq 是 m7G46 的核心测序证据。",
        ],
    },
    {
        "title": "3. tRNA 可作为病毒 vRNAP 装配伴侣",
        "source": "PMID 40908366 | Fig. 2 | Nature Structural & Molecular Biology 2025",
        "image": FIG_DIR / "PMC12618233_41594_2025_1653_Fig2_HTML.jpg",
        "bullets": [
            "图中展示 early PIC / vRNAP 的 tRNA-chaperoned 装配中间体。",
            "tRNAGln/Arg 的构象和修饰状态决定其是否适合装配，而不仅影响翻译。",
            "该论文提醒：修饰“缺失”也可能是功能性指纹，例如缺少 mcm5s2U34。",
            "数据库需要支持 non-translational function 和 viral RNP assembly 标签。",
        ],
    },
    {
        "title": "4. PANDORA-seq 锁定 cholesterol-responsive tsRNA",
        "source": "PMID 41398161 | Fig. 1 | Nature Communications 2025",
        "image": FIG_DIR / "PMC12706008_41467_2025_67387_Fig1_HTML.jpg",
        "bullets": [
            "图中比较传统 RNA-seq 与 PANDORA-seq，显示 PANDORA-seq 更能捕获肝脏 tsRNA。",
            "tsRNA-Glu-CTC 是高丰度、胆固醇响应的小 RNA。",
            "内源 tsRNA 带修饰，MLC-seq 可解析修饰型 tsRNA，且功能强于未修饰合成体。",
            "指纹库应把 tsRNA fragment 作为独立实体：来源 tRNA、端点、修饰、组织、表型。",
        ],
    },
    {
        "title": "5. ALKB-1/m1A 影响线粒体遗传质量控制",
        "source": "PMID 41611679 | Fig. 2 | Nature Communications 2026",
        "image": FIG_DIR / "PMC12957315_41467_2026_68813_Fig2_HTML.jpg",
        "bullets": [
            "图中显示 ALKB-1 失活后 tRNA m1A 增加，并伴随翻译与线粒体蛋白稳态改变。",
            "m1A 异常引发 ROS、SKN-1、UPRmt，最终延迟 paternal mitochondrial elimination。",
            "指纹库应区分 writer 与 eraser，并记录 demethylation-dependent 状态。",
            "功能字段建议加入 mitochondrial proteostasis、PME、fertility/embryo viability。",
        ],
    },
    {
        "title": "6. 冠状病毒重编程宿主 tRNA 修饰景观",
        "source": "PMID 41714626 | Fig. 3 | Nature Communications 2026",
        "image": FIG_DIR / "PMC13031925_41467_2026_69700_Fig3_HTML.jpg",
        "bullets": [
            "图中用 LC-MS/MS 展示 SARS-CoV-2 与 HCoV-OC43 感染后的 tRNA 修饰变化。",
            "关键修饰包括 I、Q、mcm5U/mcm5s2U、m5C/f5C，主要服务于病毒偏好的 suboptimal codons。",
            "mim-tRNAseq 与 LC-MS/MS 可互补建立动态修饰指纹。",
            "数据库需要把 condition=viral infection/stress 与 codon usage pressure 关联。",
        ],
    },
    {
        "title": "7. 氧化脱硫把 mcm5s2U34 转成动态翻译指纹",
        "source": "PMID 41807381 | Fig. 1 | Nature Communications 2026",
        "image": FIG_DIR / "PMC12976133_41467_2026_70126_Fig1_HTML.jpg",
        "bullets": [
            "图中说明 U34 wobble 2-thiouridine derivatives 可在氧化条件下脱硫为 h2U derivatives。",
            "mcm5h2U34 会降低氨酰化效率和 A-site codon recognition。",
            "这是“同一位点修饰状态转换”的典型动态指纹。",
            "建议字段：redox_state、precursor_modification、converted_modification、aminoacylation effect。",
        ],
    },
    {
        "title": "8. SPORTS1.1 用于 tsRNA/tRF 来源注释",
        "source": "GitHub: junchaoshi/sports1.1",
        "image": PROJECT_DIAGRAMS["sports"],
        "bullets": [
            "适合处理 PANDORA-seq 或 small RNA-seq，识别 tRNA 5 end、3 end、internal fragments。",
            "输出 reads、length、Match_Genome、Annotation 和长度分布。",
            "可把 tsRNA-Glu-CTC 这类 fragment 接入指纹库。",
            "注意：mismatch 只能提示候选修饰，不能直接作为修饰定论。",
        ],
    },
    {
        "title": "9. mim-tRNAseq 产出 RT/misincorporation 修饰指纹",
        "source": "GitHub: nedialkova-lab/mim-tRNAseq",
        "image": PROJECT_DIAGRAMS["mim"],
        "bullets": [
            "自动完成 tRNA 聚类、modification index、SNP-tolerant alignment 和 transcript-level deconvolution。",
            "核心输出是 coverage、expression、3'-CCA completeness、misincorporation signatures。",
            "可用 SLAC 分析修饰-修饰、修饰-氨酰化之间的 single-read crosstalk。",
            "适合作为指纹库中 sequencing-inferred modification evidence 的主线工具。",
        ],
    },
    {
        "title": "10. tRNA004 提供 Nanopore RNA004 修饰 benchmark",
        "source": "GitHub: rnabioco/tRNA004",
        "image": PROJECT_DIAGRAMS["trna004"],
        "bullets": [
            "项目比较 RNA002 与 RNA004 chemistry 在直接 tRNA 测序中的表现。",
            "把 basecalling error data 映射到 MODOMICS 修饰位点和 tRNA 结构。",
            "覆盖 40+ / 43 类 RNA 修饰，是纳米孔修饰指纹的参考训练集。",
            "数据库应记录 chemistry、basecaller、sequence context、coverage、error profile。",
        ],
    },
    {
        "title": "11. TRAC-Seq 是 m7G46 的切割型指纹方法",
        "source": "GitHub: rnabioinfor/TRAC-Seq + PMID 38977661 方法",
        "image": PROJECT_DIAGRAMS["trac"],
        "bullets": [
            "通过 NaBH4 还原 m7G，再用 aniline 在还原位点切割。",
            "测序后以 cleavage score 定位 m7G 位点并估算修饰变化。",
            "尤其适合 METTL1-WDR4/m7G46 体系。",
            "仓库 README 很少，数据库实现应以论文方法和补充材料为准。",
        ],
    },
    {
        "title": "12. Galaxy workflow 组织传统 NGS 修饰映射",
        "source": "GitHub: jfallmann/tRNA_ngs_mod_map_call_galaxy",
        "image": PROJECT_DIAGRAMS["galaxy"],
        "bullets": [
            "提供 addCCA、clustering、multimapperPhasing、GATK wrapper 等工具。",
            "多重比对 reads 只有在 misincorporation pattern 一致时才保留。",
            "有助于把传统短读长 mismatch/variant 信号标准化。",
            "局限是依赖 GATK 3.6 和 Galaxy 环境，现代化需要重写。",
        ],
    },
    {
        "title": "13. Modomics_Decoder 可导入参考修饰符号",
        "source": "GitHub: monimaanam/Modomics_Decoder",
        "image": PROJECT_DIAGRAMS["modomics"],
        "bullets": [
            "把 MODOMICS unicode sequence 转成 position-symbol-short name 表。",
            "例如 7=m7G，?=m5C，Ѣ=m1A，Щ=m3C。",
            "适合作为 reference modification import 的轻量工具。",
            "仍需补充物种、tRNA 名称、标准坐标、证据来源和 MODOMICS 版本。",
        ],
    },
    {
        "title": "14. Nanopore 2-Read 强调 termination/read-through 指纹",
        "source": "GitHub: AteeshaNegi/nanopore-2Read-trna-pipeline",
        "image": PROJECT_DIAGRAMS["tworead"],
        "bullets": [
            "思路是用 ONT basecalling error、polymerase read-through/error 和 stop sites 识别修饰。",
            "强调 per-read、per-site、single-nucleotide resolution。",
            "适合与 tRNA004 benchmark 合并，形成纳米孔指纹层。",
            "需要进一步检查脚本完整度、示例数据和统计阈值。",
        ],
    },
    {
        "title": "15. tRNA reference construction 解决归属和版本问题",
        "source": "GitHub: FlorianPichot/tRNA_reference_construction",
        "image": PROJECT_DIAGRAMS["reference"],
        "bullets": [
            "从 genomic tRNA predictions 构建 non-duplicated、merged、optimized references。",
            "通过 reads mapping 验证参考质量。",
            "这不是修饰 caller，但决定后续 reads 和修饰位点如何归属。",
            "指纹库应固定 reference_version，并保留 mature CCA / genomic / merged 坐标映射。",
        ],
    },
]


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=14):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(7)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0, 0, 0)
    return box


def fit_image(slide, img_path, x, y, w, h):
    img_path = Path(img_path)
    with Image.open(img_path) as img:
        iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        disp_w = w
        disp_h = w / img_ratio
    else:
        disp_h = h
        disp_w = h * img_ratio
    px = x + (w - disp_w) / 2
    py = y + (h - disp_h) / 2
    pic = slide.shapes.add_picture(str(img_path), Inches(px), Inches(py), width=Inches(disp_w), height=Inches(disp_h))
    pic.line.color.rgb = RGBColor(0, 0, 0)
    pic.line.width = Pt(0.75)
    return pic


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, slide_data in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        add_textbox(slide, 0.35, 0.18, 12.65, 0.45, slide_data["title"], size=22, bold=True)
        add_textbox(slide, 0.38, 0.67, 12.2, 0.25, slide_data["source"], size=9)
        fit_image(slide, slide_data["image"], 0.45, 1.05, 7.25, 5.85)
        add_bullets(slide, 7.95, 1.08, 4.95, 5.65, slide_data["bullets"], size=14)
        add_textbox(slide, 12.55, 7.12, 0.45, 0.2, str(idx), size=8, align=PP_ALIGN.RIGHT)

    prs.save(PPTX_PATH)

    reopened = Presentation(str(PPTX_PATH))
    slide_count = len(reopened.slides)
    bad_text = []
    for i, slide in enumerate(reopened.slides, 1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and "????" in shape.text:
                bad_text.append(i)
                break

    with zipfile.ZipFile(PPTX_PATH) as zf:
        names = zf.namelist()
        media = [n for n in names if n.startswith("ppt/media/")]
        slide_xml = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]

    MANIFEST_PATH.write_text(
        "# Asset Manifest\n\n"
        + f"PPTX: `{PPTX_PATH}`\n\n"
        + "\n".join(f"- Slide {i}: `{s['image']}` | {s['source']}" for i, s in enumerate(SLIDES, 1))
        + "\n",
        encoding="utf-8",
    )

    result = "PASS" if slide_count == len(SLIDES) and len(media) >= len(SLIDES) and not bad_text else "CHECK"
    QA_PATH.write_text(
        "# PPTX QA Report\n\n"
        + f"- Output: `{PPTX_PATH}`\n"
        + f"- Slide count after reopen: {slide_count}\n"
        + f"- Expected slide count: {len(SLIDES)}\n"
        + f"- Package slide XML files: {len(slide_xml)}\n"
        + f"- Embedded media files: {len(media)}\n"
        + f"- Slides containing literal `????`: {bad_text or 'none'}\n"
        + "- Style check: white background, black text, one image per resource slide, no decorative theme.\n"
        + f"\nValidation result: {result}\n",
        encoding="utf-8",
    )

    print(f"WROTE {PPTX_PATH}")
    print(f"WROTE {QA_PATH}")
    print(f"slides={slide_count} media={len(media)} bad_question_mark_slides={bad_text}")


if __name__ == "__main__":
    build()
